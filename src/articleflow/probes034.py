"""Independent synthetic compiler regression cases, never selected by the repair LLM."""
import json,random
from pathlib import Path
from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Pt,RGBColor
from .agent023 import extract_input,build_package
from .rules023 import semantic_spans
from .agent028 import restyle


def create(folder,seed=34001):
    folder=Path(folder);folder.mkdir(parents=True,exist_ok=False);rng=random.Random(seed);cases=[]
    for index in range(7):
        local=folder/str(index);local.mkdir();break_first=index in {1,3};blocked=index==5
        roles=['title','body'];text={r:f'{r} {rng.randrange(100000,999999)} '+('line\ncontinuation' if r=='body' and index in {0,2,6} else 'synthetic content') for r in roles}
        def make(path,order,styled=True):
            d=Document()
            if break_first:d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            for j,role in enumerate(order):
                p=d.add_paragraph(text[role]);p.paragraph_format.space_after=Pt(9 if role=='title' else 3)
                p.runs[0].font.size=Pt(16 if role=='title' else 11);p.runs[0].bold=role=='title'
                if blocked and j==0:d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            d.save(path)
        training=local/'training.docx';make(training,roles)
        text={r:text[r].replace('synthetic content', 'unseen wording '+str(rng.randrange(100000,999999))).replace('line\ncontinuation','new line\nnew continuation') for r in roles}
        golden=local/'golden.docx';make(golden,roles)
        c=extract_input(training,local/'extracted');sequence=(['divider'] if break_first else [])+(['title','divider','body'] if blocked else roles)
        labels={b['id']:r for b,r in zip(c['blocks'],sequence)}
        samples=[{'source':training,'content':c,'extracted':local/'extracted','semantics':semantic_spans(training)}]
        definitions={r:r for r in sequence};package,conflicts=build_package(samples,[labels],definitions,local/'template')
        if conflicts:raise RuntimeError('Synthetic reference should be unambiguous')
        incoming_roles=roles if index==4 else list(reversed(roles));source=local/'source.docx';make(source,incoming_roles)
        incoming_sequence=(['divider'] if break_first else [])+([incoming_roles[0],'divider',incoming_roles[1]] if blocked else incoming_roles)
        response={'groups':[{'start':i,'end':i,'role':r} for i,r in enumerate(incoming_sequence,1)]}
        restyle(source,local/'damaged.docx')
        cases.append({'id':'probe-'+str(index),'split':'development' if index<2 else 'sealed','source':local/'damaged.docx','template':local/'template','response':response,'golden':golden,'expect_blocked':blocked,'training':{'samples':samples,'labels':[labels],'roles':definitions}})
    return cases
