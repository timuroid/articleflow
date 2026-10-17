"""Verify actual renderer output, not just fc-match's independent catalogue."""
from functools import lru_cache
from pathlib import Path
import tempfile


@lru_cache(maxsize=16)
def audit(families=("Arial", "Times New Roman")):
    from docx import Document
    from .research import render_pdf
    import pymupdf
    result = {f: False for f in families}
    observed = {}
    try:
        with tempfile.TemporaryDirectory(prefix="articleflow-fonts-") as folder:
            root = Path(folder)
            for family in families:
                doc = Document()
                for bold, italic in ((False,False),(True,False),(False,True),(True,True)):
                    r = doc.add_paragraph().add_run("Font probe ABC кириллица 123")
                    r.font.name = family
                    r.bold, r.italic = bold, italic
                path = root / "probe.docx"
                doc.save(path)
                pdf = render_pdf(path, root / "rendered")
                with pymupdf.open(pdf) as rendered:
                    names = sorted({s["font"] for p in rendered for b in p.get_text("dict")["blocks"] if b["type"] == 0 for l in b["lines"] for s in l["spans"]})
                observed[family] = names
                normalized = family.lower().replace(" ", "")
                result[family] = bool(names) and all(n.lower().replace(" ", "").startswith(normalized) for n in names)
    except Exception:
        return {"families": result, "exact": False, "renderer_fonts": observed, "reason": "RENDERER_FONT_CHECK_UNAVAILABLE"}
    return {"families": result, "exact": all(result.values()), "renderer_fonts": observed}


def required_fonts(styles):
    from .docx_engine import parse_xml, NS, W
    names = set()
    profiles = list(styles.values())
    while profiles:
        profile = profiles.pop()
        profiles.extend(v for v in profile.values() if isinstance(v, dict))
        for key in ("rPr", "label_rPr"):
            if not profile.get(key):
                continue
            root = parse_xml(profile[key].encode())
            for node in root.findall("w:rFonts", NS):
                names.update(node.get(f"{{{W}}}{k}") for k in ("ascii", "hAnsi") if node.get(f"{{{W}}}{k}"))
    return tuple(sorted(names)) or ("Arial", "Times New Roman")
