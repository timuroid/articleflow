"""Build original research figures and an editable paper; no external calls."""
import argparse,json,re,statistics,textwrap
from pathlib import Path
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch,FancyArrowPatch
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]
NAVY='#17324D';TEAL='#087F8C';RED='#BB493B';GRAY='#65788A';GOLD='#C68A27'
plt.rcParams.update({'font.family':'DejaVu Sans','font.size':11,'axes.spines.top':False,'axes.spines.right':False,'axes.labelcolor':NAVY,'text.color':NAVY,'axes.titleweight':'bold','svg.fonttype':'none'})

def figures(data,out):
 out.mkdir(parents=True,exist_ok=True)
 def save(fig,name):
  fig.savefig(out/(name+'.svg'),bbox_inches='tight',facecolor='white');fig.savefig(out/(name+'.png'),dpi=200,bbox_inches='tight',facecolor='white');plt.close(fig)
 def canvas(size=(11,4)):
  f,a=plt.subplots(figsize=size);a.set_xlim(0,11);a.set_ylim(0,4);a.axis('off');return f,a
 def box(a,x,y,w,h,text,color=TEAL):
  a.add_patch(FancyBboxPatch((x,y),w,h,boxstyle='round,pad=.09',facecolor=color+'12',edgecolor=color,linewidth=1.4));a.text(x+w/2,y+h/2,'\n'.join(textwrap.fill(line,width=max(15,int(w*7))) for line in text.split('\n')),ha='center',va='center',fontsize=13)
 def arrow(a,start,end,color=GRAY):a.add_patch(FancyArrowPatch(start,end,arrowstyle='-|>',mutation_scale=14,color=color,linewidth=1.3))
 f,a=canvas()
 for x,t in [(0,'Эталоны DOCX\n+ контекст'),(2.7,'Большая модель\nроли и правила'),(5.4,'Малая модель\nновая разметка'),(8.1,'Сборка и проверка\nвсех страниц')]:box(a,x+.1,2,2.45,1,t)
 for x in (2.55,5.25,7.95):arrow(a,(x,2.5),(x+.25,2.5))
 box(a,3.5,.25,4,1,'Ошибки → уточнение → новый эксперимент',RED);arrow(a,(9.3,2),(9.3,.75));arrow(a,(9.3,.75),(7.6,.75));arrow(a,(3.4,.75),(3.4,1.9))
 a.text(.15,3.55,'Обучение заканчивается проверкой, а не решением модели',fontsize=14,weight='bold');save(f,'01-loop')
 f,a=canvas()
 box(a,.15,1.5,3.1,1.5,'Обучение\nOOXML + эталонные PDF\nроли, свойства, ошибки')
 box(a,3.95,1.5,3.1,1.5,'Замороженный пакет\nкаталоги + условия\nбез назначений эталона')
 box(a,7.75,1.5,3.1,1.5,'Применение\nновое содержание + Qwen\nдетерминированная сборка')
 arrow(a,(3.35,2.25),(3.85,2.25));arrow(a,(7.15,2.25),(7.65,2.25))
 a.text(5.5,.55,'Контрольный PDF доступен оценщику; он не является входом малой модели.',ha='center',fontsize=11,color=RED);save(f,'02-isolation')
 f,a=plt.subplots(figsize=(10,4.3))
 names=['Статья','Вложенные таблицы','Две колонки','Колонтитулы','Публичные DOCX','Эталон владельца'];counts=[3,3,3,3,3,1]
 colors=[TEAL,TEAL,GOLD,RED,RED,GOLD]
 a.barh(names[::-1],counts[::-1],color=colors[::-1]);a.set_xlim(0,4.6);a.set_xticks([0,1,2,3]);a.set_xlabel('Число различных документов')
 notes=['LLM: обучение + перенос','LLM: обучение + перенос','Контроль сборщика','Импорт не поддержан','3 отказа импорта','Неоднозначные правила']
 for y,(n,note) in enumerate(zip(counts[::-1],notes[::-1])):a.text(n+.07,y,note,va='center',fontsize=11)
 a.set_title('Корпус 023: 16 документов, разные уровни проверки',loc='left',pad=16);save(f,'03-corpus')
 f,ax=plt.subplots(1,2,figsize=(10,3.9))
 for a,series,title in zip(ax,data['fresh_heldout_trials'],['Новая статья · 10 страниц','Вложенные таблицы · 16 страниц']):
  rows=series['rows'];a.bar(range(1,11),[r['exact_pages'] for r in rows],color=TEAL,width=.7);a.set_xticks(range(1,11));a.set_ylim(0,18);a.set_title(title,fontsize=12);a.set_xlabel('Свежий вызов Qwen');a.set_ylabel('Точно совпавшие страницы');a.grid(axis='y',alpha=.15)
 f.suptitle('Повторяемость на двух отложенных документах: 10/10 + 10/10',fontsize=14,y=1.04);f.tight_layout();save(f,'04-trials')
 f,ax=plt.subplots(1,2,figsize=(10,3.8))
 for a,values,title in [(ax[0],[2,0],'Ошибки смысловых ролей'),(ax[1],[8,0],'Неразрешённые назначения')]:
  a.bar(['Первый цикл','После обратной связи'],values,color=[RED,TEAL]);a.set_ylim(0,max(values)*1.35);a.set_title(title,fontsize=12)
  for x,v in enumerate(values):a.text(x,v+.1,str(v),ha='center',weight='bold',fontsize=14)
 f.tight_layout();save(f,'05-feedback')
 f,ax=plt.subplots(1,2,figsize=(10,4))
 for i,s in enumerate(data['fresh_heldout_trials']):
  vals=[r['seconds'] for r in s['rows']];ax[0].scatter([i]*len(vals),vals,color=TEAL if i==0 else NAVY,alpha=.7);ax[0].plot([i-.15,i+.15],[statistics.median(vals)]*2,color=RED,lw=3)
 ax[0].set_xticks([0,1],['Статья','Вложенные таблицы']);ax[0].set_ylabel('Время вызова модели, с');ax[0].set_title('Все 20 задержек; красная линия — медиана',fontsize=11)
 costs=[s['total_cost_usd']/10 for s in data['fresh_heldout_trials']];ax[1].bar(['Статья','Вложенные таблицы'],costs,color=[TEAL,NAVY]);ax[1].set_ylabel('Средняя стоимость вызова, USD');ax[1].set_ylim(0,.005)
 for i,v in enumerate(costs):ax[1].text(i,v+.0001,f'${v:.5f}',ha='center')
 ax[1].set_title('Стоимость только малой модели',fontsize=11);f.tight_layout();save(f,'06-cost')
 f,a=canvas()
 box(a,.2,1.35,2.6,1.4,'Один и тот же текст\nбез оформления',NAVY);box(a,4,2.4,3,1,'Правило A: интервал 0 pt');box(a,4,.55,3,1,'Правило B: интервал 12 pt',GOLD)
 arrow(a,(2.9,2.05),(3.9,2.9));arrow(a,(2.9,2.05),(3.9,1.05));box(a,8,1.3,2.65,1.5,'Разные правильные\nстраницы при одинаковом\nнаблюдаемом входе',RED)
 arrow(a,(7.1,2.9),(7.9,2.3));arrow(a,(7.1,1.05),(7.9,1.8));a.text(5.5,3.8,'Без дополнительного сигнала выбор нельзя гарантировать',ha='center',fontsize=14,weight='bold');save(f,'07-ambiguity')
 # A diagram of the measured mutation, not a fake screenshot.
 f,ax=plt.subplots(1,2,figsize=(10,3.3))
 for a,gap,title in [(ax[0],.55,'Правило: 12 pt после таблицы'),(ax[1],.18,'Мутация: интервал удалён')]:
  a.set_xlim(0,6);a.set_ylim(0,3);a.axis('off');a.set_title(title,fontsize=12)
  for y in (2.5,2.1,1.7):a.plot([.3,5.7],[y,y],color=NAVY,lw=1)
  for x in (.3,3,5.7):a.plot([x,x],[1.7,2.5],color=NAVY,lw=1)
  y=1.7-gap;a.text(.6,y,'Анализ после таблицы…',va='top',fontsize=13)
  for j in range(1,3):a.plot([.4,5.5],[y-.27*j,y-.27*j],color=GRAY,lw=2)
  a.annotate('',xy=(5.9,1.7),xytext=(5.9,y),arrowprops={'arrowstyle':'<->','color':RED})
 f.text(.5,.02,'Схематическое объяснение: изменение сдвигает весь следующий блок.',ha='center',fontsize=13);save(f,'08-spacing')

def paper(source,out,figdir):
 d=Document();s=d.sections[0];s.page_width=Cm(21);s.page_height=Cm(29.7);s.top_margin=Cm(1.8);s.bottom_margin=Cm(1.8);s.left_margin=Cm(2.2);s.right_margin=Cm(2.2)
 normal=d.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(11);normal.paragraph_format.line_spacing=1.13;normal.paragraph_format.space_after=Pt(5);normal.paragraph_format.first_line_indent=Cm(.7)
 for name,size in [('Title',21),('Heading 1',15),('Heading 2',12)]:
  st=d.styles[name];st.font.name='Arial';st.font.size=Pt(size);st.font.color.rgb=RGBColor.from_string('17324D');st.paragraph_format.first_line_indent=Cm(0);st.paragraph_format.space_before=Pt(14);st.paragraph_format.space_after=Pt(7);st.paragraph_format.keep_with_next=True
 header=s.header.paragraphs[0];header.text='ARTICLEFLOW  /  ВОСПРОИЗВОДИМОЕ ОФОРМЛЕНИЕ ДОКУМЕНТОВ';header.paragraph_format.first_line_indent=Cm(0)
 for r in header.runs:r.font.name='Arial';r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string('65788A')
 footer=s.footer.paragraphs[0];footer.paragraph_format.first_line_indent=Cm(0);footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
 footer.add_run('Исследовательская рукопись · 17.10.2026  |  ').font.size=Pt(8)
 fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
 lines=source.read_text().splitlines();i=0;references=False
 def text(p,t):
  # Preserve link labels in readable print, exact URLs stay in bibliography.
  t=re.sub(r'\[([^]]+)\]\(([^)]+)\)',r'\1',t)
  for n,part in enumerate(t.split('**')):r=p.add_run(part);r.bold=bool(n%2)
 while i<len(lines):
  line=lines[i].strip();i+=1
  if not line:continue
  if line.startswith('# '):d.add_paragraph(line[2:],'Title');continue
  if line.startswith('## '):
   references=line[3:]=='Список литературы'
   heading=d.add_paragraph(line[3:],'Heading 1')
   if references:heading.paragraph_format.page_break_before=True
   continue
  if line.startswith('### '):d.add_paragraph(line[4:],'Heading 2');continue
  m=re.match(r'!\[([^]]*)\]\([^)]*/([^/)]+)\.svg\)',line)
  if m:
   p=d.add_paragraph();p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.keep_with_next=True;p.add_run().add_picture(str(figdir/(m.group(2)+'.png')),width=Cm(16.5))
   p=d.add_paragraph(m.group(1));p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.space_after=Pt(10)
   for r in p.runs:r.italic=True;r.font.size=Pt(9)
   continue
  if line.startswith('|'):
   rows=[line]
   while i<len(lines) and lines[i].strip().startswith('|'):rows.append(lines[i].strip());i+=1
   rows=[r for r in rows if not re.fullmatch(r'[| :\-]+',r)]
   vals=[[x.strip() for x in r.strip('|').split('|')] for r in rows];table=d.add_table(rows=0,cols=len(vals[0]));table.style='Light Shading Accent 1'
   for ri,row in enumerate(vals):
    cells=table.add_row().cells
    for cell,value in zip(cells,row):
     p=cell.paragraphs[0];p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.space_after=Pt(3);text(p,value)
     p.paragraph_format.keep_with_next=ri<len(vals)-1
     for run in p.runs:run.font.size=Pt(9.5);run.bold=ri==0
    trpr=cells[0]._tc.getparent().get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
    if ri==0:trpr.append(OxmlElement('w:tblHeader'))
   d.add_paragraph().paragraph_format.space_after=Pt(2);continue
  if line.startswith('- '):p=d.add_paragraph(style='List Bullet');p.paragraph_format.first_line_indent=Cm(0);text(p,line[2:]);continue
  p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
  if references:
   p.alignment=WD_ALIGN_PARAGRAPH.LEFT;p.paragraph_format.left_indent=Cm(.7);p.paragraph_format.first_line_indent=Cm(-.7);p.paragraph_format.tab_stops.add_tab_stop(Cm(.7))
   line=re.sub(r'^(\d+)\. ',r'\1.\t',line)
  text(p,line)
 d.core_properties.title='ArticleFlow: агентное извлечение шаблонов DOCX с проверяемым переносом';d.core_properties.author='';d.core_properties.subject='Исследовательская рукопись серии 023; авторство подлежит заполнению владельцем'
 out.parent.mkdir(parents=True,exist_ok=True);d.save(out)

if __name__=='__main__':
 p=argparse.ArgumentParser();p.add_argument('--output',type=Path,required=True);a=p.parse_args()
 fig=ROOT/'docs/research/figures023';data=json.loads((ROOT/'docs/research/evaluation-023.json').read_text());figures(data,fig)
 paper(ROOT/'docs/research/manuscript-023.md',a.output,fig)
