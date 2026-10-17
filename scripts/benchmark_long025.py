"""Long authored stress references; known roles isolate extraction from LLM accuracy."""
import argparse,json,time
from pathlib import Path
from collections import Counter
from docx import Document
from docx.shared import Pt,RGBColor,Twips
from docx.oxml.ns import qn
from build_complex023 import build
from articleflow.agent023 import extract_input,build_package,apply
from articleflow.rules023 import semantic_spans
from articleflow.catalogue import render_catalogue
from articleflow.docx_engine import dump,DocumentError,write_package
from articleflow.integrity023 import check
from articleflow.roundtrip022 import neutralize
from articleflow.research import render_pdf
from probe_style_equivalence024 import compare_bounded


def generate(path,family,seed):
    meta=build(path,'columns' if family=='columns' else 'stories' if family=='stories' else 'nested',seed)
    d=Document(path)
    if family in ('typography','typography_marked','ambiguous'):
        body=0
        for i,p in enumerate(d._element.body.iter(qn('w:p')),1):
            if meta['labels'][f'b-{i:06d}']!='body':continue
            from docx.text.paragraph import Paragraph
            para=Paragraph(p,d);body+=1
            run=para.add_run(' Термин');run.bold=True;run.font.color.rgb=RGBColor.from_string('204080')
            run=para.add_run(' — определение');run.italic=True
            run=para.add_run(' x')
            if family!='typography_marked':run.font.name='Arial'
            run=para.add_run('2');run.font.superscript=True
            run=para.add_run(' H')
            if family!='typography_marked':run.font.name='Arial'
            run=para.add_run('2');run.font.subscript=True
            if family=='ambiguous':para.paragraph_format.first_line_indent=Twips(709 if body%2 else 720)
        d.save(path)
        import zipfile
        with zipfile.ZipFile(path) as z:parts={n:z.read(n) for n in z.namelist()}
        write_package(path,parts)
    # SHA from the final fixture, never the unmodified precursor.
    import hashlib
    meta['sha256']=hashlib.sha256(path.read_bytes()).hexdigest();meta['family']=family
    return meta


def run(output):
    output.mkdir(parents=True,exist_ok=False)
    report={'schema':'long-stress-025','mode':'known_roles_not_model_trial','api_attempts':0,
            'fresh_model_trials':0,'split':'seed 1 reference; seed 17 held out','cases':[]}
    manifest=[]
    for family in ('nested','typography_marked','typography','ambiguous','columns','stories'):
        template=output/f'{family}-template'
        for seed in (1,17):
            start=time.monotonic();folder=output/f'{family}-{seed}';folder.mkdir()
            source=folder/'reference.docx';meta=generate(source,family,seed);manifest.append(meta)
            row={'family':family,'seed':seed,'split':'train' if seed==1 else 'held_out','paragraphs':meta['paragraphs']}
            try:
                extracted=folder/'extracted';content=extract_input(source,extracted)
                plan=json.loads((extracted/'oracle-plan.json').read_text())
                candidate=folder/'oracle.docx'
                render_catalogue(extracted/'template',extracted/'content',plan,candidate)
                golden=render_pdf(source,folder/'golden')
                pixels=compare_bounded(golden,render_pdf(candidate,folder/'oracle-render'))
                row['oracle']={'integrity':check(source,candidate),'pages':pixels['golden_pages'],
                               'pixel_exact':pixels['pixel_exact'],'differing_pixel_ratio':pixels['differing_pixel_ratio']}
                if seed==1:
                    altered=Document(source)
                    altered.paragraphs[0].runs[0].font.size=Pt(24)
                    bad=folder/'negative-font.docx';altered.save(bad)
                    negative=compare_bounded(golden,render_pdf(bad,folder/'negative-render'))
                    row['negative_control']={'mutation':'title size 16 to 24pt',
                                             'detected':not negative['pixel_exact']}
                    if negative['pixel_exact']:raise RuntimeError('Comparator missed negative control')
                labels=meta['labels'];definitions={r:r for r in set(labels.values())}
                if seed==1:
                    package,conflicts=build_package([{'content':content,'extracted':extracted,'semantics':semantic_spans(source)}],
                                                   [labels],definitions,template)
                    row['training_conflicts']=len(conflicts);row['conflict_kinds']=dict(Counter(c['kind'] for c in conflicts))
                    dump(folder/'conflicts.json',conflicts)
                neutral=folder/'neutral.docx';neutralize(source,neutral)
                incoming=extract_input(neutral,folder/'alignment');old=iter(content['blocks']);roles=[]
                for block in incoming['blocks']:
                    original=next(old,None)
                    while original and original['text']!=block['text'] and not original['text'] and not original['atoms']:original=next(old,None)
                    if original is None or original['text']!=block['text']:raise RuntimeError('Alignment mismatch')
                    roles.append(labels[original['id']])
                def caller(*args,**kwargs):
                    response={'groups':[{'start':i,'end':i,'role':r} for i,r in enumerate(roles,1)]}
                    kwargs['validate'](response);return response,{'mode':'known_roles','api_attempts':0}
                result=apply(neutral,template,folder/'application',{},caller=caller)
                row['transfer']={'status':result['status'],'unresolved':len(result['unresolved'])}
                if result.get('pdf'):
                    comparison=compare_bounded(golden,Path(result['pdf']))
                    row['transfer'].update(pixel_exact=comparison['pixel_exact'],differing_pixel_ratio=comparison['differing_pixel_ratio'])
                row['status']='evaluated'
            except DocumentError as e:row.update(status='rejected',error=str(e))
            row['seconds']=round(time.monotonic()-start,3);report['cases'].append(row)
            dump(output/'results.json',report);dump(output/'manifest.json',{'license':'CC0-1.0','entries':manifest})
            print(json.dumps(row),flush=True)
    return report

if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__);parser.add_argument('--output',type=Path,required=True)
    run(parser.parse_args().output)
