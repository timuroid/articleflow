"""Exhaustive synthetic order/format corruption test; no model claims or API calls."""
import argparse
import itertools
import json
from pathlib import Path
from docx import Document
from docx.shared import Pt
from articleflow.agent023 import extract_input,build_package,apply
from articleflow.rules023 import semantic_spans
from articleflow.roundtrip022 import neutralize
from articleflow.research import render_pdf
from articleflow.templateagent022 import compare_all

ROLES=['title','abstract','keywords','body']

def make(path,seed):
    texts=[f'Исследование метода {seed}',f'Аннотация. Оценены {seed*7} вариантов модели.',
           f'Ключевые слова: шаблон; опыт; серия {seed}.',
           f'Основной текст. Получены новые результаты в серии {seed}. '+'Проверяется перенос оформления. '*seed]
    d=Document()
    for i,text in enumerate(texts):
        p=d.add_paragraph(text);p.paragraph_format.space_after=Pt(4+i*2)
        p.paragraph_format.first_line_indent=Pt(12 if i==3 else 0)
        r=p.runs[0];r.font.name='Liberation Serif';r.font.size=Pt(16 if i==0 else 12);r.bold=i==0
    d.save(path)


def run(root,seeds):
    root.mkdir(parents=True,exist_ok=False);train=root/'train.docx';make(train,1)
    content=extract_input(train,root/'training');labels=dict(zip([b['id'] for b in content['blocks']],ROLES))
    _,conflicts=build_package([{'content':content,'extracted':root/'training','semantics':semantic_spans(train)}],
                             [labels],dict(zip(ROLES,['Название','Аннотация','Ключевые слова','Основной текст'])),root/'template')
    if conflicts:raise RuntimeError('Synthetic training conflict')
    summary={'schema':'order-024-benchmark','mode':'known_roles_compiler_test','api_attempts':0,
             'fresh_model_trials':0,'held_out_text_seeds':seeds,'cases':[]}
    for seed in seeds:
        golden=root/f'golden-{seed}.docx';make(golden,seed);golden_pdf=render_pdf(golden,root/f'golden-{seed}-pdf')
        neutral=root/f'neutral-{seed}.docx';neutralize(golden,neutral)
        for index,permutation in enumerate(itertools.permutations(range(4))):
            case=root/f'case-{seed}-{index}';case.mkdir();source=case/'input.docx'
            d=Document(neutral);body=d.element.body;nodes=list(body)[:4]
            for n in nodes:body.remove(n)
            for i,old in enumerate(permutation):body.insert(i,nodes[old])
            d.save(source)
            def caller(*args,**kwargs):
                response={'groups':[{'start':i+1,'end':i+1,'role':ROLES[old]} for i,old in enumerate(permutation)]}
                kwargs['validate'](response);return response,{'mode':'fixture_known_roles','api_attempts':0}
            result=apply(source,root/'template',case/'application',{},caller=caller)
            pixels=compare_all(golden_pdf,Path(result['pdf']),case/'diff') if result.get('pdf') else {}
            passed=bool(pixels.get('pixel_exact')) and all(result.get(k) for k in ('content_exact','placements_exact','nested_topology_exact','protected_vertical_exact'))
            summary['cases'].append({'seed':seed,'permutation':permutation,'passed':passed,
                                     'pixel_exact':pixels.get('pixel_exact',False),'moves':len(result['structure_moves'])})
            (root/'summary.json').write_text(json.dumps(summary,indent=2))
        print(json.dumps({'seed':seed,'cases':24,'passed':sum(c['passed'] for c in summary['cases'] if c['seed']==seed)}),flush=True)
    if not all(c['passed'] for c in summary['cases']):raise SystemExit(1)

if __name__=='__main__':
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('--output',required=True,type=Path)
    p.add_argument('--seeds',nargs='+',type=int,default=[3,7,11]);args=p.parse_args();run(args.output,args.seeds)
