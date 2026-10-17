"""Export the reconstructed Markdown manuscript as an editable reading draft.

This is not the lost manuscript and not a validated journal-template layout.
"""
import argparse
import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml.ns import qn


def export(source, output):
    doc=Document();section=doc.sections[0]
    section.page_width=Cm(21);section.page_height=Cm(29.7)
    section.top_margin=section.bottom_margin=Cm(2)
    section.left_margin=section.right_margin=Cm(2)
    for name in ['Normal','Title','Heading 1','Heading 2']:
        style=doc.styles[name];style.font.name='Times New Roman';style.font.size=Pt(12);style.font.color.rgb=RGBColor(0,0,0)
        fonts=style.element.find('.//'+qn('w:rFonts'))
        if fonts is not None:
            for key in list(fonts.attrib):
                if 'theme' in key.lower(): del fonts.attrib[key]
        style.paragraph_format.space_after=Pt(6)
    doc.styles['Normal'].paragraph_format.line_spacing=1.15
    doc.styles['Title'].font.size=Pt(16)
    doc.styles['Heading 1'].font.size=Pt(14)
    def plain(text):
        text=re.sub(r'\[([^\]]+)\]\(([^)]+)\)',r'\1',text)
        return text.replace('**','').replace('`','')
    table=None; fenced=False
    for line in source.read_text().splitlines():
        if line.startswith('```'):
            if not fenced:
                doc.add_paragraph('Эталонные DOCX → наблюдения → версионированный шаблон. Новый DOCX + шаблон → план ролей → проверка ссылок → сборка → контроль содержания, структуры и страниц → готовый DOCX или черновик.')
            fenced=not fenced;continue
        if fenced:continue
        if not line.strip(): table=None;continue
        if line.startswith('|'):
            cells=[plain(x.strip()) for x in line.strip('|').split('|')]
            if all(re.fullmatch(r'[:\- ]+',x) for x in cells):continue
            if table is None:table=doc.add_table(rows=0,cols=len(cells));table.style='Table Grid'
            row=table.add_row()
            for cell,text in zip(row.cells,cells):cell.text=text
        elif line.startswith('# '):doc.add_paragraph(plain(line[2:]),'Title')
        elif line.startswith('## '):doc.add_heading(plain(line[3:]),level=1)
        elif line.startswith('### '):doc.add_heading(plain(line[4:]),level=2)
        else:doc.add_paragraph(plain(line))
    doc.core_properties.title='ArticleFlow: реконструированный научный черновик 019'
    doc.core_properties.author=''
    output.parent.mkdir(parents=True,exist_ok=True);doc.save(output)

if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--source',type=Path,default=Path('docs/research/manuscript-reconstructed.md'));p.add_argument('--output',type=Path,required=True);a=p.parse_args();export(a.source,a.output)
