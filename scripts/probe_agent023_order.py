"""Known-role diagnostic for the missing ordering contract in agent023.

No provider calls. This records current behavior; it is not a model benchmark
and does not claim that preserving the incoming order satisfies the product.
"""
import argparse
import json
from pathlib import Path

from docx import Document
from articleflow.agent023 import apply, build_package, extract_input
from articleflow.rules023 import semantic_spans


def probe(folder: Path) -> dict:
    folder.mkdir(parents=True, exist_ok=False)
    rows = [
        ('title', 'Название исследования'),
        ('abstract', 'Аннотация. Описание метода.'),
        ('keywords', 'Ключевые слова: проверка; шаблон.'),
    ]
    incoming = [rows[1], rows[0], rows[2]]

    def make(path, values):
        document = Document()
        for role, text in values:
            paragraph = document.add_paragraph(text)
            if role == 'title':
                paragraph.runs[0].bold = True
        document.save(path)

    make(folder / 'reference.docx', rows)
    make(folder / 'input.docx', incoming)
    content = extract_input(folder / 'reference.docx', folder / 'train')
    labels = {b['id']: role for b, (role, _) in zip(content['blocks'], rows)}
    _, conflicts = build_package(
        [{'content': content, 'extracted': folder / 'train',
          'semantics': semantic_spans(folder / 'reference.docx')}],
        [labels], {'title': 'Название статьи', 'abstract': 'Аннотация',
                   'keywords': 'Ключевые слова'}, folder / 'template',
    )

    def caller(*args, **kwargs):
        response = {'groups': [
            {'start': i, 'end': i, 'role': role}
            for i, (role, _) in enumerate(incoming, 1)
        ]}
        kwargs['validate'](response)
        return response, {'api_attempts': 0, 'mode': 'diagnostic_known_roles'}

    result = apply(folder / 'input.docx', folder / 'template',
                   folder / 'application', {}, caller=caller)
    output = [p.text for p in Document(folder / 'application/result.docx').paragraphs]
    evidence = {
        'mode': 'diagnostic_known_roles_not_model_trial',
        'training_conflicts': len(conflicts), 'unresolved': len(result['unresolved']),
        'status': result['status'], 'expected_order': [text for _, text in rows],
        'actual_order': output,
        'input_order_preserved': output == [text for _, text in incoming],
        'target_order_repaired': output == [text for _, text in rows],
    }
    (folder / 'summary.json').write_text(
        json.dumps(evidence, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    return evidence


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    print(json.dumps(probe(args.output), ensure_ascii=False))
