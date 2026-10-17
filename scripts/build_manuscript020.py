"""Build our manuscript in arbitrary typography, without reading any target DOCX.

Public Markdown is the content source. Figures are original explanatory diagrams.
The gold role list is evaluator-only and must never be sent to the planner.
"""
from pathlib import Path
import argparse
import hashlib
import json
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
FONT = next((p for p in [Path('/System/Library/Fonts/Supplemental/Arial.ttf'), Path('/usr/share/fonts/truetype/msttcorefonts/Arial.ttf'), Path('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf')] if p.is_file()), None)

def figures():
    if FONT is None: raise SystemExit('A Cyrillic-capable Arial or DejaVu Sans font is required for figures')
    folder=ROOT/'docs/research/figures020';folder.mkdir(exist_ok=True)
    def draw(name, rows, caption):
        width=1600;height=160+len(rows)*170
        im=Image.new('RGB',(width,height),'white');g=ImageDraw.Draw(im)
        font=ImageFont.truetype(str(FONT),30)
        small=ImageFont.truetype(str(FONT),26)
        from xml.sax.saxutils import escape
        svg=[f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}"><rect width="100%" height="100%" fill="white"/>']
        for y,row in enumerate(rows):
            top=30+y*170;cell=width/len(row)
            for x,lines in enumerate(row):
                left=int(x*cell+22);right=int((x+1)*cell-22)
                g.rounded_rectangle((left,top,right,top+125),radius=12,fill='#f3f6fa',outline='#26384f',width=3)
                svg.append(f'<rect x="{left}" y="{top}" width="{right-left}" height="125" rx="12" fill="#f3f6fa" stroke="#26384f" stroke-width="3"/>')
                for j,line in enumerate(lines):
                    cy=top+23+j*35;g.text(((left+right)/2,cy),line,font=font,fill='#172a42',anchor='mt')
                    svg.append(f'<text x="{(left+right)/2}" y="{cy+26}" text-anchor="middle" font-family="Arial,sans-serif" font-size="30" fill="#172a42">{escape(line)}</text>')
                if x+1<len(row):
                    cx=right+22;cy=top+62
                    g.line((right+3,cy,cx+15,cy),fill='#26384f',width=3);g.polygon([(cx+15,cy),(cx+5,cy-6),(cx+5,cy+6)],fill='#26384f')
                    svg.append(f'<path d="M {right+3} {cy} H {cx+15} l -10 -6 m 10 6 l -10 6" fill="none" stroke="#26384f" stroke-width="3"/>')
        g.text((width/2,height-75),caption,font=small,fill='#26384f',anchor='mt')
        svg.append(f'<text x="800" y="{height-49}" text-anchor="middle" font-family="Arial,sans-serif" font-size="26" fill="#26384f">{escape(caption)}</text></svg>')
        im.save(folder/f'{name}.png');(folder/f'{name}.svg').write_text('\n'.join(svg))
    draw('architecture',[
        [['Эталоны DOCX','разбор и измерения'],['Роли и варианты','проверка конфликтов'],['Утверждённый шаблон','правила и происхождение']],
        [['Новая рукопись','содержание и объекты'],['LLM: роли и контексты','план без нового текста'],['Сборщик','применение свойств']],
        [['Проверка содержания','связей и структуры'],['Рендер и шрифты','визуальная проверка'],['Результат и отчёт','либо явный отказ']]],
        'Шаблон первой стадии используется планировщиком и сборщиком второй стадии.')
    draw('evidence',[
        [['Контроль сборщика','известный правильный план'],['Проверка исполнения','не измеряет модель']],
        [['Набор разработки','ошибки уже использованы'],['Проверка регрессии','не оценивает обобщение']],
        [['Независимый корпус','фиксируется заранее'],['Проверка переноса','следующий этап исследования']]],
        'Единица оценки — документ; отказы и ручные исправления учитываются отдельно.')

def blocks(text):
    rows=[]
    for part in re.split(r'\n\s*\n',text.strip()):
        if part.startswith('# '):rows.append(('title',part[2:]))
        elif part.startswith('## '):rows.append(('heading',part[3:]))
        elif part.startswith('!['):rows.append(('image',re.search(r'\]\(([^)]+)\)',part).group(1)))
        elif part.startswith('|'):
            table=[[c.strip() for c in line.strip().strip('|').split('|')] for line in part.splitlines() if not re.fullmatch(r'[| :\-]+',line)]
            rows.append(('table',table))
        else:rows.append(('paragraph',part))
    return rows

def build(output):
    figures();source=ROOT/'docs/research/manuscript-reconstructed.md'
    rows=blocks(source.read_text());d=Document();sec=d.sections[0]
    sec.page_width=Cm(21.59);sec.page_height=Cm(27.94);sec.left_margin=sec.right_margin=Cm(3)
    for name in ['Normal','Title','Heading 1']:
        s=d.styles[name];s.font.name='Courier New';s.font.size=Pt(13)
        s.font.color.rgb=RGBColor.from_string('274766');s.paragraph_format.line_spacing=1.4
    d.styles['Title'].font.size=Pt(19);d.styles['Heading 1'].font.size=Pt(16)
    # Deliberately permute abstract pieces; all scientific text remains unchanged.
    rows[2:6]=[rows[4],rows[5],rows[2],rows[3]]
    gold=[];references=False
    for kind,value in rows:
        role='body';context='body'
        if kind=='title':d.add_paragraph(value,'Title');role='title';context='title'
        elif kind=='heading':
            d.add_paragraph(value,'Heading 1');context='body.heading.1'
            references=value=='Список литературы'
        elif kind=='image':d.add_picture(str(source.parent/value),width=Cm(14.2));context='body.figure'
        elif kind=='table':
            table=d.add_table(rows=0,cols=len(value[0]));table.style='Table Grid'
            for row in value:
                for cell,txt in zip(table.add_row().cells,row):cell.text=txt
            context='body.table'
        else:
            d.add_paragraph(value)
            for prefix,r in [('Аннотация','abstract_heading'),('Постановка проблемы.','abstract_problem'),('Цель.','abstract_goal'),('Результаты.','abstract_results'),('Практическая значимость.','abstract_practical'),('Ключевые слова:','keywords')]:
                if value.startswith(prefix):role=r;context=r;break
            if value.startswith('Рис. '):context='body.figure_caption'
            elif value.startswith('Таблица '):context='body.table_caption'
            elif references:context='body.references'
        gold.append({'id':f'n-{len(gold)+1:06d}','role':role,'context':context})
    d.core_properties.author='';d.core_properties.title=rows[0][1]
    output.parent.mkdir(parents=True,exist_ok=True)
    if output.exists():raise SystemExit('Source exists: choose a new path instead of overwriting a frozen trial')
    d.save(output)
    output.with_suffix('.gold.json').write_text(json.dumps(gold,ensure_ascii=False,indent=2))
    output.with_suffix('.manifest.json').write_text(json.dumps({'source_sha256':hashlib.sha256(output.read_bytes()).hexdigest(),'markdown_sha256':hashlib.sha256(source.read_bytes()).hexdigest(),'blocks':len(rows),'words':len(source.read_text().split()),'input_mode':'arbitrary DOCX with original structure and styles visible; not neutral I1','gold':'evaluator only; excluded from model input'},ensure_ascii=False,indent=2))
    print(json.dumps({'blocks':len(rows),'output':str(output)}))

if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--output',type=Path,required=True);build(p.parse_args().output)
