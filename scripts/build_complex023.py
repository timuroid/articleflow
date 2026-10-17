"""Seeded, entirely authored long DOCX cases; ZIP timestamps are not frozen."""
import argparse
import hashlib
import json
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def build(path,family,seed):
    d=Document();s=d.sections[0]
    s.page_width=Cm(21);s.page_height=Cm(29.7)
    s.left_margin=Cm(3);s.right_margin=Cm(1.5);s.top_margin=s.bottom_margin=Cm(2)
    normal=d.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(12)
    normal.paragraph_format.space_after=Pt(0);normal.paragraph_format.line_spacing=1.15
    normal.paragraph_format.first_line_indent=Cm(1.25)
    labels={}
    def paragraph(text,role,parent=d):
        p=parent.add_paragraph(text);labels[p._p]=role
        f=p.paragraph_format
        if role=='title':
            f.alignment=WD_ALIGN_PARAGRAPH.CENTER;f.first_line_indent=Cm(0);f.space_after=Pt(18)
            for r in p.runs:r.bold=True;r.font.size=Pt(16)
        elif role=='heading':
            f.first_line_indent=Cm(0);f.space_before=Pt(14);f.space_after=Pt(7);f.keep_with_next=True
            for r in p.runs:r.bold=True
        elif role in {'caption','keywords'}:
            f.first_line_indent=Cm(0);f.space_before=Pt(6);f.space_after=Pt(6)
            for r in p.runs:r.bold=True
        elif role=='reference':
            f.left_indent=Cm(.9);f.first_line_indent=Cm(-.9);f.space_after=Pt(4)
        elif role=='after_table':
            f.space_before=Pt(12)
        if role not in {'title','heading'}:f.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
        return p
    paragraph(f'Структурная проверка документов: эксперимент {seed}', 'title')
    paragraph('Аннотация','heading')
    paragraph('Исследуется сохранность структуры, числовых данных и оформления при преобразовании научных документов. Сопоставляются исходный документ, нейтральное представление и восстановленная версия. Условия эксперимента задаются явно, результаты не заменяются предположениями.','body')
    paragraph('Ключевые слова: документ; структура; таблица; воспроизводимость; контроль.','keywords')
    if family=='columns':
        cols=s._sectPr.find(qn('w:cols'));cols.set(qn('w:num'),'2');cols.set(qn('w:space'),'420')
        normal.font.size=Pt(10);normal.paragraph_format.first_line_indent=Cm(.5)
    topics=['Постановка задачи','Материалы и методы','Экспериментальная установка','Проверка структуры','Анализ результатов','Ограничения метода','Обсуждение','Заключение']
    for chapter,title in enumerate(topics):
        paragraph(f'{chapter+1}. {title}','heading')
        for j in range(5+(seed%3)):
            value=(seed*17+chapter*11+j*7)%103
            paragraph((f'В серии {seed} на этапе {chapter+1} рассматривается наблюдение {j+1}. Значение контрольного параметра составляет {value}, а размер выборки равен {120+seed*3+j}. '
                'Структура входного документа определяет порядок измерений. Проверка сопоставляет отдельные абзацы и объекты, сохраняя текст без научного редактирования. '
                'Расхождение интервалов влияет на положение последующих строк и границы страниц. Поэтому средняя оценка близости не заменяет проверку каждой страницы. ')* (2 if family=='nested' else 1),'body')
        if chapter in (1,4,6):
            paragraph(f'Таблица {chapter+1}. Контрольные измерения','caption')
            rows=34 if family=='nested' and chapter==4 else 8
            table=d.add_table(rows=rows,cols=3);table.style='Table Grid';table.autofit=False
            for ri,row in enumerate(table.rows):
                for ci,cell in enumerate(row.cells):
                    cell.width=Cm(5 if ci==0 else 4)
                    p=cell.paragraphs[0];p.text=['Показатель','Значение','Отклонение'][ci] if ri==0 else (f'Наблюдение {ri}' if ci==0 else str((ri*13+ci+seed)%211))
                    labels[p._p]='table_header' if ri==0 else 'table_cell'
                    p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.space_after=Pt(0)
                    p.paragraph_format.line_spacing=1
                    for r in p.runs:r.font.size=Pt(10);r.bold=(ri==0)
                if ri==0:row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
            if family=='nested' and chapter==1:
                host=table.cell(3,0);nested=host.add_table(rows=3,cols=2);nested.style='Table Grid'
                for ri,row in enumerate(nested.rows):
                    for ci,cell in enumerate(row.cells):
                        p=cell.paragraphs[0];p.text=f'Подуровень {ri+1}.{ci+1}'
                        labels[p._p]='nested_cell';p.paragraph_format.first_line_indent=Cm(0)
                        for r in p.runs:r.font.size=Pt(8)
                for p in host.paragraphs:
                    if p._p not in labels:labels[p._p]='empty_cell'
            paragraph('После таблицы продолжается анализ измерений. Интервал над этим абзацем задаётся отдельным правилом и должен сохраняться при переносе текста.','after_table')
    paragraph('Список литературы','heading')
    for i in range(1,13):paragraph(f'{i}. Учебный источник {i}. Синтетическая запись для проверки висячего отступа. Серия {seed}. 2026. С. {10*i}–{10*i+9}.','reference')
    if family=='stories':
        d.sections[0].header.paragraphs[0].text='Исследовательский отчёт'
        d.sections[0].footer.paragraphs[0].text='Контрольная версия'
    path.parent.mkdir(parents=True,exist_ok=True);d.save(path)
    from articleflow.docx_engine import load_package,write_package
    # Repack very repetitive legitimate synthetic XML under the loader ratio cap.
    import zipfile
    with zipfile.ZipFile(path) as z:parts={n:z.read(n) for n in z.namelist()}
    write_package(path,parts)
    ps=list(d._element.body.iter(qn('w:p')))
    return {'path':path.name,'family':family,'seed':seed,'paragraphs':len(ps),
            'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),
            'labels':{f'b-{i+1:06d}':labels.get(p,'body') for i,p in enumerate(ps)}}


if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--output',type=Path,required=True);a=p.parse_args()
    a.output.mkdir(parents=True,exist_ok=False);entries=[]
    for family in ('article','nested','columns','stories'):
        for seed in (1,2,17):
            row=build(a.output/f'{family}-{seed}.docx',family,seed)
            row['split']='held_out' if seed==17 else 'train';entries.append(row)
    (a.output/'manifest.json').write_text(json.dumps({'series':'complex-023','license':'CC0-1.0','entries':entries},ensure_ascii=False,indent=2))
    print(json.dumps({'documents':len(entries),'paragraphs':sum(e['paragraphs'] for e in entries)}))
