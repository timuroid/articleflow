"""Build and verify ArticleFlow's fully synthetic DOCX corpus.

This corpus is for structural regression only. It does not reproduce the lost
conference corpus and must not be reported as independent real-world transfer.
"""

from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import io
import json
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.opc.constants import RELATIONSHIP_TYPE
from PIL import Image, ImageDraw


ROOT = Path(__file__).resolve().parents[1] / "datasets" / "synthetic-v1"
FIXED_TIME = (2020, 1, 1, 0, 0, 0)
STYLES = {
    "north": {"font": "Times New Roman", "title": 14, "body": 12, "heading": 12, "align": "center", "label": "Аннотация"},
    "east": {"font": "Arial", "title": 15, "body": 11, "heading": 12, "align": "left", "label": "Реферат"},
    "south": {"font": "Calibri", "title": 16, "body": 11, "heading": 13, "align": "center", "label": "Краткое содержание"},
}
PARTS = ("Постановка проблемы", "Цель", "Результаты", "Практическая значимость")


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def chart(seed: int) -> io.BytesIO:
    image = Image.new("RGB", (640, 230), "white")
    draw = ImageDraw.Draw(image)
    draw.line((50, 185, 610, 185), fill="#394d68", width=3)
    for index, height in enumerate((70 + seed * 5, 100 + seed * 3, 55 + seed * 7)):
        x = 100 + index * 160
        draw.rectangle((x, 185 - height, x + 65, 185), fill=(35 + 30 * index, 90 + 25 * index, 150))
    result = io.BytesIO()
    image.save(result, format="PNG")
    result.seek(0)
    return result


def add_hyperlink(paragraph, text: str) -> None:
    rel = paragraph.part.relate_to("https://example.org/articleflow-synthetic", RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel)
    run = OxmlElement("w:r")
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    link.append(run)
    paragraph._p.append(link)


def add_math(paragraph) -> None:
    math = OxmlElement("m:oMath")
    run = OxmlElement("m:r")
    text = OxmlElement("m:t")
    text.text = "q = a + b"
    run.append(text)
    math.append(run)
    paragraph._p.append(math)


def add_text(paragraph, text: str, *, bold=False, italic=False, superscript=False):
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.superscript = superscript
    return run


def style_document(document: Document, style: str | None) -> None:
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.7 if style == "north" else 2.1)
    section.right_margin = Cm(1.9)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    normal = document.styles["Normal"]
    normal.font.name = STYLES[style]["font"] if style else "Calibri"
    normal.font.size = Pt(STYLES[style]["body"] if style else 11)
    normal.paragraph_format.line_spacing = 1.5 if style == "north" else (1.15 if style == "east" else 1.3)
    normal.paragraph_format.first_line_indent = Cm(0.7 if style == "north" else 0.4)


def paragraph(document, text: str = "", *, kind: str = "body", style: str | None = None):
    p = document.add_paragraph()
    if style:
        profile = STYLES[style]
        if kind == "title":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if profile["align"] == "center" else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.first_line_indent = Cm(0)
            p.add_run(text).bold = True
            p.runs[0].font.size = Pt(profile["title"])
            return p
        if kind in {"heading", "abstract_heading", "keywords_heading"}:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if style == "north" else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.first_line_indent = Cm(0)
            p.add_run(text).bold = True
            p.runs[0].font.size = Pt(profile["heading"])
            return p
        if kind == "caption":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if style == "east" else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.first_line_indent = Cm(0)
    p.add_run(text)
    return p


def abstract_part(document, label: str, text: str, style: str | None):
    p = document.add_paragraph()
    if style == "south":
        p.add_run(label + ". ").italic = True
    else:
        p.add_run(label + ". ").bold = bool(style)
    p.add_run(text)
    return p


def make_document(style: str, number: int, *, source: bool = False, missing: bool = False) -> Document:
    document = Document()
    document.core_properties.author = "ArticleFlow synthetic corpus"
    document.core_properties.title = f"СИНТЕТИЧЕСКИЙ ТЕСТ {style}-{number:02d}"
    document.core_properties.created = dt.datetime(2020, 1, 1, 0, 0, 0)
    document.core_properties.modified = dt.datetime(2020, 1, 1, 0, 0, 0)
    applied = None if source else style
    style_document(document, applied)
    data = {
        "Постановка проблемы": f"В тестовом наборе {number} требуется упорядочить блоки без изменения чисел и формул.",
        "Цель": f"Проверить перенос смысловых частей для профиля {style} и варианта {number}.",
        "Результаты": f"В синтетическом опыте получены значения {number + 3}, {number + 7} и {number + 11} условных единиц.",
        "Практическая значимость": "Проверка выявляет потери фрагментов и подмену типографики в тестовой сборке.",
    }
    if source:
        paragraph(document, "РАБОЧАЯ ЗАМЕТКА: синтетический материал, не научная публикация.")
    paragraph(document, f"СИНТЕТИЧЕСКИЙ ТЕСТ: исследование системы {style.upper()}-{number:02d}", kind="title", style=applied)
    if not missing:
        paragraph(document, f"Автор тестового примера {number}; Лаборатория синтетических документов", style=applied)
    paragraph(document, STYLES[style]["label"], kind="abstract_heading", style=applied)
    order = ("Результаты", "Постановка проблемы", "Практическая значимость", "Цель") if source else PARTS
    for label in order:
        if missing and label == "Практическая значимость":
            continue
        abstract_part(document, label, data[label], applied)
    paragraph(document, "Ключевые слова", kind="keywords_heading", style=applied)
    paragraph(document, f"синтетический корпус; DOCX; профиль {style}; контроль {number}", style=applied)
    paragraph(document, "1. Проблема и постановка", kind="heading", style=applied)
    body = document.add_paragraph()
    add_text(body, "В документе определён ")
    add_text(body, "контрольный термин", bold=True)
    add_text(body, " и выбран ")
    add_text(body, "проверочный метод", italic=True)
    add_text(body, f". Измерение x{number} сохраняет индекс ")
    add_text(body, "2", superscript=True)
    add_text(body, ".")
    paragraph(document, "2. Метод и результаты", kind="heading", style=applied)
    table = document.add_table(rows=3, cols=2)
    table.style = "Table Grid"
    for row, values in enumerate((("Показатель", "Значение"), ("Число блоков", str(5 + number)), ("Контроль", str(10 + number)))):
        for col, value in enumerate(values):
            table.cell(row, col).text = value
            if row == 0:
                for run in table.cell(row, col).paragraphs[0].runs:
                    run.bold = True
    paragraph(document, f"Таблица 1. Синтетические измерения {number}", kind="caption", style=applied)
    document.add_picture(chart(number), width=Cm(11))
    paragraph(document, f"Рис. 1. Синтетическая схема {style}-{number:02d}", kind="caption", style=applied)
    p = paragraph(document, "Формула контроля: ", style=applied)
    add_math(p)
    paragraph(document, "3. Выводы", kind="heading", style=applied)
    paragraph(document, f"Для варианта {number} ожидается сохранение всех исходных научных узлов.", style=applied)
    ref = paragraph(document, "Синтетический справочный адрес: ", style=applied)
    add_hyperlink(ref, "example.org")
    if source:
        paragraph(document, "Редакционная заметка: рабочая версия для теста.")
    return document


def save_normalized(document: Document, target: Path) -> None:
    target.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as temporary:
        raw = Path(temporary) / "raw.docx"
        document.save(raw)
        with zipfile.ZipFile(raw) as source, zipfile.ZipFile(target, "w") as output:
            for name in sorted(source.namelist()):
                info = zipfile.ZipInfo(name, FIXED_TIME)
                info.compress_type = zipfile.ZIP_DEFLATED
                info.external_attr = 0o600 << 16
                output.writestr(info, source.read(name), compresslevel=9)


def build() -> dict:
    entries = []
    for style in STYLES:
        for number in range(1, 5):
            split = "train" if number <= 2 else "evaluation"
            reference = ROOT / split / f"{style}-{number:02d}.docx"
            save_normalized(make_document(style, number), reference)
            entries.append({"id": f"{style}-{number:02d}-reference", "split": split, "kind": "reference", "style": style, "path": reference.relative_to(ROOT).as_posix(), "sha256": digest(reference), "expected_abstract_order": list(PARTS)})
            if split == "evaluation":
                source = ROOT / "inputs" / f"{style}-{number:02d}.docx"
                save_normalized(make_document(style, number, source=True), source)
                entries.append({"id": f"{style}-{number:02d}-source", "split": split, "kind": "source", "style": style, "path": source.relative_to(ROOT).as_posix(), "sha256": digest(source), "paired_reference": reference.relative_to(ROOT).as_posix(), "expected_abstract_order": list(PARTS), "source_abstract_order": [PARTS[i] for i in (2, 0, 3, 1)]})
    missing = ROOT / "inputs" / "north-missing.docx"
    save_normalized(make_document("north", 5, source=True, missing=True), missing)
    entries.append({"id": "north-missing", "split": "evaluation", "kind": "missing_required", "style": "north", "path": missing.relative_to(ROOT).as_posix(), "sha256": digest(missing), "missing_roles": ["authors", "abstract.significance"], "expected_status": "draft_needs_review"})
    manifest = {"schema_version": 1, "name": "articleflow-synthetic-v1", "origin": "generated entirely for ArticleFlow; no real articles", "purpose": "structural development/regression, not independent conference transfer", "styles": list(STYLES), "entries": entries}
    (ROOT / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return manifest


def verify() -> dict:
    manifest = json.loads((ROOT / "manifest.json").read_text(encoding="utf-8"))
    seen = set()
    for entry in manifest["entries"]:
        path = ROOT / entry["path"]
        if entry["id"] in seen or not path.is_file() or digest(path) != entry["sha256"]:
            raise ValueError(f"missing, duplicate, or changed fixture: {entry['id']}")
        seen.add(entry["id"])
        with zipfile.ZipFile(path) as package:
            if package.testzip() is not None:
                raise ValueError(f"invalid DOCX: {entry['id']}")
            document_xml = package.read("word/document.xml")
            if b"oMath" not in document_xml or b"hyperlink" not in document_xml:
                raise ValueError(f"fixture lost OOXML features: {entry['id']}")
    return {"entries": len(seen), "styles": len(manifest["styles"]), "verified": True}


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--verify", action="store_true")
    args = parser.parse_args()
    print(json.dumps(verify() if args.verify else {"entries": len(build()["entries"]), "generated": True}, ensure_ascii=False))
