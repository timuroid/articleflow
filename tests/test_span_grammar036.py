"""Content-derived span grammar: transfer and ambiguity without source style cues."""
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document

from articleflow.agent023 import build_package, extract_input
from articleflow.inline026 import overlay
from articleflow.rules023 import (compile_plan, fit_span_grammar, semantic_spans,
                                  span_grammar_examples, span_grammar_segments)
from articleflow.sandbox034 import use


def marked(semantics, block_id, start, end):
    return overlay(semantics, {block_id:[{'start':start,'end':end,'role':'term'}]})


class SpanGrammarTests(unittest.TestCase):
    def test_segmentation_ignores_source_run_boundaries_and_typography(self):
        content={'blocks':[{'id':'new-id','text':'Ab-12 Ц','atoms':[]}],
                 'structure':[{'kind':'paragraph','id':'new-id'}]}
        plain={'new-id':[{'start':0,'end':7,'features':{
            'bold':False,'italic':False,'language':'en-US','inline_role':'plain'}}]}
        corrupted={'new-id':[{'start':0,'end':2,'features':{
            'bold':True,'italic':True,'language':'ru-RU','inline_role':'plain'}},
            {'start':2,'end':4,'features':{
                'bold':False,'italic':True,'language':'zh-CN','inline_role':'plain'}},
            {'start':4,'end':7,'features':{
                'bold':True,'italic':False,'language':'en-US','inline_role':'plain'}}]}
        a=span_grammar_segments(content,{'new-id':'body'},plain)
        b=span_grammar_segments(content,{'new-id':'body'},corrupted)
        self.assertEqual(a,b)
        self.assertEqual([(s['start'],s['end'],s['when']['character_class']) for s in a['new-id']],
                         [(0,2,'latin'),(2,3,'dash'),(3,5,'number'),
                          (5,6,'space'),(6,7,'cyrillic')])
        self.assertTrue(all(set(s['when'])=={'role','inline_role','character_class'}
                            for s in a['new-id']))

    def test_same_observable_span_class_with_different_targets_abstains(self):
        with tempfile.TemporaryDirectory() as td:
            root=Path(td);source=root/'reference.docx';doc=Document()
            p=doc.add_paragraph();p.add_run('Alpha').bold=True
            p.add_run(' Beta').bold=False;doc.save(source)
            content=extract_input(source,root/'extracted')
            bid=content['blocks'][0]['id'];labels={bid:'body'}
            # The two Latin words have indistinguishable input selectors.
            package,conflicts=build_package(
                [{'extracted':root/'extracted','content':content,
                  'semantics':semantic_spans(source)}],[labels],
                {'body':'Main document body'},root/'template',span_grammar=True)
            self.assertEqual(package['span_grammar']['status'],'ambiguous')
            self.assertTrue(any(c['kind']=='character' and
                                c['when'].get('character_class')=='latin' for c in conflicts))
            template=json.loads((root/'template/rules.json').read_text())
            plan,issues=compile_plan(template,content,package,labels,semantic_spans(source))
            self.assertIsNone(plan)
            self.assertEqual(issues[0]['code'],'SPAN_GRAMMAR_AMBIGUOUS')

    def test_semantic_mark_recovers_target_style_after_source_corruption(self):
        with tempfile.TemporaryDirectory() as td:
            root=Path(td);source=root/'reference.docx';doc=Document()
            p=doc.add_paragraph();p.add_run('Alpha').bold=True
            p.add_run(' Beta').bold=False;doc.save(source)
            content=extract_input(source,root/'extracted')
            bid=content['blocks'][0]['id'];labels={bid:'body'}
            annotations=marked(semantic_spans(source),bid,0,5)
            package,conflicts=build_package(
                [{'extracted':root/'extracted','content':content,'semantics':annotations}],
                [labels],{'body':'Main document body'},root/'template',
                {'term':'Technical term'},span_grammar=True)
            self.assertEqual(conflicts,[])
            self.assertEqual(package['span_grammar']['status'],'complete')
            self.assertFalse(package['span_grammar']['source_typography_used'])
            self.assertFalse(any(rule['kind']=='character' for rule in package['rules']))
            rules_json=json.dumps(package['span_grammar'],ensure_ascii=False)
            self.assertNotIn('Alpha',rules_json)
            self.assertNotIn(bid,rules_json)
            self.assertNotIn('bold',rules_json)
            oracle=json.loads((root/'extracted/oracle-plan.json').read_text())
            target_term=next(s['style'] for s in oracle['blocks'][0]['spans'] if s['start']==0)
            target_body=next(s['style'] for s in oracle['blocks'][0]['spans'] if s['end']==len('Alpha Beta'))

            # One italic source run erases the original bold/run segmentation;
            # new words prove the grammar does not memorize training text.
            incoming=root/'damaged.docx';doc=Document();p=doc.add_paragraph()
            run=p.add_run('Omega delta');run.bold=False;run.italic=True
            run.font.name='Courier New';doc.save(incoming)
            extracted=extract_input(incoming,root/'incoming')
            incoming_id=extracted['blocks'][0]['id']
            incoming_marks=marked(semantic_spans(incoming),incoming_id,0,5)
            template=json.loads((root/'template/rules.json').read_text())
            plan,issues=compile_plan(template,extracted,package,{incoming_id:'body'},incoming_marks)
            self.assertEqual(issues,[])
            self.assertEqual([(s['start'],s['end'],s['style']) for s in plan['blocks'][0]['spans']],
                             [(0,5,target_term),(5,11,target_body)])

    def test_sandbox_build_receives_opt_in_without_changing_legacy_default(self):
        fake=object()
        with patch('articleflow.engine034.build',return_value=({'schema':'fake'},[])) as builder:
            with use(fake):
                package,_=build_package([],[],{},Path('/tmp/not-created'),span_grammar=True)
            self.assertEqual(package['schema'],'fake')
            self.assertIs(builder.call_args.args[0],fake)
            self.assertIs(builder.call_args.args[-1],True)


if __name__=='__main__':unittest.main()
