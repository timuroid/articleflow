import copy
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from docx.shared import Pt
from lxml import etree as ET

from articleflow import typography019 as t
from articleflow.docx_engine import DocumentError, load_package, parse_xml, NS, xml_bytes, write_package
from articleflow.structure017 import inspect_templates, inventory, suggest_roles, compile_document
from articleflow.model019 import validate_plan, plan_document, save_private


class TypographyTests(unittest.TestCase):
    def test_inherited_font_toggle_and_partial_attributes(self):
        with tempfile.TemporaryDirectory() as d:
            path=Path(d)/'a.docx';doc=Document()
            doc.styles['Normal'].font.name='Arial';doc.styles['Normal'].font.size=Pt(13)
            style=doc.styles.add_style('Derived',1);style.base_style=doc.styles['Normal'];style.font.bold=True
            p=doc.add_paragraph('content',style='Derived');p.runs[0].italic=True;doc.save(path)
            parts=load_package(path);r=t.Resolver(parts);p=parse_xml(parts['word/document.xml']).find('.//w:p',NS)
            _,props=r.resolve(p,p.find('w:r',NS))
            self.assertEqual(t.child(props,'rFonts').get(t.tag('ascii')),'Arial')
            self.assertEqual(t.value(props,'sz'),'26');self.assertEqual(t.value(props,'b','1'),'1')
            self.assertEqual(t.value(props,'i','1'),'1')

    def test_scientific_signature_detects_number_and_link_changes(self):
        a=ET.fromstring(f'<w:p xmlns:w="{t.W}"><w:hyperlink w:anchor="one"><w:r><w:t>42</w:t></w:r></w:hyperlink></w:p>')
        b=copy.deepcopy(a);b.find('.//w:t',NS).text='43'
        self.assertNotEqual(t.scientific_signature([a]),t.scientific_signature([b]))
        b=copy.deepcopy(a);b.find('w:hyperlink',NS).set(t.tag('anchor'),'two')
        self.assertNotEqual(t.scientific_signature([a]),t.scientific_signature([b]))

    def test_missing_data_is_user_supplied_and_original_preserved(self):
        with tempfile.TemporaryDirectory() as d:
            path=Path(d)/'a.docx';out=Path(d)/'b.docx';doc=Document()
            doc.add_paragraph('Title');doc.add_paragraph('Аннотация');doc.save(path)
            before=path.read_bytes();roles,_=suggest_roles(inventory(path)['blocks'])
            report=compile_document(path,out,assignments=roles,required=['title','authors'],supplied={'authors':'Иванов И. И.'},approved=True,fonts_ok=True)
            self.assertEqual(report['status'],'ready');self.assertEqual(path.read_bytes(),before)
            self.assertIn('Иванов И. И.', ''.join(parse_xml(load_package(out)['word/document.xml']).itertext()))
            with self.assertRaises(DocumentError):
                compile_document(path,out,required=['title'],supplied={'title':'replacement'})

    def test_template_profiles_exclude_text_and_keep_conflicts(self):
        with tempfile.TemporaryDirectory() as d:
            paths=[]
            for i,font in enumerate(['Arial','Times New Roman']):
                p=Path(d)/f'{i}.docx';doc=Document();doc.styles['Normal'].font.name=font
                doc.add_paragraph('PRIVATE SCIENCE TITLE');doc.save(p);paths.append(p)
            spec=inspect_templates(paths)
            self.assertEqual(spec['schema_version'],'web-019')
            self.assertIn('__base__',spec['conflicts'])
            self.assertNotIn('PRIVATE SCIENCE',json.dumps(spec))

    def test_label_inside_single_run_and_semantic_emphasis(self):
        with tempfile.TemporaryDirectory() as d:
            path=Path(d)/'a.docx';doc=Document();doc.add_paragraph('Цель. Проверить число 42.');doc.save(path)
            parts=load_package(path);p=parse_xml(parts['word/document.xml']).find('.//w:p',NS)
            before=copy.deepcopy(p);resolver=t.Resolver(parts)
            profile={'pPr':f'<w:pPr xmlns:w="{t.W}"/>','rPr':f'<w:rPr xmlns:w="{t.W}"/>','label_rPr':f'<w:rPr xmlns:w="{t.W}"><w:b/></w:rPr>'}
            t.apply(p,profile,resolver,'abstract_goal')
            self.assertEqual(t.scientific_signature([before]),t.scientific_signature([p]))
            runs=p.findall('w:r',NS);self.assertEqual(len(runs),2)
            self.assertEqual(t.value(t.child(runs[0],'rPr'),'b','1'),'1')
            self.assertEqual(t.value(t.child(runs[1],'rPr'),'b'),'0')
            body=Document();p0=body.add_paragraph();p0.add_run('important').italic=True
            p0.add_run('2').font.superscript=True;body.save(path)
            parts=load_package(path);p=parse_xml(parts['word/document.xml']).find('.//w:p',NS)
            t.apply(p,profile,t.Resolver(parts),'body')
            self.assertEqual(t.value(p.find('w:r/w:rPr',NS),'i','1'),'1')
            self.assertEqual(t.value(p.findall('w:r',NS)[1].find('w:rPr',NS),'vertAlign'),'superscript')

    def test_mixed_abstract_ranges_reorder_without_text_loss(self):
        with tempfile.TemporaryDirectory() as d:
            path=Path(d)/'a.docx';out=Path(d)/'out.docx';doc=Document()
            doc.add_paragraph('Title');doc.add_paragraph('Автор А. Б.');doc.add_paragraph('Аннотация')
            text='Результаты. 42. Цель. Проверить. Постановка проблемы. Разметка. Практическая значимость. Тест.'
            doc.add_paragraph(text);doc.add_paragraph('Ключевые слова: тест');doc.save(path)
            inv=inventory(path);frags=[b for b in inv['blocks'] if b.get('range')]
            self.assertEqual(len(frags),4);self.assertEqual(''.join(b['text'] for b in frags),text)
            roles,_=suggest_roles(inv['blocks'])
            report=compile_document(path,out,assignments=roles,approved=True,fonts_ok=True)
            self.assertEqual(report['status'],'ready',report)
            paragraphs=Document(out).paragraphs
            ordered=[p.text for p in paragraphs if any(p.text.startswith(label) for label in ('Постановка','Цель','Результаты','Практическая'))]
            self.assertEqual([x.split('.')[0] for x in ordered],['Постановка проблемы','Цель','Результаты','Практическая значимость'])

    def test_front_matter_and_blank_anchor_preserve_order(self):
        with tempfile.TemporaryDirectory() as d:
            p=Path(d)/'source.docx';doc=Document()
            values=['УДК 004','Название','','Автор А. Б.','Тестовый университет','(Москва, Россия)','author@example.org','Аннотация','Цель. Проверка','Ключевые слова: тест','Основной текст']
            for text in values:doc.add_paragraph(text)
            doc.save(p);blocks=inventory(p)['blocks'];roles=suggest_roles(blocks)[0]
            self.assertEqual([roles[b['id']] for b in blocks[:8]],['udc','title','title','authors','affiliation','city','email','abstract_heading'])
            out=Path(d)/'result.docx';compile_document(p,out,assignments=roles,required=[])
            self.assertEqual([x.text for x in Document(out).paragraphs],values)

    def test_legacy_profile_cannot_silently_pass_ready(self):
        with tempfile.TemporaryDirectory() as d:
            p=Path(d)/'a.docx';doc=Document();doc.add_paragraph('Title');doc.save(p)
            roles=suggest_roles(inventory(p)['blocks'])[0]
            r=compile_document(p,Path(d)/'out.docx',assignments=roles,required=['title'],approved=True,fonts_ok=True,styles={'title':{'pPr':None,'rPr':None}})
            self.assertEqual(r['status'],'draft')
            self.assertTrue(any('Старый профиль' in q for q in r['questions']))

    def test_numbering_binding_is_protected(self):
        a=ET.fromstring(f'<w:p xmlns:w="{t.W}"><w:pPr><w:numPr><w:numId w:val="1"/></w:numPr></w:pPr><w:r><w:t>42</w:t></w:r></w:p>')
        b=copy.deepcopy(a);b.find('.//w:numId',NS).set(t.tag('val'),'2')
        self.assertNotEqual(t.scientific_signature([a]),t.scientific_signature([b]))


class ModelTests(unittest.TestCase):
    def test_plan_coverage_rejects_duplicates_unknown_and_cross_role_context(self):
        blocks=[{'id':'a','kind':'paragraph'},{'id':'b','kind':'paragraph'}]
        valid={'blocks':[{'id':'a','role':'title','context':'auto'},{'id':'b','role':'body','context':'body.heading.1'}]}
        self.assertEqual(validate_plan(valid,blocks,['auto','body.heading.1'])[0]['a'],'title')
        for change in ['duplicate','unknown','context','missing']:
            p=copy.deepcopy(valid)
            if change=='duplicate':p['blocks'][1]['id']='a'
            elif change=='unknown':p['blocks'][1]['id']='z'
            elif change=='context':p['blocks'][1]['role']='authors'
            else:p['blocks'].pop()
            with self.assertRaises(DocumentError):validate_plan(p,blocks,['auto','body.heading.1'])

    def test_migrated_ledger_never_sends_request(self):
        with tempfile.TemporaryDirectory() as d:
            root=Path(d);ledger=root/'ledger.json';prompt=root/'prompt';prompt.write_text('private')
            save_private(ledger,{'read_only':True})
            config={'ledger':str(ledger),'prompt':str(prompt),'key_file':'unused','model':'test'}
            with patch('articleflow.model019.read_key',return_value='unused'),patch('articleflow.model019.api') as api:
                with self.assertRaisesRegex(DocumentError,'LEDGER_READ_ONLY'):plan_document([],[],config,root/'run')
                api.assert_not_called()

    def test_budget_reserves_failure_and_prevents_second_call(self):
        with tempfile.TemporaryDirectory() as d:
            root=Path(d);ledger=root/'ledger.json';prompt=root/'prompt.txt';prompt.write_text('test')
            save_private(ledger,{'baseline':{'usage_usd':'12'},'additional_cap_usd':'0.02','attempts':[]})
            cfg={'ledger':str(ledger),'prompt':str(prompt),'key_file':'unused','model':'test'}
            def api(key, endpoint, body=None):
                if endpoint=='models':return 200,{'data':[{'id':'test','pricing':{'prompt':'0.000001','completion':'0.000002'}}]}
                return 429,{'error':{'code':429}}
            with patch('articleflow.model019.read_key',return_value='unused'),patch('articleflow.model019.account',return_value={'usage_usd':'12','available_usd':'3','at':'test'}),patch('articleflow.model019.api',side_effect=api) as network:
                with self.assertRaisesRegex(DocumentError,'PROVIDER_HTTP_429'):
                    plan_document([],[],cfg,root/'first')
                state=json.loads(ledger.read_text());self.assertEqual(len(state['attempts']),1)
                self.assertIsNone(state['attempts'][0]['cost_usd'])
                with self.assertRaisesRegex(DocumentError,'BUDGET_EXCEEDED'):
                    plan_document([],[],cfg,root/'second')
                self.assertEqual(sum(c.args[1]=='chat/completions' for c in network.call_args_list),1)


class WebModelTests(unittest.TestCase):
    def test_consent_model_revision_and_supplied_fields(self):
        from fastapi.testclient import TestClient
        from articleflow.web017 import create_app
        from test_web017 import article
        with tempfile.TemporaryDirectory() as d, patch('articleflow.web017.font_audit',return_value={'exact':True}), patch('articleflow.web017.model_config',return_value={'model':'small','extractor_model':'large'}) :
            with TestClient(create_app(Path(d))) as client:
                h={'X-CSRF-Token':client.get('/api/v017/config').json()['csrf_token']}
                r=client.post('/api/v017/templates',headers=h,data={'name':'Synthetic'},files={'files':('a.docx',article())});self.assertEqual(r.status_code,200)
                tid=r.json()['id'];url='/api/v017/templates/'+tid
                def planned(blocks,styles,config,directory):
                    return suggest_roles(blocks)[0],{}, {'status':'valid','api_attempts':1}
                with patch('articleflow.model019.plan_document',side_effect=planned) as model:
                    self.assertEqual(client.post(url+'/analyze',headers=h,json={'consent_external':False}).status_code,400)
                    model.assert_not_called()
                    r=client.post(url+'/analyze',headers=h,json={'consent_external':True});self.assertEqual(r.status_code,200,r.text)
                    self.assertEqual(model.call_args.args[2]['model'],'large')
                    self.assertEqual(r.json()['version'],2)
                    r=client.post(url+'/approve',headers=h,json={'selected':{}});self.assertEqual(r.status_code,200,r.text)
                    input_bytes=article(authors=False)
                    files={'files':('missing.docx',input_bytes)}
                    r=client.post('/api/v017/batches',headers=h,data={'template_id':tid,'use_llm':'true'},files=files);self.assertEqual(r.status_code,400)
                    r=client.post('/api/v017/batches',headers=h,data={'template_id':tid,'use_llm':'true','consent_external':'true'},files=files);self.assertEqual(r.status_code,202,r.text)
                    jid=r.json()['jobs'][0];job=client.get('/api/v017/jobs/'+jid).json()
                    self.assertEqual(job['status'],'draft');self.assertIn('authors',job['report']['missing'])
                    calls=model.call_count
                    r=client.patch('/api/v017/jobs/'+jid+'/decisions',headers=h,json={'expected_revision':0,'assignments':job['report']['assignments'],'supplied':{'authors':'Петров П. П.'}})
                    self.assertEqual(r.status_code,200,r.text);self.assertEqual(model.call_count,calls)
                    self.assertEqual(client.get('/api/v017/jobs/'+jid).json()['status'],'ready')
                    self.assertEqual(client.get('/api/v017/jobs/'+jid+'/files/source').content,input_bytes)


if __name__=='__main__':unittest.main()
