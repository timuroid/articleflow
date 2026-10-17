import sys
import tempfile
import unittest
from pathlib import Path
from docx import Document
from docx.shared import Inches
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
from experiment_real027 import content_signature, restyle
from articleflow.roundtrip022 import neutralize


class CorruptionGuardTests(unittest.TestCase):
    def test_tab_stop_is_layout_but_tab_character_is_content(self):
        with tempfile.TemporaryDirectory() as tmp:
            base=Path(tmp);source=base/'source.docx';neutral=base/'neutral.docx';other=base/'other.docx'
            doc=Document();p=doc.add_paragraph('Alpha\tBeta')
            p.paragraph_format.tab_stops.add_tab_stop(Inches(1))
            p.add_run('Power').bold=True
            doc.save(source);neutralize(source,neutral);restyle(neutral,other)
            self.assertEqual(content_signature(source),content_signature(neutral))
            self.assertEqual(content_signature(source),content_signature(other))
            doc=Document(other);doc.paragraphs[0].text='Alpha BetaPower';doc.save(other)
            self.assertNotEqual(content_signature(source),content_signature(other))
