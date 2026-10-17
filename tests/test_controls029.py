import json,tempfile,unittest
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from articleflow.agent023 import extract_input
from articleflow.catalogue import render_catalogue,placement_fingerprint,control_text
from articleflow.docx_engine import DocumentError,load_package,parse_xml,xml_bytes,write_package,NS
from articleflow.integrity023 import check

class ControlsTests(unittest.TestCase):
    def test_page_column_breaks_and_symbol_reconstructed_and_tampering_detected(self):
        with tempfile.TemporaryDirectory() as td:
            root=Path(td);src=root/'src.docx';d=Document();p=d.add_paragraph('Before')
            for kind in ['page','column']:
                br=OxmlElement('w:br');br.set(qn('w:type'),kind);p.add_run()._r.append(br)
            sym=OxmlElement('w:sym');sym.set(qn('w:font'),'Symbol');sym.set(qn('w:char'),'F061');p.add_run()._r.append(sym)
            p.add_run('After');d.save(src);extract_input(src,root/'input')
            plan=json.loads((root/'input/oracle-plan.json').read_text());target=root/'rebuilt.docx'
            render_catalogue(root/'input/template',root/'input/content',plan,target)
            self.assertTrue(all(check(src,target).values()))
            parts=load_package(target);tree=parse_xml(parts['word/document.xml']);tree.find('.//w:sym',NS).set(qn('w:char'),'F062');parts['word/document.xml']=xml_bytes(tree);write_package(target,parts)
            self.assertNotEqual(placement_fingerprint(src),placement_fingerprint(target))
            self.assertFalse(check(src,target)['placements_exact'])

    def test_controls_cannot_embed_arbitrary_xml_or_invalid_glyphs(self):
        for raw in [b'<w:object xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">DDE</w:object>',b'<w:br xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:t>text</w:t></w:br>',b'<w:sym xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:font="Symbol" w:char="not-hex"/>']:
            with self.assertRaisesRegex(DocumentError,'CONTENT_CONTROL_INVALID'):control_text(parse_xml(raw))
