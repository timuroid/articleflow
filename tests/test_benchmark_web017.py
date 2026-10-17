"""Negative controls for evaluator 018: never accept lost scientific payload."""
import importlib.util
import tempfile
import unittest
from pathlib import Path

from docx import Document
from articleflow.docx_engine import NS, load_package, parse_xml, write_package, xml_bytes

spec = importlib.util.spec_from_file_location("benchmark018", Path(__file__).resolve().parents[1] / "scripts/benchmark_web017.py")
bench = importlib.util.module_from_spec(spec)
spec.loader.exec_module(bench)


class BenchmarkEvaluatorTests(unittest.TestCase):
    def test_text_loss_and_semantic_span_loss_are_detected(self):
        with tempfile.TemporaryDirectory() as d:
            source, target = Path(d) / "source.docx", Path(d) / "target.docx"
            doc = Document()
            doc.add_paragraph().add_run("Payload 2").font.superscript = True
            doc.save(source)
            parts = load_package(source)
            root = parse_xml(parts["word/document.xml"])
            node = root.find(".//w:vertAlign", NS)
            node.getparent().remove(node)
            parts["word/document.xml"] = xml_bytes(root)
            write_package(target, parts)
            self.assertFalse(bench.conservation(source, target)["payload_exact"])
            root.find(".//w:t", NS).text = "Payload 3"
            parts["word/document.xml"] = xml_bytes(root)
            write_package(target, parts)
            self.assertFalse(bench.conservation(source, target)["all_payload_checks"])

    def test_duplicate_block_is_not_hidden_by_same_text_set(self):
        with tempfile.TemporaryDirectory() as d:
            source, target = Path(d) / "source.docx", Path(d) / "target.docx"
            doc = Document()
            doc.add_paragraph("Payload")
            doc.save(source)
            doc.add_paragraph("Payload")
            doc.save(target)
            self.assertFalse(bench.conservation(source, target)["payload_exact"])

    def test_identity_and_allowed_formatting_do_not_fail_payload(self):
        with tempfile.TemporaryDirectory() as d:
            source, target = Path(d) / "source.docx", Path(d) / "target.docx"
            doc = Document()
            run = doc.add_paragraph().add_run("Payload")
            doc.save(source)
            run.font.name = "Arial"
            doc.save(target)
            self.assertTrue(bench.conservation(source, target)["payload_exact"])


if __name__ == "__main__":
    unittest.main()
