"""Independent synthetic regressions for template completeness and list geometry."""
import copy
import tempfile
import unittest
from pathlib import Path
from docx import Document
from articleflow.docx_engine import load_package, write_package, parse_xml, xml_bytes, NS, paragraph_text
from articleflow import typography019 as t, representation021 as rep
from articleflow.structure017 import compile_document, inventory, inspect_templates


def profile():
    return {'schema':'typography-019', 'pPr':t.stable_xml(t.prop('pPr')), 'rPr':t.stable_xml(t.prop('rPr'))}


def native_source(path, values=('First reference with a long wrapped line', 'Second reference'), *, inherited=False):
    doc=Document()
    for text in values: doc.add_paragraph(text)
    doc.save(path); parts=load_package(path); root=parse_xml(parts['word/document.xml']); ps=root.findall('w:body/w:p',NS)
    numbering=t.prop('numbering'); a=t.prop('abstractNum',abstractNumId=77); lvl=t.prop('lvl',ilvl=0)
    for name,val in [('start','9'),('numFmt','decimal'),('lvlText','%1.'),('suff','tab'),('lvlJc','left')]:lvl.append(t.prop(name,val=val))
    pp=t.prop('pPr');pp.append(t.prop('ind',left=900,hanging=450));tabs=t.prop('tabs');tabs.append(t.prop('tab',val='num',pos=900));pp.append(tabs);lvl.append(pp);a.append(lvl);numbering.append(a)
    num=t.prop('num',numId=42);num.append(t.prop('abstractNumId',val=77));numbering.append(num);parts['word/numbering.xml']=xml_bytes(numbering)
    for p in ps:
        pp=t.prop('pPr');np=t.prop('numPr');np.append(t.prop('ilvl',val=0));np.append(t.prop('numId',val=42));pp.append(np);p.insert(0,pp)
    if inherited:
        styles=parse_xml(parts['word/styles.xml']); normal=styles.find('w:style[@w:styleId="Normal"]',NS); normal.append(copy.deepcopy(ps[0][0]))
        for p in ps:p.remove(p[0])
        parts['word/styles.xml']=xml_bytes(styles)
    parts['word/document.xml']=xml_bytes(root);write_package(path,parts);return parts,ps[0]


class Transfer021Tests(unittest.TestCase):
    def test_heading_boundary_captures_paragraph_mark_and_avoids_double_gap(self):
        from articleflow import layout021
        with tempfile.TemporaryDirectory() as d:
            d=Path(d);parts,_=native_source(d/'parts.docx');resolver=t.Resolver(parts)
            body=t.prop('body');body.append(t.prop('p'));blank=t.prop('p');pp=t.prop('pPr');mark=t.prop('rPr');mark.append(t.prop('sz',val=24));pp.append(mark);blank.append(pp);body.append(blank)
            head=t.prop('p');r=t.prop('r');txt=t.prop('t');txt.text='Section';r.append(txt);head.append(r);body.append(head)
            measured=layout021.preceding_profiles(body,2,resolver)
            self.assertEqual(t.value(parse_xml(measured[-1]['rPr'].encode()),'sz'),'24')
            style={**profile(),'leading_spacing':[measured[-1]]};pending=[(head,{'id':'n-3'},'body','body.heading.1',style)]
            output=[copy.deepcopy(body[0]),head];changes,warnings=layout021.apply_boundaries(output,pending,resolver)
            # The existing empty source paragraph is reused, not doubled.
            self.assertEqual(len(output),2);self.assertEqual(changes,[]);self.assertEqual(warnings,[])
            output=[head];changes,warnings=layout021.apply_boundaries(output,pending,resolver)
            self.assertEqual(len(output),2);self.assertEqual(changes[0]['inserted_empty_paragraphs'],1)
            self.assertEqual(t.value(output[0].find('w:pPr/w:rPr',NS),'sz'),'24')

    def test_section_grid_and_layout_tabs_do_not_leak_into_content(self):
        section=t.prop('sectPr');section.append(t.prop('docGrid',linePitch=700))
        t.apply_section(section,{'section':t.stable_xml(t.prop('sectPr'))})
        self.assertIsNone(t.child(section,'docGrid'))
        p=t.prop('p');pp=t.prop('pPr');tabs=t.prop('tabs');tabs.append(t.prop('tab',val='left',pos=900));pp.append(tabs);p.append(pp)
        r=t.prop('r');txt=t.prop('t');txt.text='Text';r.append(txt);r.append(t.prop('tab'));p.append(r)
        self.assertEqual(paragraph_text(p),'Text\t')

    def test_numbering_geometry_inherited_and_portable_without_ids(self):
        with tempfile.TemporaryDirectory() as d:
            parts,p=native_source(Path(d)/'n.docx');r=t.Resolver(parts);pr=t.measure(p,r,'body.references')
            pp=parse_xml(pr['pPr'].encode());self.assertEqual(t.child(pp,'ind').get(t.tag('left')),'900')
            self.assertEqual(pp.find('w:tabs/w:tab',NS).get(t.tag('pos')),'900')
            self.assertNotIn('numId',str(pr));self.assertNotIn('abstractNumId',str(pr));self.assertEqual(pr['numbering']['schema'],'numbering-021')

    def test_number_font_comes_from_reference_paragraph_mark(self):
        with tempfile.TemporaryDirectory() as d:
            parts,p=native_source(Path(d)/'n.docx');mark=t.prop('rPr');mark.append(t.prop('sz',val=24));p.find('w:pPr',NS).append(mark)
            pr=t.measure(p,t.Resolver(parts),'body.references')
            self.assertEqual(t.value(parse_xml(pr['numbering']['rPr'].encode()),'sz'),'24')

    def test_manual_numbers_preserved_including_discontinuous_and_three_digits(self):
        with tempfile.TemporaryDirectory() as d:
            d=Path(d);parts,p=native_source(d/'ref.docx');pr=t.measure(p,t.Resolver(parts),'body.references')
            doc=Document()
            for n in (1,9,10,100):
                p=doc.add_paragraph();p.add_run(str(n));p.add_run('. ');p.add_run('Exact DOI 10.1000/xyz. Figure 42.').italic=True
            source=d/'in.docx';doc.save(source);roles={b['id']:'body' for b in inventory(source)['blocks']};contexts={k:'body.references' for k in roles}
            report=compile_document(source,d/'out.docx',assignments=roles,contexts=contexts,styles={'body.references':pr},required=[],approved=True,fonts_ok=True)
            out=load_package(d/'out.docx');body=parse_xml(out['word/document.xml']).find('w:body',NS)
            self.assertEqual(list(rep.source_ordinals(body,t.Resolver(out)).values()),[1,9,10,100])
            self.assertEqual([paragraph_text(p) for p in body if p.tag==t.tag('p')],['Exact DOI 10.1000/xyz. Figure 42.']*4)
            self.assertEqual(len(report['representation_changes']),4);self.assertEqual(report['status'],'ready')
            for p in body.findall('w:p',NS):self.assertIsNotNone(p.find('.//w:i',NS))

    def test_native_source_counter_and_style_binding_preserved(self):
        for inherited in (False,True):
            with self.subTest(inherited=inherited), tempfile.TemporaryDirectory() as d:
                d=Path(d);parts,p=native_source(d/'in.docx',inherited=inherited);pr=t.measure(p,t.Resolver(parts),'body.references')
                ids=[b['id'] for b in inventory(d/'in.docx')['blocks']]
                compile_document(d/'in.docx',d/'out.docx',assignments=dict.fromkeys(ids,'body'),contexts=dict.fromkeys(ids,'body.references'),styles={'body.references':pr},required=[])
                out=load_package(d/'out.docx');body=parse_xml(out['word/document.xml']).find('w:body',NS)
                self.assertEqual(list(rep.source_ordinals(body,t.Resolver(out)).values()),[9,10])

    def test_keywords_heading_only_or_empty_does_not_satisfy_presence(self):
        for words in ('Ключевые слова:', '   ', ''):
            with self.subTest(words=words), tempfile.TemporaryDirectory() as d:
                d=Path(d);doc=Document();doc.add_paragraph(words);doc.save(d/'in.docx')
                report=compile_document(d/'in.docx',d/'out.docx',assignments={'n-000001':'keywords'},styles={'keywords':profile(),'keywords.heading':profile()},required=['keywords'],approved=True,fonts_ok=True)
                self.assertEqual(report['missing'],['keywords']);self.assertEqual(report['status'],'draft')
                self.assertIn('[ЗАПОЛНИТЬ: keywords]',[p.text for p in Document(d/'out.docx').paragraphs])

    def test_inline_keywords_split_exactly_and_separate_not_duplicated(self):
        for texts in (['Ключевые слова: DOCX; ИИ; контроль.'], ['Ключевые слова:', 'DOCX; ИИ; контроль.']):
            with self.subTest(texts=texts), tempfile.TemporaryDirectory() as d:
                d=Path(d);doc=Document()
                for text in texts:doc.add_paragraph(text)
                doc.save(d/'in.docx');ids=[b['id'] for b in inventory(d/'in.docx')['blocks']]
                report=compile_document(d/'in.docx',d/'out.docx',assignments=dict.fromkeys(ids,'keywords'),styles={'keywords':profile(),'keywords.heading':profile()},required=['keywords'],approved=True,fonts_ok=True)
                out=Document(d/'out.docx');self.assertEqual(len(out.paragraphs),2);self.assertEqual(''.join(p.text for p in out.paragraphs),''.join(texts));self.assertEqual(report['missing'],[]);self.assertEqual(report['status'],'ready')

    def test_observed_but_not_required_is_not_silently_omitted(self):
        with tempfile.TemporaryDirectory() as d:
            d=Path(d);doc=Document();doc.add_paragraph('Article');doc.save(d/'in.docx')
            report=compile_document(d/'in.docx',d/'out.docx',assignments={'n-000001':'title'},styles={'title':profile(),'keywords':profile(),'email':profile()},required=['title'],approved=True,fonts_ok=True)
            self.assertFalse(report['coverage']['keywords']['present']);self.assertFalse(report['coverage']['email']['present']);self.assertEqual(report['status'],'draft');self.assertEqual(len(report['questions']),2)

    def test_supplied_keywords_get_only_approved_label_and_reject_heading_only_value(self):
        with tempfile.TemporaryDirectory() as d:
            d=Path(d);doc=Document();doc.add_paragraph('Article');doc.save(d/'in.docx')
            kwargs={'assignments':{'n-000001':'title'},'styles':{'title':profile(),'keywords':profile(),'keywords.heading':{**profile(),'literal':'Ключевые слова:'}},'required':['keywords'],'approved':True,'fonts_ok':True}
            r=compile_document(d/'in.docx',d/'out.docx',supplied={'keywords':'DOCX; анализ.'},**kwargs)
            self.assertEqual(r['status'],'ready');self.assertEqual([p.text for p in Document(d/'out.docx').paragraphs],['Article','Ключевые слова:','DOCX; анализ.'])
            from articleflow.docx_engine import DocumentError
            with self.assertRaisesRegex(DocumentError,'SUPPLIED_DATA_INVALID'):
                compile_document(d/'in.docx',d/'bad.docx',supplied={'keywords':'Ключевые слова:'},**kwargs)

    def test_extraction_proposes_keywords_as_required(self):
        with tempfile.TemporaryDirectory() as d:
            p=Path(d)/'ref.docx';doc=Document()
            for s in ['Article','Author','Аннотация','Цель. Проверка.','Ключевые слова:', 'DOCX; контроль.']:doc.add_paragraph(s)
            doc.save(p);spec=inspect_templates([p]);self.assertIn('keywords',spec['structure']['required']);self.assertIn('keywords',spec['structure']['observed'])

    def test_unsupported_or_ambiguous_list_keeps_source_and_reports_review(self):
        for case in ('roman','double','leading_zero','hyperlink'):
            with self.subTest(case=case), tempfile.TemporaryDirectory() as d:
                d=Path(d);ref,refp=native_source(d/'ref.docx');pr=t.measure(refp,t.Resolver(ref),'body.references')
                parts,p=native_source(d/'in.docx',values=('1. Citation.' if case=='double' else 'Citation.',))
                if case=='roman':
                    n=parse_xml(parts['word/numbering.xml']);n.find('.//w:numFmt',NS).set(t.tag('val'),'upperRoman');parts['word/numbering.xml']=xml_bytes(n)
                if case in ('leading_zero','hyperlink'):
                    root=parse_xml(parts['word/document.xml']);p=root.find('w:body/w:p',NS);p.remove(p.find('w:pPr',NS));p.find('w:r/w:t',NS).text='01. Citation.' if case=='leading_zero' else '1. Citation.'
                    if case=='hyperlink':
                        link=t.prop('hyperlink',anchor='local');link.append(p.find('w:r',NS));p.append(link)
                    parts['word/document.xml']=xml_bytes(root)
                write_package(d/'in.docx',parts)
                r=compile_document(d/'in.docx',d/'out.docx',assignments={'n-000001':'body'},contexts={'n-000001':'body.references'},styles={'body.references':pr},required=[],approved=True,fonts_ok=True)
                self.assertEqual(r['status'],'draft');self.assertEqual(r['representation_changes'],[])
                before=parse_xml(parts['word/document.xml']).find('w:body/w:p',NS);after=parse_xml(load_package(d/'out.docx')['word/document.xml']).find('w:body/w:p',NS)
                self.assertEqual(t.scientific_signature([before]),t.scientific_signature([after]))

    def test_legacy_numbering_profile_is_explicit_review(self):
        with tempfile.TemporaryDirectory() as d:
            d=Path(d);doc=Document();doc.add_paragraph('1. Exact citation.');doc.save(d/'in.docx')
            r=compile_document(d/'in.docx',d/'out.docx',assignments={'n-000001':'body'},contexts={'n-000001':'body.references'},styles={'body.references':profile()},required=[],approved=True,fonts_ok=True)
            self.assertEqual(r['status'],'draft');self.assertIn('нумерации',r['questions'][0]);self.assertEqual(Document(d/'out.docx').paragraphs[0].text,'1. Exact citation.')

if __name__=='__main__':unittest.main()

class TemplateApproval021Tests(unittest.TestCase):
    def test_visual_failure_or_unavailable_cannot_approve_version(self):
        import io
        from unittest.mock import patch
        from fastapi.testclient import TestClient
        from articleflow.web017 import create_app
        with tempfile.TemporaryDirectory() as d, TestClient(create_app(Path(d))) as client:
            headers={'X-CSRF-Token':client.get('/api/v017/config').json()['csrf_token']}
            doc=Document()
            for text in ['Title','Author','Аннотация','Цель. Проверить.','Ключевые слова: тест.','', 'Введение','Body']:doc.add_paragraph(text)
            buf=io.BytesIO();doc.save(buf)
            response=client.post('/api/v017/templates',headers=headers,data={'name':'Synthetic'},files=[('files',('sample.docx',buf.getvalue()))]);self.assertEqual(response.status_code,200)
            data=response.json();tid=data['id'];selected={k:next(iter(v)) for k,v in data['spec']['variants'].items()}
            for args,code in [({'return_value':{'status':'failed','rows':[]}},'TEMPLATE_LAYOUT_MISMATCH'),({'side_effect':RuntimeError('renderer absent')},'TEMPLATE_LAYOUT_CHECK_UNAVAILABLE')]:
                with patch('articleflow.templatecheck021.validate_samples',**args):
                    r=client.post(f'/api/v017/templates/{tid}/approve',headers=headers,json={'selected':selected})
                self.assertEqual(r.status_code,409);self.assertEqual(r.json()['detail']['code'],code)
                current=client.get(f'/api/v017/templates/{tid}').json();self.assertFalse(current['approved']);self.assertEqual(current['version'],1)
            with patch('articleflow.templatecheck021.validate_samples',return_value={'status':'passed','rows':[],'scope':'test stub'}):
                r=client.post(f'/api/v017/templates/{tid}/approve',headers=headers,json={'selected':selected})
            self.assertEqual(r.status_code,200);current=client.get(f'/api/v017/templates/{tid}').json();self.assertEqual(current['spec']['layout_validation']['status'],'passed')
