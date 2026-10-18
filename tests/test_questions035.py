"""Conservative clarification of persistent training contradictions."""
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from docx.shared import Pt

from articleflow.agent023 import run
from articleflow.questions035 import clarify, conflict_signature


class ConflictClarificationTests(unittest.TestCase):
    def test_persistent_contrast_names_exact_places_and_property_alternatives(self):
        conflict={'kind':'character','when':{'role':'body'},'variants':2,
                  'differing_properties':['0/rPr/0/sz/@val'],
                  'property_variants':{'small':{'0/rPr/0/sz/@val':'22'},
                                       'large':{'0/rPr/0/sz/@val':'34'}},
                  'examples':[{'set':{'style':'small'}},{'set':{'style':'large'}}]}
        locations=[{'document':1,'block_id':'b-000002','paragraph':2,'range':[0,5],
                    'excerpt':'alpha','target':{'style':'small'}},
                   {'document':1,'block_id':'b-000003','paragraph':3,'range':[0,4],
                    'excerpt':'beta','target':{'style':'large'}}]
        persistent={conflict_signature(conflict)}
        questions=clarify([conflict],{'groups':[{'locations':locations}]},
                          {'body':'Ordinary body text'},persistent)
        self.assertEqual(len(questions),1)
        self.assertIn('абз. 2, символы 0–5',questions[0]['question'])
        self.assertIn('абз. 3, символы 0–4',questions[0]['question'])
        self.assertIn('11 пт',questions[0]['question'])
        self.assertIn('17 пт',questions[0]['question'])
        self.assertEqual([e['variant'] for e in questions[0]['evidence']],['A','B'])
        self.assertEqual(questions[0]['source'],'persistent_training_conflict')
        self.assertEqual(clarify([conflict],{'groups':[{'locations':locations}]},
                                 {'body':'Ordinary body text'},set()),[])
        self.assertEqual(clarify([conflict],{'groups':[{'locations':locations[:1]}]},
                                 {'body':'Ordinary body text'},persistent),[])

    def test_agent_waits_once_after_third_persistent_conflict_and_uses_answer(self):
        with tempfile.TemporaryDirectory() as td:
            root=Path(td);source=root/'reference.docx';doc=Document()
            for text,size in [('Head',11),('alpha',11),('beta',17),('gamma',11),('tail',11)]:
                doc.add_paragraph(text).runs[0].font.size=Pt(size)
            doc.save(source)
            feedback={'revision':0,'text':''};teacher_prompts=[];wait_calls=[]
            def caller(config,messages,directory,**kwargs):
                self.assertEqual(config['model'],'teacher','Conflicting package must skip student')
                teacher_prompts.append(json.loads(messages[1]['content']))
                role=f'body_{len(teacher_prompts)}'
                return {'roles':{role:'Ordinary body paragraph'},
                        'groups':[{'start':1,'end':5,'role':role}],
                        'questions':[],'inline_roles':{},'inline':[]},{}
            def fake_wait(questions,initial,provider,emit,*,seconds):
                wait_calls.append((questions,seconds))
                emit({'status':'waiting','deadline_unix':0,'questions':questions})
                feedback.update(revision=1,text='Вариант Б — опечатка автора.')
                state={'status':'answered','received_revision':1,'questions':questions}
                emit(state)
                return feedback.copy(),state
            with patch('articleflow.agent023.render_pdf',return_value=root/'unused.pdf'),\
                 patch('articleflow.questions035.wait',side_effect=fake_wait):
                report=run([source],root/'run',{'model':'teacher','triage_conflicts':True,
                                                'question_wait_seconds':30},
                           {'model':'student'},max_iterations=4,caller=caller,
                           feedback_provider=lambda:feedback.copy(),adaptive=True)
            self.assertEqual(report['status'],'needs_review')
            self.assertEqual(len(teacher_prompts),4)
            self.assertEqual(len(wait_calls),1)
            self.assertEqual(wait_calls[0][1],30)
            self.assertEqual(report['iterations'][2]['derived_questions'][0]['source'],
                             'persistent_training_conflict')
            self.assertEqual(report['iterations'][2]['clarification_revision_received'],1)
            self.assertEqual(report['question_window']['status'],'answered')
            self.assertEqual(teacher_prompts[3]['observations']['user_feedback'],
                             'Вариант Б — опечатка автора.')
            self.assertEqual([row['code'] for row in report['iterations']],
                             ['TRAINING_CONFLICTS']*4)
            self.assertTrue(all(not row['applications'] for row in report['iterations']))

    def test_no_feedback_keeps_question_and_unresolved_result(self):
        with tempfile.TemporaryDirectory() as td:
            root=Path(td);source=root/'reference.docx';doc=Document()
            for text,size in [('Head',11),('alpha',11),('beta',17),('gamma',11),('tail',11)]:
                doc.add_paragraph(text).runs[0].font.size=Pt(size)
            doc.save(source);calls=[]
            def caller(config,messages,directory,**kwargs):
                calls.append(1);role=f'body_{len(calls)}'
                return {'roles':{role:'Ordinary body paragraph'},
                        'groups':[{'start':1,'end':5,'role':role}],
                        'questions':[],'inline_roles':{},'inline':[]},{}
            with patch('articleflow.agent023.render_pdf',return_value=root/'unused.pdf'):
                report=run([source],root/'run',{'model':'teacher','triage_conflicts':True},
                           {'model':'student'},max_iterations=3,caller=caller,adaptive=True)
            self.assertEqual(report['status'],'needs_review')
            self.assertEqual(len(report['questions']),1)
            self.assertEqual(report['questions'][0]['source'],'persistent_training_conflict')
            self.assertNotIn('question_window',report)
            self.assertFalse(list((root/'run').rglob('student/validated.json')))

    def test_timeout_does_not_resolve_or_repeat_the_question_window(self):
        with tempfile.TemporaryDirectory() as td:
            root=Path(td);source=root/'reference.docx';doc=Document()
            for text,size in [('Head',11),('alpha',11),('beta',17),('gamma',11),('tail',11)]:
                doc.add_paragraph(text).runs[0].font.size=Pt(size)
            doc.save(source);calls=[];windows=[]
            def caller(config,messages,directory,**kwargs):
                calls.append(json.loads(messages[1]['content']))
                role=f'body_{len(calls)}'
                return {'roles':{role:'Ordinary body paragraph'},
                        'groups':[{'start':1,'end':5,'role':role}],
                        'questions':[],'inline_roles':{},'inline':[]},{}
            def fake_timeout(questions,initial,provider,emit,*,seconds):
                windows.append(questions)
                state={'status':'continued_without_answer','questions':questions}
                emit(state)
                return initial,state
            with patch('articleflow.agent023.render_pdf',return_value=root/'unused.pdf'),\
                 patch('articleflow.questions035.wait',side_effect=fake_timeout):
                report=run([source],root/'run',{'model':'teacher','triage_conflicts':True,
                                                'question_wait_seconds':30},
                           {'model':'student'},max_iterations=4,caller=caller,
                           feedback_provider=lambda:{'revision':0,'text':''},adaptive=True)
            self.assertEqual(len(windows),1)
            self.assertEqual(report['question_window']['status'],'continued_without_answer')
            self.assertEqual(calls[3]['observations']['user_feedback'],'')
            self.assertEqual(report['status'],'needs_review')
            self.assertNotIn('template',report)


if __name__=='__main__':unittest.main()
