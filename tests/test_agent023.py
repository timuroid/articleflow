import copy
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from docx.shared import Pt
from articleflow.agent023 import extract_input, build_package
from articleflow.rules023 import semantic_spans, compile_plan, fit, validate_rules
from articleflow.semantic023 import expand_groups, student_messages
from articleflow.catalogue import render_catalogue
from articleflow.docx_engine import DocumentError, content_fingerprint, digest
from articleflow import provider023


class SemanticRulesTests(unittest.TestCase):
    def test_character_rules_use_neighbor_context_without_source_ids(self):
        examples=[]
        for previous,style in [('heading','lead-font'),('body','normal-font')]:
            examples.append({'kind':'character','when':{'role':'body','bold':False,'italic':False,'vertical':'baseline','hyperlink':False,'previous_role':previous},'set':{'style':style},'id':previous})
        rules,conflicts=fit(examples)
        self.assertEqual(conflicts,[])
        self.assertEqual({r['when']['previous_role'] for r in rules},{'heading','body'})

    def test_missing_layout_paragraph_after_table_is_restored_without_copying_text(self):
        from articleflow.agent023 import apply
        from articleflow.roundtrip022 import neutralize
        with tempfile.TemporaryDirectory() as td:
            root=Path(td);source=root/'reference.docx';d=Document()
            d.add_paragraph('Заголовок');d.add_table(rows=1,cols=1).cell(0,0).text='42'
            p=d.add_paragraph();p.paragraph_format.space_after=Pt(9)
            d.add_paragraph('Текст после таблицы');d.save(source)
            content=extract_input(source,root/'train');roles=['title','cell','spacer','body']
            labels={b['id']:r for b,r in zip(content['blocks'],roles)}
            build_package([{'extracted':root/'train','content':content,'semantics':semantic_spans(source)}],[labels],{r:r for r in roles},root/'template')
            neutralize(source,root/'input.docx')
            def caller(config,messages,directory,**kwargs):
                return {'groups':[{'start':1,'end':1,'role':'title'},{'start':2,'end':2,'role':'cell'},{'start':3,'end':3,'role':'body'}]}, {'api_attempts':0}
            with patch('articleflow.agent023.render_pdf',return_value=root/'mock.pdf'):
                report=apply(root/'input.docx',root/'template',root/'apply',{},caller=caller)
            self.assertEqual(report['status'],'rendered_unverified',report)
            self.assertEqual(report['inserted_layout_paragraphs'],1)
            self.assertTrue(report['content_exact']);self.assertTrue(report['nested_topology_exact'])
            self.assertEqual([p.text for p in Document(root/'apply/result.docx').paragraphs],['Заголовок','','Текст после таблицы'])

    def test_transfer_rebuilds_nested_table_without_reading_training_document(self):
        with tempfile.TemporaryDirectory() as td:
            root=Path(td)
            def make(path,title,number):
                d=Document();p=d.add_paragraph(title);p.paragraph_format.space_after=Pt(19)
                t=d.add_table(rows=2,cols=2);t.style='Table Grid'
                for i,row in enumerate(t.rows):
                    for j,c in enumerate(row.cells):c.text=f'{number+i+j}'
                nested=t.cell(1,1).add_table(rows=1,cols=2)
                nested.cell(0,0).text='Вложенный';nested.cell(0,1).text=str(number)
                p=d.add_paragraph('Продолжение исследования');p.paragraph_format.space_before=Pt(13)
                d.save(path)
            make(root/'train.docx','Один заголовок',7)
            content=extract_input(root/'train.docx',root/'train')
            labels={b['id']:('title' if i==0 else 'after_table' if i==len(content['blocks'])-1 else 'cell') for i,b in enumerate(content['blocks'])}
            semantics=semantic_spans(root/'train.docx')
            package,conflicts=build_package([{'extracted':root/'train','content':content,'semantics':semantics}],
                [labels],{r:r for r in set(labels.values())},root/'template')
            self.assertEqual(conflicts,[])
            make(root/'new.docx','Другой заголовок с новым текстом',53)
            content=extract_input(root/'new.docx',root/'input')
            template=json.loads((root/'template/rules.json').read_text())
            # Erase every training artifact before runtime compilation.
            import shutil
            shutil.rmtree(root/'train');(root/'train.docx').unlink()
            plan,issues=compile_plan(template,content,package,labels,semantic_spans(root/'new.docx'))
            self.assertEqual(issues,[])
            plan.update(template_sha256=digest((root/'template/rules.json').read_bytes()),input_sha256=digest((root/'input/content/content.json').read_bytes()))
            render_catalogue(root/'template',root/'input/content',plan,root/'output.docx')
            self.assertEqual(content_fingerprint(root/'new.docx'),content_fingerprint(root/'output.docx'))
            self.assertNotIn('Один заголовок',json.dumps(package,ensure_ascii=False))

    def test_conflicting_same_context_never_becomes_majority_rule(self):
        examples=[{'kind':'paragraph','when':{'role':'body'},'set':{'style':s},'id':str(i)} for i,s in enumerate(['a','a','b'])]
        rules,conflicts=fit(examples)
        self.assertEqual(rules,[]);self.assertEqual(len(conflicts),1)

    def test_selector_cannot_memorize_reference(self):
        with self.assertRaisesRegex(DocumentError,'RULE_SELECTOR_FORBIDDEN'):
            validate_rules([{'kind':'paragraph','when':{'text':'secret'},'set':{'style':'a'}}],{})

    def test_group_coverage_and_overlap(self):
        blocks=[{'id':'x'},{'id':'y'}];roles={'body':'Основной текст'}
        with self.assertRaisesRegex(DocumentError,'COVERAGE'):
            expand_groups({'groups':[{'start':1,'end':1,'role':'body'}]},blocks,roles)
        with self.assertRaisesRegex(DocumentError,'OVERLAP'):
            expand_groups({'groups':[{'start':1,'end':2,'role':'body'},{'start':2,'end':2,'role':'body'}]},blocks,roles)

    def test_student_does_not_receive_original_styles_or_assignments(self):
        content={'blocks':[{'id':'x','text':'Текст','atoms':[],'secret_oracle':'bad'}],
                 'structure':[{'kind':'paragraph','id':'x'}]}
        prompt=json.dumps(student_messages({'body':'Основной текст'},content))
        self.assertNotIn('bad',prompt);self.assertNotIn('secret_oracle',prompt)


class BudgetTests(unittest.TestCase):
    def test_unknown_error_retains_reserve_and_does_not_retry(self):
        with tempfile.TemporaryDirectory() as td:
            root=Path(td);ledger=root/'ledger.json'
            ledger.write_text(json.dumps({'baseline':{'usage_usd':'2'},'additional_cap_usd':'1','attempts':[]}))
            calls=[]
            def api(key,endpoint,body=None):
                calls.append(endpoint)
                if endpoint=='models':return 200,{'data':[{'id':'test','pricing':{'prompt':'0.000001','completion':'0.000001'}}]}
                return 429,{'error':{'message':'limited'}}
            with patch.object(provider023,'read_key',return_value='not-a-real-key'),patch.object(provider023,'account',return_value={'usage_usd':'2','available_usd':'1','at':'now'}),patch.object(provider023,'api',side_effect=api):
                with self.assertRaisesRegex(DocumentError,'PROVIDER_HTTP_429'):
                    provider023.call({'key_file':str(root/'key'),'ledger':str(ledger),'model':'test'},[],root/'run',max_tokens=256)
            state=json.loads(ledger.read_text())
            self.assertEqual(calls.count('chat/completions'),1)
            self.assertIsNone(state['attempts'][0]['cost_usd'])
            self.assertGreater(float(state['attempts'][0]['reserve_usd']),0)


class WebLearningTests(unittest.TestCase):
    def test_consent_questions_and_package_hash_are_enforced(self):
        from fastapi.testclient import TestClient
        from articleflow.web017 import create_app
        with tempfile.TemporaryDirectory() as td:
            root=Path(td);source=root/'source.docx';doc=Document();doc.add_paragraph('Название статьи');doc.save(source)
            def fake_run(paths,folder,teacher,student,**kwargs):
                package=folder/'iteration-1/template';package.mkdir(parents=True)
                for name,data in [('rules.json','{}'),('semantic.json',json.dumps({'roles':{'title':'Название статьи'}})),('resources.zip','test')]:
                    (package/name).write_text(data)
                return {'schema':'teacher-student-023','status':'reference_verified','template':str(package),
                        'iterations':[{'number':1,'applications':[]}],'source_hashes':[digest(p.read_bytes()) for p in paths],
                        'questions':[{'role':'title','question':'Заголовок обязателен?'}],
                        'negative_controls':[{'detected':True}],'student_model':'qwen/qwen3.8-27b','held_out_transfer':'not_measured'}
            with patch('articleflow.web017.model_config',return_value={'model':'test'}),patch('articleflow.templateagent022.run',return_value={'status':'needs_review'}),patch('articleflow.agent023.run',side_effect=fake_run):
                with TestClient(create_app(root/'web')) as client:
                    headers={'X-CSRF-Token':client.get('/api/v017/config').json()['csrf_token']}
                    response=client.post('/api/v017/templates',data={'name':'Test'},files={'files':('test.docx',source.read_bytes())},headers=headers)
                    self.assertEqual(response.status_code,200,response.text);tid=response.json()['id']
                    endpoint=f'/api/v017/templates/{tid}'
                    self.assertEqual(client.post(endpoint+'/learn',json={},headers=headers).status_code,400)
                    response=client.post(endpoint+'/learn',json={'consent_external':True,'mode':'assisted'},headers=headers)
                    self.assertEqual(response.status_code,200,response.text);rid=response.json()['id']
                    self.assertEqual(client.post(endpoint+'/approve-agent',json={'run_id':rid},headers=headers).status_code,409)
                    response=client.post(endpoint+'/approve-agent',json={'run_id':rid,'answers':{'title':True}},headers=headers)
                    self.assertEqual(response.status_code,200,response.text)
                    self.assertEqual(client.post(endpoint+'/requirements',json={'expected_version':1,'answers':{'title':False}},headers=headers).status_code,409)
                    self.assertEqual(client.post(endpoint+'/requirements',json={'expected_version':2,'answers':{'title':False}},headers=headers).status_code,200)
                    version=client.get(endpoint).json();self.assertEqual(version['spec']['agent023']['required_roles'],[])
                    self.assertEqual(client.post('/api/v017/batches',data={'template_id':tid},files={'files':('test.docx',source.read_bytes())},headers=headers).status_code,400)
                    def finish_job(store,job,*args):
                        store.execute('UPDATE jobs SET status=?,report_json=? WHERE id=?',('draft',json.dumps({'schema':'application-023','questions':[]}),job['id']))
                    with patch('articleflow.webagent023.process_job',side_effect=finish_job) as worker:
                        response=client.post('/api/v017/batches',data={'template_id':tid,'use_llm':'true','consent_external':'true'},files={'files':('test.docx',source.read_bytes())},headers=headers)
                        self.assertEqual(response.status_code,202,response.text);jid=response.json()['jobs'][0]
                        self.assertEqual(client.patch(f'/api/v017/jobs/{jid}/decisions',json={'expected_revision':0,'assignments':{}},headers=headers).status_code,409)
                        self.assertEqual(worker.call_count,1)


if __name__=='__main__':unittest.main()
