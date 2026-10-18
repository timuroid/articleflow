"""A known conflicting template must be refined before a paid student call."""
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from docx import Document
from docx.shared import Pt
from articleflow.agent023 import run


class TriageTests(unittest.TestCase):
    def test_conflict_triage_reuses_teacher_feedback_without_student(self):
        with tempfile.TemporaryDirectory() as td:
            base=Path(td);source=base/'reference.docx';document=Document()
            a=document.add_paragraph('Первый абзац');a.runs[0].font.size=Pt(9)
            b=document.add_paragraph('Второй абзац');b.runs[0].font.size=Pt(17)
            document.save(source);teacher_feedback=[]
            def caller(config,messages,directory,**kwargs):
                if config['model']!='teacher':raise AssertionError('Student must be skipped while training conflicts remain')
                payload=json.loads(messages[1]['content']);teacher_feedback.append(payload['feedback'])
                return {'roles':{'body':'Обычный абзац текста'},'groups':[{'start':1,'end':2,'role':'body'}],'questions':[]},{}
            conflict={'kind':'character','when':{'role':'body'},'examples':[],
                      'differing_properties':['rPr/sz'],'variants':2}
            with patch('articleflow.agent023.build_package',return_value=({},[conflict])):
                result=run([source],base/'result',{'model':'teacher','triage_conflicts':True},{'model':'student'},
                           max_iterations=2,caller=caller,adaptive=True)
            self.assertEqual(result['status'],'needs_review')
            self.assertEqual(len(teacher_feedback),2)
            self.assertTrue(teacher_feedback[1]['conflicts'])
            self.assertEqual([row['code'] for row in result['iterations']],['TRAINING_CONFLICTS','SEMANTIC_PLATEAU'])
            self.assertTrue(all(not row['applications'] for row in result['iterations']))
            self.assertFalse(list((base/'result').rglob('student/validated.json')))
