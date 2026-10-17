import json,tempfile,unittest
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from articleflow.agent023 import extract_input
from articleflow.catalogue import render_catalogue,placement_fingerprint
from articleflow.docx_engine import DocumentError,load_package,parse_xml,xml_bytes,write_package,NS
from articleflow.integrity023 import check
from articleflow.fields030 import validate_fields

class FieldTests(unittest.TestCase):
    def add_field(self,p,instruction='PAGE'):
        for name,value in [('fldChar','begin'),('instrText',instruction),('fldChar','separate'),('t','1'),('fldChar','end')]:
            n=OxmlElement('w:'+name)
            if name=='fldChar':n.set(qn('w:fldCharType'),value)
            else:n.text=value
            p.add_run()._r.append(n)

    def test_fields_and_content_control_reconstruct_without_text_loss(self):
        with tempfile.TemporaryDirectory() as td:
            base=Path(td);src=base/'source.docx';d=Document();p=d.add_paragraph('Page ');self.add_field(p)
            wrapper=OxmlElement('w:sdt');props=OxmlElement('w:sdtPr');wrapper.append(props);content=OxmlElement('w:sdtContent');wrapper.append(content)
            d._element.body.remove(p._p);content.append(p._p);d._element.body.insert(0,wrapper);d.save(src)
            extract_input(src,base/'extracted');plan=json.loads((base/'extracted/oracle-plan.json').read_text())
            target=base/'result.docx';render_catalogue(base/'extracted/template',base/'extracted/content',plan,target)
            self.assertTrue(all(check(src,target).values()))
            parts=load_package(target);root=parse_xml(parts['word/document.xml']);root.find('.//w:instrText',NS).text='NUMPAGES';parts['word/document.xml']=xml_bytes(root);write_package(target,parts)
            self.assertFalse(check(src,target)['placements_exact'])

    def test_unsafe_and_unbalanced_fields_rejected_on_intake(self):
        with tempfile.TemporaryDirectory() as td:
            src=Path(td)/'source.docx'
            for instruction in ['DDEAUTO server topic','INCLUDETEXT https://example.org/input','INCLUDEPICTURE https://example.org/image']:
                d=Document();self.add_field(d.add_paragraph(),instruction);d.save(src)
                with self.assertRaisesRegex(DocumentError,'FIELD_INSTRUCTION_UNSUPPORTED'):load_package(src,allow_unsupported=True)
            d=Document();p=d.add_paragraph();self.add_field(p);p._p.remove(p._p[-1]);d.save(src)
            with self.assertRaisesRegex(DocumentError,'FIELD_STRUCTURE_INVALID'):load_package(src)
