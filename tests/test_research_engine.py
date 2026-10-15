"""Synthetic tests only. No manuscript or paid API access is needed."""
import copy
import io
import json
import tempfile
import unittest
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from PIL import Image
from lxml import etree as ET

from articleflow.docx_engine import (
    DocumentError, NS, build_template, compare_content, digest, inspect_document,
    load_package, lossless_split, lossless_rebuild, parse_xml, render_roles,
    validate_plan, write_package, xml_bytes,
)
from articleflow.research import reserve_budget, settle_budget, one_model_call, plan_schema, validate_model_choice


class DocumentTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.source = self.root / "source.docx"
        d = Document()
        p = d.add_paragraph()
        p.add_run("Определение").bold = True
        p.add_run(" — точный текст 1.25.\tИндекс ")
        p.add_run("2").font.superscript = True
        table = d.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Значение"
        table.cell(0, 1).text = "Единица"
        table.cell(1, 0).text = "3.1415"
        table.cell(1, 1).text = "м"
        image = io.BytesIO()
        Image.new("RGB", (10, 10), "red").save(image, format="PNG")
        image.seek(0)
        d.add_picture(image)
        d.save(self.source)

    def tearDown(self):
        self.temp.cleanup()

    def test_instance_roundtrip_preserves_payload_and_not_a_model_trial(self):
        split = self.root / "split"
        manifest = lossless_split(self.source, split)
        target = self.root / "rebuilt.docx"
        lossless_rebuild(split, target)
        self.assertTrue(compare_content(self.source, target)["content_exact"])
        self.assertFalse(manifest["eligible_for_model_benchmark"])
        self.assertNotIn("точный текст", (split / "format-catalogue.json").read_text())

    def test_content_validator_detects_single_digit_change(self):
        parts = load_package(self.source)
        d = parse_xml(parts["word/document.xml"])
        for node in d.findall(".//w:t", NS):
            if node.text and "1.25" in node.text:
                node.text = node.text.replace("1.25", "1.26")
        parts["word/document.xml"] = xml_bytes(d)
        target = self.root / "changed.docx"
        write_package(target, parts)
        self.assertFalse(compare_content(self.source, target)["content_exact"])

    def test_content_validator_detects_changed_image_bytes(self):
        parts = load_package(self.source)
        media = next(k for k in parts if k.startswith("word/media/"))
        parts[media] += b"changed"
        target = self.root / "changed.docx"
        write_package(target, parts)
        self.assertFalse(compare_content(self.source, target)["checks"]["media"])

    def test_content_validator_protects_superscript(self):
        parts = load_package(self.source)
        d = parse_xml(parts["word/document.xml"])
        for node in d.findall(".//w:vertAlign", NS):
            node.getparent().remove(node)
        parts["word/document.xml"] = xml_bytes(d)
        target = self.root / "changed.docx"
        write_package(target, parts)
        self.assertFalse(compare_content(self.source, target)["content_exact"])
        self.assertFalse(compare_content(self.source, target)["checks"]["protected_vertical"])

    def test_xml_entities_are_rejected(self):
        with self.assertRaisesRegex(DocumentError, "XML_DTD"):
            parse_xml(b'<!DOCTYPE x [<!ENTITY e SYSTEM "file:///etc/passwd">]><x>&e;</x>')

    def test_zip_traversal_is_rejected(self):
        p = self.root / "bad.docx"
        with ZipFile(p, "w", ZIP_DEFLATED) as z:
            z.writestr("../outside", "x")
        with self.assertRaisesRegex(DocumentError, "ZIP_PATH"):
            load_package(p)

    def test_renderer_preserves_structured_content_with_oracle_inline_spans(self):
        work = self.root / "content"
        inventory = inspect_document(self.source, work)
        self.assertEqual(inventory["media_count"], 1)
        evidence = json.loads((work / "evidence.json").read_text())
        assignments = {b["id"]: "body" for b in evidence["blocks"]}
        template = build_template(work, assignments, {"body": "Synthetic body"}, self.root / "template.json")
        plan = {"blocks": [{"id": b["id"], "role": "body", "spans": [{"start": s["start"], "end": s["end"], "bold": s.get("b"), "italic": s.get("i"), "vertical": s.get("vertAlign")} for s in b["source_spans"]]} for b in evidence["blocks"]]}
        target = self.root / "rendered.docx"
        render_roles(work, template, plan, target)
        self.assertTrue(compare_content(self.source, target)["content_exact"])
        content = json.loads((work / "content.json").read_text())
        self.assertNotIn("source_spans", content["blocks"][0])
        self.assertNotIn("ppr", content["blocks"][0])


class PlanTests(unittest.TestCase):
    def setUp(self):
        self.content = {"blocks": [{"id": "b-1", "text": "Термин e\u0301"}]}
        self.template = {"roles": {"body": {}}}
        self.plan = {"blocks": [{"id": "b-1", "role": "body", "spans": []}]}

    def test_missing_and_duplicate_blocks_rejected(self):
        for blocks in [[], self.plan["blocks"] * 2]:
            with self.assertRaises(DocumentError):
                validate_plan(self.content, self.template, {"blocks": blocks})

    def test_unknown_role_rejected(self):
        self.plan["blocks"][0]["role"] = "invented"
        with self.assertRaisesRegex(DocumentError, "PLAN_ROLE"):
            validate_plan(self.content, self.template, self.plan)

    def test_span_bounds_and_combining_character_rejected(self):
        for start, end in [(0, 100), (-1, 2), (0, 8)]:
            self.plan["blocks"][0]["spans"] = [{"start": start, "end": end, "bold": True, "italic": None, "vertical": None}]
            with self.assertRaises(DocumentError):
                validate_plan(self.content, self.template, self.plan)

    def test_overlapping_spans_rejected(self):
        self.plan["blocks"][0]["spans"] = [{"start": 0, "end": 5, "bold": True, "italic": None, "vertical": None}, {"start": 4, "end": 6, "bold": None, "italic": True, "vertical": None}]
        with self.assertRaisesRegex(DocumentError, "OVERLAP"):
            validate_plan(self.content, self.template, self.plan)


class BudgetTests(unittest.TestCase):
    def test_cache_write_price_is_reserved_before_any_http_call(self):
        with tempfile.TemporaryDirectory() as temp, patch("articleflow.research.httpx.Client") as client:
            base = Path(temp)
            with self.assertRaisesRegex(DocumentError, "BUDGET_EXCEEDED"):
                one_model_call(model="synthetic", messages=[], schema=plan_schema(["body"]),
                    pricing={"prompt": "0.000001", "completion": "0", "input_cache_write": "0.0001"},
                    key_file=None, run_dir=base / "run", ledger=base / "ledger.json", cap_usd="0.1", max_tokens=1)
            client.assert_not_called()
            self.assertFalse((base / "run/request.json").exists())

    def test_budget_rejects_nonfinite_and_negative_values(self):
        with tempfile.TemporaryDirectory() as temp:
            for amount, cap in [("-1", "5"), ("1", "NaN"), ("1", "0")]:
                with self.assertRaisesRegex(DocumentError, "INVALID_BUDGET"):
                    reserve_budget(Path(temp) / "ledger.json", Decimal(amount), Decimal(cap), "a")

    def test_unknown_cost_keeps_reservation(self):
        with tempfile.TemporaryDirectory() as temp:
            ledger = Path(temp) / "ledger.json"
            reserve_budget(ledger, Decimal("0.8"), Decimal("1"), "a")
            settle_budget(ledger, "a", None, "unknown")
            with self.assertRaisesRegex(DocumentError, "BUDGET_EXCEEDED"):
                reserve_budget(ledger, Decimal("0.3"), Decimal("1"), "b")

    def test_invalid_output_does_not_trigger_retry_or_leak_key(self):
        class Response:
            status_code = 200
            def json(self):
                return {"model": "synthetic", "usage": {"cost": 0.01}, "choices": [{"finish_reason": "stop", "message": {"content": "not JSON"}}]}
        with tempfile.TemporaryDirectory() as temp:
            base = Path(temp)
            with patch.dict("os.environ", {"OPENROUTER_API_KEY": "synthetic-test-secret"}), patch("articleflow.research.httpx.Client") as cls:
                client = cls.return_value.__enter__.return_value
                client.post.return_value = Response()
                result = one_model_call(model="synthetic", messages=[{"role": "user", "content": "synthetic"}], schema=plan_schema(["body"]), pricing={"prompt": "0.000001", "completion": "0.000001"}, key_file=None, run_dir=base / "run", ledger=base / "ledger.json", cap_usd="1")
                self.assertEqual(client.post.call_count, 1)
                self.assertEqual(result["status"], "invalid_model_output")
                self.assertEqual(result["repair_calls"], 0)
                for p in base.rglob("*.json"):
                    self.assertNotIn("synthetic-test-secret", p.read_text())


class ModelEligibilityTests(unittest.TestCase):
    def test_cutoff_is_exclusive_and_unverified_models_are_rejected(self):
        registry = {"release_after": "2026-06-30", "models": [{"id": "candidate", "release_date": "2026-06-30", "source_type": "official_release_log", "source": "https://example.org/releases"}]}
        with self.assertRaisesRegex(DocumentError, "MODEL_TOO_OLD"):
            validate_model_choice("candidate", registry)
        registry["models"][0]["release_date"] = "2026-07-01"
        self.assertEqual(validate_model_choice("candidate", registry)["id"], "candidate")
        with self.assertRaisesRegex(DocumentError, "MODEL_RELEASE_NOT_VERIFIED"):
            validate_model_choice("unknown", registry)

    def test_current_registry_excludes_old_baseline(self):
        path = Path(__file__).resolve().parents[1] / "docs/research/models.json"
        registry = json.loads(path.read_text())
        validate_model_choice("openai/gpt-6.1-sol", registry)
        with self.assertRaisesRegex(DocumentError, "MODEL_RELEASE_NOT_VERIFIED"):
            validate_model_choice("openai/gpt-4.1", registry)


if __name__ == "__main__":
    unittest.main()
