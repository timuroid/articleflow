import tempfile,unittest
from pathlib import Path
from docx import Document
from docx.enum.text import WD_BREAK
from articleflow.agent023 import extract_input
from articleflow.order024 import arrange
from articleflow.agent028 import permute

class BreakOrderingTests(unittest.TestCase):
    def test_plain_line_break_does_not_prevent_moving_text(self):
        with tempfile.TemporaryDirectory() as td:
            root=Path(td);src=root/'source.docx';d=Document();d.add_paragraph('Second\ncontinued');d.add_paragraph('First');d.save(src)
            c=extract_input(src,root/'input');labels=dict(zip([b['id'] for b in c['blocks']],['second','first']))
            contract={'schema':'role-order-024','status':'observed','edges':[{'before':'first','after':'second'}]}
            _,order,moves,issues=arrange(c,labels,contract,src)
            self.assertEqual(order,[1,0]);self.assertFalse(issues);self.assertTrue(moves)

    def test_movement_within_page_is_allowed_but_crossing_break_is_not(self):
        with tempfile.TemporaryDirectory() as td:
            root=Path(td);src=root/'source.docx';d=Document();d.add_paragraph().add_run().add_break(WD_BREAK.PAGE);d.add_paragraph('Second');d.add_paragraph('First');d.save(src)
            c=extract_input(src,root/'input');labels=dict(zip([b['id'] for b in c['blocks']],['divider','second','first']))
            contract={'schema':'role-order-024','status':'observed','edges':[{'before':'first','after':'second'}]}
            _,order,_,issues=arrange(c,labels,contract,src);self.assertEqual(order,[0,2,1]);self.assertFalse(issues)
            contract['edges']=[{'before':'first','after':'divider'}]
            _,order,_,issues=arrange(c,labels,contract,src);self.assertEqual(order,[0,1,2]);self.assertTrue(issues)

    def test_corruptor_does_not_move_a_paragraph_containing_page_break(self):
        with tempfile.TemporaryDirectory() as td:
            root=Path(td);src=root/'source.docx';d=Document();d.add_paragraph('Fixed').add_run().add_break(WD_BREAK.PAGE);d.add_paragraph('Second');d.add_paragraph('First');d.save(src)
            result=permute(src,root/'mutated.docx');self.assertEqual(result['swapped'],[1,2])
            d=Document();d.add_paragraph('Same');d.add_paragraph('Same');d.save(src)
            self.assertFalse(permute(src,root/'same.docx')['applied'])
