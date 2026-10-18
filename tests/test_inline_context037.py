"""Paragraph-role constraints for one-pass inline predictions."""
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document

from articleflow.agent023 import build_package, extract_input
from articleflow.docx_engine import DocumentError
from articleflow.inline026 import overlay
from articleflow.rules023 import semantic_spans
from articleflow.sandbox034 import use
from articleflow.semantic023 import student_messages


class InlineContextTests(unittest.TestCase):
    def test_package_learns_body_only_dash_without_capturing_reference_text(self):
        with tempfile.TemporaryDirectory() as td:
            root=Path(td)
            source=root/'reference.docx'
            doc=Document()
            body='Описание — пояснение'
            caption='Рисунок 1 — описание'
            doc.add_paragraph(body)
            doc.add_paragraph(caption)
            doc.save(source)
            extracted=root/'extracted'
            content=extract_input(source,extracted)
            ids=[block['id'] for block in content['blocks']]
            labels={ids[0]:'body',ids[1]:'figure_caption'}
            dash=body.index('—')
            annotations={ids[0]:[{'start':dash,'end':dash+1,
                                  'role':'explanatory_dash'}],ids[1]:[]}
            semantics=overlay(semantic_spans(source),annotations)
            inline_roles={'explanatory_dash':'Explanatory dash in body prose',
                          'unobserved':'Defined without a demonstrated occurrence'}
            package,_=build_package([{'source':source,'extracted':extracted,
                                      'content':content,'semantics':semantics}],
                                    [labels],{'body':'Main prose','figure_caption':'Figure caption'},
                                    root/'template',inline_roles)
            expected={'explanatory_dash':['body'],'unobserved':[]}
            self.assertEqual(package['inline_role_contexts'],expected)
            self.assertEqual(json.loads((root/'template/semantic.json').read_text())
                             ['inline_role_contexts'],expected)
            self.assertNotIn(body,json.dumps(expected,ensure_ascii=False))
            self.assertNotIn(caption,json.dumps(expected,ensure_ascii=False))
            messages=student_messages(package['roles'],content,inline_roles,
                                      package['inline_role_contexts'])
            request=json.loads(messages[1]['content'])
            self.assertEqual(request['inline_role_contexts'],expected)
            self.assertIn('predicted role appears',messages[0]['content'])
            self.assertIn('same punctuation or words',messages[0]['content'])
            self.assertNotIn('inline_role_contexts',
                             student_messages(package['roles'],content,inline_roles)[1]['content'])

    def test_malformed_context_map_fails_closed(self):
        content={'blocks':[{'id':'b-000001','text':'Plain','atoms':[]}],
                 'structure':[{'kind':'paragraph','id':'b-000001'}]}
        with self.assertRaisesRegex(DocumentError,'INLINE_CONTEXT_SCHEMA'):
            student_messages({'body':'Prose'},content,
                             {'dash':'Semantic dash'},
                             {'dash':['figure_caption']})

    def test_sandbox_candidate_package_gets_supervisor_context_hint(self):
        with tempfile.TemporaryDirectory() as td:
            destination=Path(td)/'template'
            sample={'content':{'blocks':[{'id':'one'}]},
                    'semantics':{'one':[{'features':{'inline_role':'dash'}}]}}
            def fake_build(engine,samples,labels,definitions,target,inline_roles,span_grammar):
                target.mkdir()
                (target/'semantic.json').write_text(json.dumps({'schema':'candidate'}))
                return {'schema':'candidate'},[]
            with patch('articleflow.engine034.build',side_effect=fake_build),use(object()):
                package,_=build_package([sample],[{'one':'body'}],{'body':'Main prose'},
                                        destination,{'dash':'Explanatory dash'})
            self.assertEqual(package['inline_role_contexts'],{'dash':['body']})
            self.assertEqual(json.loads((destination/'semantic.json').read_text())
                             ['inline_role_contexts'],{'dash':['body']})


if __name__=='__main__':
    unittest.main()
