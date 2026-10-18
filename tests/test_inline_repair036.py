"""Offline correction tests: the second model may replace only bad inline rows."""
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document

from articleflow.agent023 import apply, build_package, extract_input
from articleflow.engine034 import apply as engine_apply
from articleflow.docx_engine import DocumentError
from articleflow.inline026 import overlay
from articleflow.inline_repair036 import prepare, messages, merge
from articleflow.rules023 import semantic_spans


class InlineRepairTests(unittest.TestCase):
    def fixture(self, root):
        source = root/'reference.docx'
        document = Document()
        paragraph = document.add_paragraph()
        paragraph.add_run('alpha').italic = True
        paragraph.add_run(' beta').italic = True
        document.save(source)
        content = extract_input(source, root/'extract')
        block = content['blocks'][0]['id']
        semantics = overlay(semantic_spans(source), {block: [
            {'start': 0, 'end': 5, 'role': 'term'},
            {'start': 6, 'end': 10, 'role': 'term'}]})
        _, conflicts = build_package([{'source': source, 'content': content,
                                       'extracted': root/'extract', 'semantics': semantics}],
                                     [{block: 'body'}], {'body': 'Body paragraph'},
                                     root/'template', {'term': 'Technical term'})
        self.assertEqual(conflicts, [])
        return source, root/'template', content

    @staticmethod
    def first_response():
        return {'groups': [{'start': 1, 'end': 1, 'role': 'body'}],
                'inline': [{'paragraph': 1, 'quote': 'alpha', 'occurrence': 1, 'role': 'term'},
                           {'paragraph': 1, 'quote': 'beta', 'occurrence': 2, 'role': 'term'}]}

    @staticmethod
    def correction():
        return {'repairs': [{'index': 1, 'row': {
            'paragraph': 1, 'quote': 'beta', 'occurrence': 1, 'role': 'term'}}]}

    def test_helper_freezes_valid_row_and_uses_exact_paragraph_text(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            _, _, content = self.fixture(root)
            original = self.first_response()
            context = prepare(original, content['blocks'], {'term': 'Technical term'})
            self.assertEqual(context['invalid'], {1: 'INLINE_ANCHOR_NOT_FOUND'})
            self.assertEqual(context['frozen'], (0,))
            prompt = messages(context)
            self.assertIn('alpha beta', prompt[1]['content'])
            self.assertIn('Technical term', prompt[1]['content'])
            merged = merge(context, self.correction(), {'body': 'Body paragraph'})
            self.assertEqual(merged['groups'], original['groups'])
            self.assertEqual(merged['inline'][0], original['inline'][0])
            self.assertEqual(merged['inline'][1]['occurrence'], 1)
            self.assertEqual(original['inline'][1]['occurrence'], 2)
            for malicious in [
                {'repairs': [{'index': 0, 'row': original['inline'][0]}]},
                {**self.correction(), 'groups': [{'start': 1, 'end': 1, 'role': 'other'}]},
                {'repairs': [{'index': 1, 'row': {**self.correction()['repairs'][0]['row'], 'paragraph': 2}}]},
            ]:
                with self.assertRaises(DocumentError):
                    merge(context, malicious, {'body': 'Body paragraph'})

    def test_apply_replays_one_model_call_without_changing_valid_rows_or_groups(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            source, template, _ = self.fixture(root)
            calls = []

            def caller(config, prompt, directory, **kwargs):
                calls.append(prompt)
                directory.mkdir(parents=True)
                response = self.first_response() if len(calls) == 1 else self.correction()
                try:
                    kwargs['validate'](response)
                except DocumentError as exc:
                    (directory/'meta.json').write_text(json.dumps({
                        'api_attempts': 1, 'status': str(exc), 'cost_usd': '0.001'}))
                    raise
                return response, {'api_attempts': 1, 'status': 'valid', 'cost_usd': '0.001'}

            with patch('articleflow.agent023.render_pdf', return_value=root/'dummy.pdf'):
                result = apply(source, template, root/'apply',
                               {'inline_repair_attempts': 1}, caller=caller)
            self.assertEqual(len(calls), 2)
            self.assertEqual(result['status'], 'rendered_unverified')
            self.assertEqual(result['first_pass']['code'], 'INLINE_ANCHOR_NOT_FOUND')
            self.assertFalse(result['strict_one_pass_eligible'])
            self.assertEqual(result['model_invocations'], 2)
            self.assertEqual(result['inline_repair']['status'], 'repaired')
            self.assertEqual(result['inline_repair']['frozen_rows'], 1)
            self.assertEqual(result['inline_repair']['corrected_rows'], 1)
            self.assertTrue((root/'apply/result.docx').exists())
            self.assertEqual(json.loads((root/'apply/inline-repair.json').read_text())['status'], 'repaired')
            effective = json.loads((root/'apply/effective-student.json').read_text())
            self.assertEqual(effective['groups'], self.first_response()['groups'])
            self.assertEqual(effective['inline'][0], self.first_response()['inline'][0])
            self.assertEqual(effective['inline'][1]['occurrence'], 1)
            self.assertFalse((root/'apply/student/validated.json').exists())

    def test_disabled_repair_keeps_first_pass_failure(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            source, template, _ = self.fixture(root)
            calls = []

            def caller(config, prompt, directory, **kwargs):
                calls.append(prompt)
                return kwargs['validate'](self.first_response()), {'api_attempts': 1}

            with self.assertRaisesRegex(DocumentError, 'INLINE_ANCHOR_NOT_FOUND'):
                apply(source, template, root/'apply', {}, caller=caller)
            self.assertEqual(len(calls), 1)
            self.assertFalse((root/'apply/result.docx').exists())

    def test_valid_first_pass_uses_same_effective_answer_path(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            source, template, _ = self.fixture(root)
            valid = {'groups': [{'start': 1, 'end': 1, 'role': 'body'}],
                     'inline': [{'paragraph': 1, 'quote': 'alpha', 'occurrence': 1, 'role': 'term'},
                                {'paragraph': 1, 'quote': 'beta', 'occurrence': 1, 'role': 'term'}]}
            calls = []

            def caller(config, prompt, directory, **kwargs):
                calls.append(prompt)
                return kwargs['validate'](valid), {'api_attempts': 1, 'status': 'valid'}

            with patch('articleflow.agent023.render_pdf', return_value=root/'dummy.pdf'):
                result = apply(source, template, root/'apply',
                               {'inline_repair_attempts': 1}, caller=caller)
            self.assertEqual(len(calls), 1)
            self.assertTrue(result['strict_one_pass_eligible'])
            self.assertEqual(json.loads((root/'apply/effective-student.json').read_text()), valid)

    def test_correction_cannot_edit_groups_or_create_output(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            source, template, _ = self.fixture(root)
            calls = []

            def caller(config, prompt, directory, **kwargs):
                calls.append(prompt)
                response = self.first_response() if len(calls) == 1 else {
                    **self.correction(), 'groups': [{'start': 1, 'end': 1, 'role': 'other'}]}
                return kwargs['validate'](response), {'api_attempts': 1}

            with self.assertRaisesRegex(DocumentError, 'INLINE_REPAIR_SCHEMA'):
                apply(source, template, root/'apply',
                      {'inline_repair_attempts': 1}, caller=caller)
            self.assertEqual(len(calls), 2)
            self.assertFalse((root/'apply/result.docx').exists())
            state = json.loads((root/'apply/inline-repair.json').read_text())
            self.assertEqual(state['status'], 'failed')
            self.assertEqual(state['code'], 'INLINE_REPAIR_SCHEMA')

    def test_non_inline_group_error_never_triggers_correction(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            source, template, _ = self.fixture(root)
            calls = []

            def caller(config, prompt, directory, **kwargs):
                calls.append(prompt)
                response = self.first_response()
                response['groups'][0]['role'] = 'other'
                return kwargs['validate'](response), {'api_attempts': 1}

            with self.assertRaisesRegex(DocumentError, 'SEMANTIC_GROUP_INVALID'):
                apply(source, template, root/'apply',
                      {'inline_repair_attempts': 1}, caller=caller)
            self.assertEqual(len(calls), 1)

    def test_sandboxed_profile_repairs_in_supervisor_with_same_budget_ledger(self):
        with tempfile.TemporaryDirectory() as td:
            root=Path(td);source,template,_=self.fixture(root)
            calls=[]
            class Engine:
                hashes={'rules023.py':'pinned'}
                applied=[]
                def invoke(self,operation,payload,inputs,destination):
                    if operation=='extract':
                        return extract_input(inputs['source.docx'],destination/'extracted')
                    self.assert_apply(payload['response'])
                    self.applied.append(payload['response'])
                    (destination/'application').mkdir(parents=True)
                    return {'status':'needs_review'}
                def assert_apply(self,response):
                    assert response['groups']==InlineRepairTests.first_response()['groups']
                    assert response['inline'][0]==InlineRepairTests.first_response()['inline'][0]
                    assert response['inline'][1]['occurrence']==1
            engine=Engine()
            config={'inline_repair_attempts':1,'ledger':str(root/'same-ledger.json'),'model':'test-model'}
            def caller(actual,prompt,directory,**kwargs):
                self.assertIs(actual,config)
                calls.append((actual['ledger'],prompt))
                directory.mkdir(parents=True)
                response=self.first_response() if len(calls)==1 else self.correction()
                return kwargs['validate'](response),{'api_attempts':1,'cost_usd':'0.001'}
            result=engine_apply(engine,source,template,root/'profile',config,caller)
            self.assertEqual(len(calls),2)
            self.assertEqual([c[0] for c in calls],[config['ledger']]*2)
            self.assertEqual(len(engine.applied),1)
            self.assertEqual(result['inline_repair']['status'],'repaired')
            self.assertEqual(result['model_invocations'],2)
            self.assertFalse(result['strict_one_pass_eligible'])
            self.assertEqual(json.loads((root/'profile/inline-repair.json').read_text())['frozen_rows'],1)
            self.assertEqual(json.loads((root/'profile/effective-student.json').read_text()),engine.applied[0])

    def test_sandboxed_profile_rejects_group_change_before_worker(self):
        with tempfile.TemporaryDirectory() as td:
            root=Path(td);source,template,_=self.fixture(root)
            class Engine:
                hashes={}
                applied=0
                def invoke(self,operation,payload,inputs,destination):
                    if operation=='extract':
                        return extract_input(inputs['source.docx'],destination/'extracted')
                    self.applied+=1
                    raise AssertionError('invalid correction reached sandbox worker')
            engine=Engine();calls=[]
            def caller(config,prompt,directory,**kwargs):
                calls.append(prompt);directory.mkdir(parents=True)
                response=self.first_response() if len(calls)==1 else {
                    **self.correction(),'groups':[{'start':1,'end':1,'role':'other'}]}
                return kwargs['validate'](response),{'api_attempts':1}
            with self.assertRaisesRegex(DocumentError,'INLINE_REPAIR_SCHEMA'):
                engine_apply(engine,source,template,root/'profile',
                             {'inline_repair_attempts':1,'ledger':str(root/'ledger.json')},caller)
            self.assertEqual(len(calls),2)
            self.assertEqual(engine.applied,0)
            self.assertEqual(json.loads((root/'profile/inline-repair.json').read_text())['status'],'failed')
            self.assertFalse((root/'profile/effective-student.json').exists())

    def test_sandboxed_profile_valid_first_pass_never_calls_repair(self):
        with tempfile.TemporaryDirectory() as td:
            root=Path(td);source,template,_=self.fixture(root)
            class Engine:
                hashes={}
                def invoke(self,operation,payload,inputs,destination):
                    if operation=='extract':
                        return extract_input(inputs['source.docx'],destination/'extracted')
                    (destination/'application').mkdir(parents=True)
                    return {'status':'needs_review'}
            calls=[]
            def caller(config,prompt,directory,**kwargs):
                calls.append(config['ledger']);directory.mkdir(parents=True)
                response=self.first_response()
                response['inline'][1]['occurrence']=1
                return kwargs['validate'](response),{'api_attempts':1}
            result=engine_apply(Engine(),source,template,root/'profile',
                                {'inline_repair_attempts':1,'ledger':'shared-ledger'},caller)
            self.assertEqual(calls,['shared-ledger'])
            self.assertEqual(result['model_invocations'],1)
            self.assertTrue(result['strict_one_pass_eligible'])
            self.assertEqual(result['inline_repair']['status'],'not_needed')


if __name__ == '__main__':
    unittest.main()
