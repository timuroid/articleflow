import tempfile,unittest
from pathlib import Path
from docx import Document
from docx.shared import RGBColor
from articleflow.format032 import audit
from articleflow.docx_engine import load_package,write_package,parse_xml,xml_bytes,NS

class PropertyAuditTests(unittest.TestCase):
    def test_font_color_and_empty_boolean_properties_are_not_pixel_aliases(self):
        with tempfile.TemporaryDirectory() as td:
            root=Path(td);a=root/'a.docx';b=root/'b.docx';d=Document();r=d.add_paragraph().add_run('Same text');r.bold=True;d.save(a)
            d.save(b);self.assertTrue(audit(a,b)['property_exact'])
            r.bold=None;r.italic=True;d.save(b);self.assertFalse(audit(a,b)['property_exact'])
            r.italic=None;r.bold=True;r.font.color.rgb=RGBColor(1,2,3);d.save(b);self.assertFalse(audit(a,b)['property_exact'])
            r.font.color.rgb=None;r.font.name='Distinct font';d.save(b);self.assertFalse(audit(a,b)['property_exact'])

    def test_style_resources_are_checked_even_when_not_rendered_in_this_sample(self):
        with tempfile.TemporaryDirectory() as td:
            root=Path(td);a=root/'a.docx';b=root/'b.docx';d=Document();d.add_paragraph('Same text');d.save(a)
            parts=load_package(a);parts['word/settings.xml']+=b'\n';write_package(b,parts)
            self.assertFalse(audit(a,b)['resources_exact'])

    def test_paragraph_mark_format_is_audited_separately_from_visible_runs(self):
        from lxml import etree
        with tempfile.TemporaryDirectory() as td:
            root=Path(td);a=root/'a.docx';b=root/'b.docx'
            document=Document();document.add_paragraph('Same visible text');document.save(a)
            parts=load_package(a);xml=parse_xml(parts['word/document.xml'])
            paragraph=xml.find('w:body/w:p',NS)
            ppr=paragraph.find('w:pPr',NS)
            if ppr is None:
                ppr=etree.Element('{%s}pPr'%NS['w']);paragraph.insert(0,ppr)
            mark=etree.SubElement(ppr,'{%s}rPr'%NS['w'])
            etree.SubElement(mark,'{%s}rFonts'%NS['w'],{'{%s}eastAsia'%NS['w']:'KaiTi'})
            etree.SubElement(mark,'{%s}bCs'%NS['w'],{'{%s}val'%NS['w']:'1'})
            parts['word/document.xml']=xml_bytes(xml);write_package(b,parts)
            report=audit(a,b)
            self.assertTrue(report['resources_exact'])
            self.assertTrue(all(part['paragraphs_exact'] for part in report['parts']))
            self.assertFalse(report['property_exact'])
            self.assertFalse(next(part for part in report['parts'] if part['part']=='word/document.xml')['paragraph_marks_exact'])
