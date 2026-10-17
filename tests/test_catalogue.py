"""Synthetic compiler controls; no real manuscript, model or PDF tool required."""
import copy
import io
import json
import shutil
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from PIL import Image

from articleflow.catalogue import extract_catalogue, render_catalogue, placement_fingerprint, validate_catalogue_plan
from articleflow.docx_engine import (
    NS, DocumentError, compare_content, content_fingerprint, digest,
    load_package, parse_xml, write_package, xml_bytes,
)
from articleflow.research import render_pdf


class CatalogueTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.source = self.root / "source.docx"
        self.make_document(self.source, "Первый синтетический текст")
        self.package = self.root / "package"
        self.extract(self.source, self.package)
        self.plan = json.loads((self.package / "oracle-plan.json").read_text())
        self.template = json.loads((self.package / "template/rules.json").read_text())
        self.content = json.loads((self.package / "content/content.json").read_text())

    def make_document(self, path, title):
        doc = Document()
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(7)
        p.add_run(title).font.name = "Arial"
        r = p.add_run(" — значение ")
        r.font.size = Pt(11)
        r.italic = True
        p.add_run("2").font.superscript = True
        p.add_run(" и 3.1415")
        bookmark = OxmlElement("w:bookmarkStart")
        bookmark.set(qn("w:id"), "7")
        bookmark.set(qn("w:name"), "synthetic_anchor")
        p._p.append(bookmark)
        p.add_run(". Ссылка: ")
        link = OxmlElement("w:hyperlink")
        link.set(qn("r:id"), doc.part.relate_to("https://example.org/synthetic", RT.HYPERLINK, is_external=True))
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = "example"
        run.append(text)
        link.append(run)
        p._p.append(link)
        end = OxmlElement("w:bookmarkEnd")
        end.set(qn("w:id"), "7")
        p._p.append(end)
        math = OxmlElement("m:oMath")
        mr = OxmlElement("m:r")
        mt = OxmlElement("m:t")
        mt.text = "α+β"
        mr.append(mt)
        math.append(mr)
        p._p.append(math)
        empty = doc.add_paragraph()
        empty.paragraph_format.space_after = Pt(4)
        other = doc.add_paragraph("Другой абзац.\tЕдиница: м")
        other.paragraph_format.space_after = Pt(12)
        other.runs[0].bold = True
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).merge(table.cell(0, 1)).text = "Общий заголовок"
        table.cell(1, 0).text = "1.25"
        table.cell(1, 1).text = "2.50"
        for color in ["red", "blue"]:
            stream = io.BytesIO()
            Image.new("RGB", (8, 6), color).save(stream, format="PNG")
            stream.seek(0)
            doc.add_picture(stream)
        doc.save(path)

    def extract(self, source, output):
        tree = parse_xml(load_package(source)["word/document.xml"])
        count = len(tree.findall(".//w:body//w:p", NS))
        annotations = {f"b-{i + 1:06d}": "body" for i in range(count)}
        return extract_catalogue(source, annotations, {"roles": {"body": "Synthetic paragraph"}}, output)

    def render(self, plan=None):
        target = self.root / "result.docx"
        render_catalogue(self.package / "template", self.package / "content", plan or self.plan, target)
        return target

    def test_rebuild_preserves_payload_mixed_styles_and_object_positions(self):
        target = self.render()
        metrics = compare_content(self.source, target)
        self.assertTrue(metrics["content_exact"])
        self.assertTrue(metrics["inline_exact"])
        self.assertEqual(placement_fingerprint(self.source), placement_fingerprint(target))
        self.assertGreater(len(self.template["styles"]["paragraph"]), 1)
        self.assertGreater(len(self.template["styles"]["character"]), 1)

    def test_renderer_does_not_need_original_or_extraction_oracle_file(self):
        expected = content_fingerprint(self.source)
        self.source.unlink()
        (self.package / "oracle-plan.json").unlink()
        (self.package / "manifest.json").unlink()
        self.assertEqual(expected, content_fingerprint(self.render()))

    def test_template_contains_no_article_text_block_ids_or_image_bytes(self):
        raw = (self.package / "template/rules.json").read_text()
        self.assertNotIn("Первый синтетический", raw)
        self.assertNotIn("synthetic_anchor", raw)
        self.assertNotIn("b-000001", raw)
        self.assertNotIn("example.org", raw)
        with ZipFile(self.package / "template/resources.zip") as z:
            self.assertNotIn("word/document.xml", z.namelist())
            self.assertFalse(any(n.startswith("word/media/") for n in z.namelist()))
        with ZipFile(self.package / "content/payload.zip") as z:
            self.assertNotIn("word/document.xml", z.namelist())
            self.assertNotIn("word/styles.xml", z.namelist())
        serialized = json.dumps(self.content)
        self.assertNotIn("pPr", serialized)
        self.assertNotIn("tblGrid", serialized)
        self.assertNotIn("paragraph_style", serialized)

    def test_catalogue_resources_and_rules_are_reproducible(self):
        second = self.root / "second"
        self.extract(self.source, second)
        for relative in ["template/resources.zip", "template/rules.json", "content/content.json", "oracle-plan.json"]:
            self.assertEqual((self.package / relative).read_bytes(), (second / relative).read_bytes())

    def test_template_reused_for_different_text_without_reextracting_it(self):
        newer = self.root / "new.docx"
        self.make_document(newer, "Совершенно новая рукопись большей длины")
        other = self.root / "other"
        self.extract(newer, other)
        plan = json.loads((other / "oracle-plan.json").read_text())
        plan["template_sha256"] = self.plan["template_sha256"]
        shutil.rmtree(other / "template")
        target = self.root / "new-result.docx"
        render_catalogue(self.package / "template", other / "content", plan, target)
        self.assertTrue(compare_content(newer, target)["content_exact"])
        self.assertNotIn("Первый синтетический текст", "".join(content_fingerprint(target)["paragraphs"]))

    def test_missing_style_and_noncontiguous_spans_rejected(self):
        plan = copy.deepcopy(self.plan)
        plan["blocks"][0]["paragraph_style"] = "unknown"
        with self.assertRaisesRegex(DocumentError, "PLAN_STYLE_ROLE"):
            self.render(plan)
        plan = copy.deepcopy(self.plan)
        plan["blocks"][0]["spans"][0]["start"] = 1
        with self.assertRaisesRegex(DocumentError, "PLAN_SPAN_COVERAGE"):
            self.render(plan)

    def test_changed_input_or_resources_rejected_before_output(self):
        content = self.package / "content/content.json"
        content.write_bytes(content.read_bytes() + b" ")
        with self.assertRaisesRegex(DocumentError, "PLAN_INPUT_HASH"):
            self.render()
        content.write_bytes(content.read_bytes()[:-1])
        resource = self.package / "template/resources.zip"
        resource.write_bytes(resource.read_bytes() + b"changed")
        with self.assertRaisesRegex(DocumentError, "RESOURCE_HASH"):
            self.render()

    def test_duplicated_object_occurrence_rejected(self):
        content = copy.deepcopy(self.content)
        block = next(b for b in content["blocks"] if any(a["kind"] == "drawing" for a in b["atoms"]))
        block["atoms"] *= 2
        with self.assertRaisesRegex(DocumentError, "CONTENT_OBJECT_COVERAGE"):
            validate_catalogue_plan(content, self.template, self.plan)

    def test_swapped_image_binding_detected_without_changed_text_or_media(self):
        target = self.render()
        parts = load_package(target)
        tree = parse_xml(parts["word/document.xml"])
        images = tree.findall(".//a:blip", NS)
        key = "{" + NS["r"] + "}embed"
        a, b = images[0].get(key), images[1].get(key)
        images[0].set(key, b)
        images[1].set(key, a)
        parts["word/document.xml"] = xml_bytes(tree)
        write_package(target, parts)
        self.assertTrue(compare_content(self.source, target)["checks"]["media"])
        self.assertNotEqual(placement_fingerprint(self.source), placement_fingerprint(target))

    def test_header_requires_explicit_annotation_coverage(self):
        doc = Document(self.source)
        doc.sections[0].header.paragraphs[0].text = "Unsupported header"
        source = self.root / "with-header.docx"
        doc.save(source)
        with self.assertRaisesRegex(DocumentError, "ANNOTATION_COVERAGE"):
            self.extract(source, self.root / "header-package")

    def test_failed_conversion_cannot_reuse_stale_pdf(self):
        output = self.root / "pdf"
        output.mkdir()
        stale = output / "source.pdf"
        stale.write_bytes(b"previous run")
        with patch("articleflow.research.subprocess.run") as run:
            run.return_value.returncode = 0
            with self.assertRaisesRegex(DocumentError, "PDF_RENDER_FAILED"):
                render_pdf(self.source, output)
        self.assertEqual(stale.read_bytes(), b"previous run")


if __name__ == "__main__":
    unittest.main()
