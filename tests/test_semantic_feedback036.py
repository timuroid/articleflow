import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from docx.shared import Pt

from articleflow.agent023 import run
from articleflow.semantic023 import compact_conflicts,teacher_messages,student_messages


class CompactConflictFeedbackTests(unittest.TestCase):
    def test_teacher_separates_definition_act_from_repeated_mentions(self):
        system=teacher_messages({'paragraphs':[],'catalogues':[]})[0]['content']
        self.assertIn('assign a separate reusable paragraph role',system)
        self.assertIn('explicit term definition or symbol explanation',system)
        self.assertIn('merely mention the same term or symbol',system)
        self.assertIn('repeated spelling elsewhere is not evidence',system)
        self.assertIn('recognizable in unseen wording',system)

    def test_pages_cover_forty_conflicts_without_repeating_first_page(self):
        conflicts=[];groups=[]
        for index in range(40):
            kind=('character','paragraph','spacer_boundary')[index%3]
            conflicts.append({'kind':kind,'when':{'role':'body'},'variants':2,
                              'differing_properties':['size','font'],
                              'property_variants':{'a':{'size':'10'},'b':{'size':'12'}}})
            groups.append({'locations':[{'document':1,'block_id':f'b-{index:06d}',
                                         'paragraph':index+1,'excerpt':'sample '*50,
                                         'target':{'style':'a'}},
                                        {'document':1,'block_id':f'b-{index:06d}',
                                         'paragraph':index+1,'excerpt':'contrast '*50,
                                         'target':{'style':'b'}}]})
        evidence={'groups':groups};assignments={f'0:b-{i:06d}':'body' for i in range(40)}
        first=compact_conflicts(conflicts,evidence,assignments,max_bytes=5000)
        pages=[compact_conflicts(conflicts,evidence,assignments,page=i,max_bytes=5000)
               for i in range(first['conflicts_pages_total'])]
        indices=[[item['conflict_index'] for item in result['conflicts']] for result in pages]
        self.assertGreater(len(pages),1)
        self.assertTrue(set(indices[0]).isdisjoint(indices[1]))
        self.assertEqual(sorted(index for group in indices for index in group),list(range(40)))
        self.assertEqual(len({index for group in indices for index in group}),40)
        self.assertEqual({item['kind'] for item in pages[0]['conflicts']},
                         {'character','paragraph','spacer_boundary'})
        for i,result in enumerate(pages):
            self.assertEqual(result['conflicts_page'],i)
            self.assertEqual(result['conflicts_total'],40)
            self.assertEqual(result['conflicts_omitted'],40-len(result['conflicts']))
            self.assertLessEqual(len(json.dumps(result,ensure_ascii=False).encode()),5000)

    def test_student_sees_drawing_structure_in_empty_text_paragraph(self):
        content={'blocks':[{'id':'b-000001','text':'','atoms':[{'kind':'drawing','children':[]}]},
                           {'id':'b-000002','text':'','atoms':[]}]}
        geometry={'paragraph':{'b-000001':{'role':'unknown','depth':0},
                               'b-000002':{'role':'unknown','depth':0}}}
        with patch('articleflow.rules023.features',return_value=geometry):
            messages=student_messages({'figure_container':'Figure placement','empty_paragraph':'Blank spacer'},content)
        rows=json.loads(messages[1]['content'])['paragraphs']
        self.assertEqual(rows[0]['atom_kinds'],['drawing'])
        self.assertEqual(rows[1]['atom_kinds'],[])
        self.assertIn('not an empty spacer',messages[0]['content'])

    def test_large_conflict_set_keeps_contrasting_evidence_with_bounded_size(self):
        conflicts=[];groups=[]
        for index in range(46):
            kind=('character','paragraph','spacer_boundary')[index%3]
            conflicts.append({'kind':kind,'when':{'role':'body'},'variants':2,
                              'differing_properties':[f'property-{j}' for j in range(20)],
                              'property_variants':{f'style-{j}':{f'property-{k}':'v'*1000 for k in range(20)} for j in range(3)}})
            groups.append({'locations':[{'document':1,'block_id':f'b-{index:06d}','paragraph':index+1,
                                         'excerpt':'text'*500,'before':'before'*100,'after':'after'*100,
                                         'target':{'style':'style-1'},'range':[0,4]},
                                        {'document':1,'block_id':f'b-{index:06d}','paragraph':index+1,
                                         'excerpt':'other'*500,'before':'before'*100,'after':'after'*100,
                                         'target':{'style':'style-2'},'range':[0,5]}],
                           'location_complete':False})
        assignments={f'0:b-{i:06d}':'body' for i in range(46)}
        summary=compact_conflicts(conflicts,{'groups':groups},assignments)
        self.assertEqual(summary['conflicts_total'],46)
        self.assertGreater(summary['conflicts_omitted'],0)
        self.assertLessEqual(len(json.dumps(summary,ensure_ascii=False).encode()),24000)
        self.assertEqual(set(summary['conflict_kinds']),{'character','paragraph','spacer_boundary'})
        self.assertEqual({item['kind'] for item in summary['conflicts']},{'character','paragraph','spacer_boundary'})
        self.assertEqual(summary['conflicts'][0]['examples'][0]['assigned_role'],'body')
        self.assertEqual({e['target']['style'] for e in summary['conflicts'][0]['examples']},
                         {'style-1','style-2'})
        self.assertLessEqual(len(summary['conflicts'][0]['examples'][0]['excerpt']),120)
        prompt=teacher_messages({'paragraphs':[],'catalogues':[]},summary)[0]['content']
        self.assertIn('changing only role descriptions cannot resolve',prompt)
        self.assertIn('change groups/role labels',prompt)

    def test_unchanged_group_and_inline_plan_stops_before_student(self):
        with tempfile.TemporaryDirectory() as td:
            root=Path(td);source=root/'reference.docx';document=Document()
            for text,size in [('Opening',11),('Middle one',11),('Middle two',16),('Closing',11)]:
                document.add_paragraph(text).runs[0].font.size=Pt(size)
            document.save(source)
            feedback=[];calls=[]
            def caller(config,messages,directory,**kwargs):
                calls.append(config['model'])
                self.assertEqual(config['model'],'teacher')
                feedback.append(json.loads(messages[1]['content'])['feedback'])
                groups=([{'start':1,'end':4,'role':'body'}] if len(calls)==1 else
                        [{'start':i,'end':i,'role':'body'} for i in range(1,5)])
                return {'roles':{'body':'Body paragraphs '+str(len(calls))},
                        'groups':groups,
                        'questions':[],'inline_roles':{},'inline':[]},{}
            with patch('articleflow.agent023.render_pdf',return_value=root/'unused.pdf'):
                result=run([source],root/'run',{'model':'teacher','triage_conflicts':True},
                           {'model':'student'},max_iterations=5,caller=caller,adaptive=True)
            self.assertEqual(result['status'],'needs_review')
            self.assertEqual(result['code'],'SEMANTIC_PLATEAU')
            self.assertEqual(calls,['teacher','teacher'])
            self.assertEqual(len(result['iterations']),2)
            self.assertEqual(result['iterations'][1]['student_skipped'],'unchanged_training_plan')
            self.assertGreater(feedback[1]['conflicts_total'],0)
            self.assertLessEqual(len(json.dumps(feedback[1],ensure_ascii=False).encode()),24000)
            self.assertNotIn('previous_groups',feedback[1])
            self.assertEqual(feedback[1]['conflicts'][0]['examples'][0]['assigned_role'],'body')

    def test_agent_rotates_feedback_before_semantic_plateau(self):
        with tempfile.TemporaryDirectory() as td:
            root=Path(td);source=root/'reference.docx';document=Document()
            document.add_paragraph('Reference text');document.save(source)
            responses=[];conflicts=[{'kind':('character','paragraph','spacer_boundary')[i%3],
                                    'when':{'role':'body'},'variants':2,
                                    'examples':[]} for i in range(40)]
            def caller(config,messages,directory,**kwargs):
                self.assertEqual(config['model'],'teacher')
                responses.append(json.loads(messages[1]['content'])['feedback'])
                return {'roles':{'body':'Body paragraph'},
                        'groups':[{'start':1,'end':1,'role':'body'}],
                        'questions':[],'inline_roles':{},'inline':[]},{}
            with patch('articleflow.agent023.render_pdf',return_value=root/'unused.pdf'),\
                 patch('articleflow.agent023.build_package',return_value=({},conflicts)),\
                 patch('articleflow.diagnostics024.evidence',return_value={'groups':[{} for _ in conflicts]}):
                result=run([source],root/'run',{'model':'teacher','triage_conflicts':True},
                           {'model':'student'},max_iterations=5,caller=caller,adaptive=True)
            self.assertEqual(result['code'],'SEMANTIC_PLATEAU')
            self.assertEqual(len(responses),5)
            self.assertEqual([f['conflicts_page'] for f in responses[1:]],[0,1,2,3])
            self.assertEqual(sorted(item['conflict_index'] for feedback in responses[1:]
                                    for item in feedback['conflicts']),list(range(40)))
