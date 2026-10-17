import json,tempfile,unittest
from pathlib import Path
from docx import Document
from docx.shared import Cm,Pt
from docx.enum.section import WD_SECTION_START,WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from articleflow.agent023 import extract_input,build_package,apply
from articleflow.catalogue import render_catalogue
from articleflow.docx_engine import load_package,write_package,parse_xml,xml_bytes,NS
from articleflow.integrity023 import check
from articleflow.rules023 import semantic_spans
from articleflow.roundtrip022 import neutralize


def make(path,seed):
    d=Document();d.sections[0].different_first_page_header_footer=True
    d.sections[0].header.paragraphs[0].text=f'Отчёт {seed}'
    d.sections[0].first_page_header.paragraphs[0].text=f'Первая страница {seed}'
    p=d.sections[0].footer.paragraphs[0]
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE')
    run=OxmlElement('w:r');t=OxmlElement('w:t');t.text='1';run.append(t);field.append(run);p._p.append(field)
    for i in range(30):d.add_paragraph(f'Проверка {seed}, наблюдение {i}. '+'Сохранность текста и оформления проверяется отдельно. '*10)
    sec=d.add_section(WD_SECTION_START.NEW_PAGE);sec.orientation=WD_ORIENT.LANDSCAPE;sec.page_width=Cm(29.7);sec.page_height=Cm(21)
    for i in range(20):d.add_paragraph(f'Альбомная часть {seed}, измерение {i}. '+'Новые значения не меняют структуру шаблона. '*10)
    for sec in d.sections:
        for header in [sec.header,sec.first_page_header,sec.footer]:
            for p in header.paragraphs:
                p.paragraph_format.space_after=Pt(3)
                for r in p.runs:r.font.size=Pt(9)
    d.save(path)

class StoriesTests(unittest.TestCase):
    def test_oracle_retains_stories_sections_fields_and_detects_text_mutation(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);source=root/'source.docx';make(source,1)
            content=extract_input(source,root/'extracted')
            self.assertEqual(len(content['stories']),3)
            plan=json.loads((root/'extracted/oracle-plan.json').read_text())
            out=root/'result.docx';render_catalogue(root/'extracted/template',root/'extracted/content',plan,out)
            self.assertTrue(all(check(source,out).values()))
            parts=load_package(out);header=parse_xml(parts['word/header1.xml']);header.find('.//w:t',NS).text='Подмена'
            parts['word/header1.xml']=xml_bytes(header);write_package(root/'changed.docx',parts)
            self.assertFalse(check(source,root/'changed.docx')['content_exact'])

    def test_section_boundary_removal_is_detected(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);source=root/'source.docx';make(source,2)
            parts=load_package(source);document=parse_xml(parts['word/document.xml'])
            boundary=document.find('.//w:p/w:pPr/w:sectPr',NS)
            boundary.getparent().remove(boundary)
            parts['word/document.xml']=xml_bytes(document);changed=root/'changed.docx';write_package(changed,parts)
            self.assertFalse(check(source,changed)['section_boundaries_exact'])
