"""External semantic candidates stay reviewable drafts in the web intake."""
import io
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from fastapi.testclient import TestClient

from articleflow.agent023 import apply as runtime_apply
from articleflow.web017 import create_app


ROOT = Path(__file__).resolve().parents[1]
AMBIGUOUS = ROOT / "datasets" / "span-grammar-v037" / "negative-ambiguous.docx"


def reference() -> bytes:
    document = Document()
    document.add_paragraph("A simple source paragraph.")
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def candidate(count: int = 1, *, questions=True) -> bytes:
    return json.dumps({
        "roles": {"body": "Ordinary body paragraph"},
        "groups": [{"start": 1, "end": count, "role": "body"}],
        "inline_roles": {}, "inline": [],
        "questions": [{"role": "body", "question": "Is this paragraph always required?"}]
        if questions else [],
    }).encode()


class WebSemanticImportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.client = TestClient(create_app(self.root))
        self.headers = {"X-CSRF-Token": self.client.get("/api/v017/config").json()["csrf_token"]}

    def tearDown(self):
        self.client.close()
        self.temp.cleanup()

    def upload(self, document: bytes, teacher: bytes, *, grammar="true", extra_documents=()):
        documents = [document, *extra_documents]
        return self.client.post("/api/v017/templates/import-semantic", headers=self.headers,
                                data={"name": "Imported candidate", "span_grammar": grammar},
                                files=[("files", (f"reference-{index}.docx", blob,
                                                  "application/vnd.openxmlformats-officedocument.wordprocessingml.document"))
                                        for index, blob in enumerate(documents, 1)] +
                                      [("teacher_json", ("teacher.json", teacher, "application/json"))])

    def test_valid_import_is_visible_but_cannot_be_approved_or_used(self):
        response = self.upload(reference(), candidate(2), extra_documents=[reference()])
        self.assertEqual(response.status_code, 200, response.text)
        payload = response.json()
        self.assertFalse(payload["approved"])
        self.assertEqual(payload["report"]["status"], "draft_compiled")
        self.assertEqual(payload["report"]["model_calls"], 0)
        self.assertEqual(payload["report"]["questions"][0]["role"], "body")
        self.assertEqual(len(payload["report"]["references"]), 2)
        template_id, run_id = payload["id"], payload["import_run_id"]
        item = self.client.get(f"/api/v017/templates/{template_id}").json()
        self.assertFalse(item["approved"])
        self.assertEqual(item["spec"]["external_import"]["run_id"], run_id)
        self.assertEqual(self.client.get(f"/api/v017/template-runs/{run_id}").json()["report"]["schema"],
                         "external-semantic-import-037")
        self.assertTrue((self.root / "template-runs" / run_id / "import" / "template" / "semantic.json").is_file())
        self.assertEqual(self.client.post(f"/api/v017/templates/{template_id}/approve",
                                          headers=self.headers, json={"selected": {}}).status_code, 409)
        self.assertEqual(self.client.post("/api/v017/batches", headers=self.headers,
                                          data={"template_id": template_id},
                                          files={"files": ("incoming.docx", reference())}).status_code, 409)

    def test_ambiguous_import_exposes_conflicts_and_bad_json_leaves_no_template(self):
        response = self.upload(AMBIGUOUS.read_bytes(), candidate(4, questions=False))
        self.assertEqual(response.status_code, 200, response.text)
        report = response.json()["report"]
        self.assertEqual(report["status"], "needs_review")
        self.assertGreater(report["training_conflicts"], 0)
        self.assertEqual(len(report["conflicts"]), report["training_conflicts"])
        invalid = self.upload(reference(), b"This is prose, not JSON")
        self.assertEqual(invalid.status_code, 400)
        self.assertEqual(invalid.json()["detail"], "IMPORT_TEACHER_JSON_INVALID")
        self.assertEqual(len(self.client.get("/api/v017/templates").json()), 1)
        self.assertIn('id="semantic-import-form"', self.client.get("/templates").text)

    def test_single_call_preview_produces_reviewable_files_without_approval(self):
        imported = self.upload(reference(), candidate()).json()
        template_id = imported["id"]
        self.assertEqual(self.client.post(f"/api/v017/templates/{template_id}/test-import",
                                          headers=self.headers, data={"consent_external": "false"},
                                          files={"file": ("incoming.docx", reference())}).status_code, 400)
        calls = []
        def apply_stub(source, package, folder, config):
            calls.append(config)
            self.assertEqual(source.name, "source.docx")
            self.assertEqual(package.name, "template")
            folder.mkdir(parents=True)
            (folder / "result.docx").write_bytes(reference())
            pdf = folder / "render" / "result.pdf"
            pdf.parent.mkdir()
            pdf.write_bytes(b"synthetic-pdf-preview")
            return {"status": "rendered_unverified", "pdf": str(pdf),
                    "unresolved": [], "model_invocations": 1,
                    "strict_one_pass_eligible": True, "structure_moves": []}
        config = {"key_file": "private-key", "ledger": "shared-budget-ledger",
                  "student023": "small-model"}
        with patch("articleflow.web017.model_config", return_value=config), \
             patch("articleflow.agent023.apply", side_effect=apply_stub), \
             patch("articleflow.documentcheck022.audit_pdf", return_value={"status": "passed", "issues": []}):
            response = self.client.post(f"/api/v017/templates/{template_id}/test-import",
                                        headers=self.headers, data={"consent_external": "true"},
                                        files={"file": ("incoming.docx", reference())})
        self.assertEqual(response.status_code, 200, response.text)
        run_id = response.json()["id"]
        report = self.client.get(f"/api/v017/template-runs/{run_id}").json()["report"]
        self.assertEqual(report["status"], "needs_review")
        self.assertEqual(report["certification"], "single_preview_not_certification")
        self.assertTrue(report["result_available"])
        self.assertEqual(report["visual_diagnostics"]["status"], "passed")
        self.assertEqual(len(calls), 1)
        self.assertEqual(calls[0]["model"], "small-model")
        self.assertEqual(calls[0]["ledger"], "shared-budget-ledger")
        self.assertEqual(calls[0]["inline_repair_attempts"], 0)
        self.assertEqual(self.client.get(f"/api/v017/import-tests/{run_id}/files/docx").status_code, 200)
        self.assertEqual(self.client.get(f"/api/v017/import-tests/{run_id}/files/pdf").content,
                         b"synthetic-pdf-preview")
        self.assertFalse(self.client.get(f"/api/v017/templates/{template_id}").json()["approved"])

    def test_exact_reference_replay_requires_answers_then_enables_imported_batch(self):
        imported = self.upload(reference(), candidate()).json()
        template_id = imported["id"]
        self.assertEqual(self.client.post(f"/api/v017/templates/{template_id}/approve-import",
                                          headers=self.headers, json={"run_id": imported["import_run_id"],
                                                                      "answers": {"body": True}}).status_code, 409)
        calls = []
        def student_call(config, messages, folder, **kwargs):
            calls.append(config["model"])
            response = {"groups": [{"start": 1, "end": 1, "role": "body"}], "inline": []}
            kwargs["validate"](response)
            return response, {"model": "offline-student", "cost_usd": "0"}
        def offline_apply(source, package, folder, config):
            return runtime_apply(source, package, folder, config, caller=student_call)
        config = {"key_file": "private-key", "ledger": "shared-budget-ledger",
                  "student023": "offline-student"}
        with patch("articleflow.web017.model_config", return_value=config), \
             patch("articleflow.webimport037.apply", side_effect=offline_apply):
            response = self.client.post(f"/api/v017/templates/{template_id}/verify-import",
                                        headers=self.headers, json={"consent_external": True})
        self.assertEqual(response.status_code, 200, response.text)
        run_id = response.json()["id"]
        report = self.client.get(f"/api/v017/template-runs/{run_id}").json()["report"]
        self.assertEqual(report["status"], "reference_replay_passed", report)
        self.assertEqual(report["samples"][0]["role_correct"], 1)
        self.assertEqual(report["samples"][0]["inline_correct"], 1)
        self.assertTrue(report["samples"][0]["pixels"]["pixel_exact"])
        self.assertTrue(report["samples"][0]["format"]["property_exact"])
        self.assertTrue(all(row["detected"] for row in report["negative_controls"]))
        self.assertEqual(report["held_out_transfer"], "not_measured")
        self.assertEqual(calls, ["offline-student"])
        self.assertEqual(self.client.post(f"/api/v017/templates/{template_id}/approve-import",
                                          headers=self.headers, json={"run_id": run_id}).status_code, 409)
        approved = self.client.post(f"/api/v017/templates/{template_id}/approve-import",
                                    headers=self.headers, json={"run_id": run_id,
                                                                "answers": {"body": True}})
        self.assertEqual(approved.status_code, 200, approved.text)
        self.assertTrue(self.client.get(f"/api/v017/templates/{template_id}").json()["approved"])
        applied = []
        def batch_apply(source, package, folder, config):
            applied.append((package, config))
            folder.mkdir(parents=True)
            (folder / "input-original.json").write_text(json.dumps({"blocks": [
                {"id": "b-000001", "text": "A simple source paragraph."}]}))
            return {"status": "needs_review", "assignments": {"b-000001": "body"},
                    "unresolved": [], "model_invocations": 1}
        with patch("articleflow.web017.model_config", return_value=config), \
             patch("articleflow.agent023.apply", side_effect=batch_apply):
            batch = self.client.post("/api/v017/batches", headers=self.headers,
                                     data={"template_id": template_id, "use_llm": "true",
                                           "consent_external": "true"},
                                     files={"files": ("incoming.docx", reference())})
        self.assertEqual(batch.status_code, 202, batch.text)
        self.assertEqual(len(applied), 1)
        self.assertEqual(applied[0][0].parts[-2:], ("import", "template"))
        self.assertEqual(applied[0][1]["model"], "offline-student")
        self.assertEqual(applied[0][1]["inline_repair_attempts"], 0)

    def test_import_approval_rejects_package_tampering_after_replay(self):
        imported = self.upload(reference(), candidate(questions=False)).json()
        template_id = imported["id"]
        config = {"key_file": "private-key", "ledger": "shared-budget-ledger",
                  "student023": "offline-student"}
        def offline_apply(source, package, folder, settings):
            def caller(config, messages, destination, **kwargs):
                response = {"groups": [{"start": 1, "end": 1, "role": "body"}], "inline": []}
                kwargs["validate"](response)
                return response, {"model": "offline-student"}
            return runtime_apply(source, package, folder, settings, caller=caller)
        with patch("articleflow.web017.model_config", return_value=config), \
             patch("articleflow.webimport037.apply", side_effect=offline_apply):
            verified = self.client.post(f"/api/v017/templates/{template_id}/verify-import",
                                        headers=self.headers, json={"consent_external": True}).json()
        run_id = verified["id"]
        self.assertEqual(self.client.get(f"/api/v017/template-runs/{run_id}").json()["status"],
                         "reference_replay_passed")
        package = self.root / "template-runs" / imported["import_run_id"] / "import" / "template" / "resources.zip"
        package.write_bytes(package.read_bytes() + b"changed")
        rejected = self.client.post(f"/api/v017/templates/{template_id}/approve-import",
                                    headers=self.headers, json={"run_id": run_id})
        self.assertEqual(rejected.status_code, 409)
        self.assertEqual(rejected.json()["detail"], "IMPORT_EVIDENCE_CHANGED")

    def test_missing_inline_role_blocks_reference_certification(self):
        document = Document()
        paragraph = document.add_paragraph()
        paragraph.add_run("Term").bold = True
        paragraph.add_run(" follows.")
        blob = io.BytesIO()
        document.save(blob)
        teacher = json.dumps({
            "roles": {"body": "Ordinary body paragraph"},
            "groups": [{"start": 1, "end": 1, "role": "body"}],
            "inline_roles": {"term": "Defined technical term"},
            "inline": [{"paragraph": 1, "start": 0, "end": 4,
                        "role": "term", "quote": "Term"}],
            "questions": [],
        }).encode()
        imported = self.upload(blob.getvalue(), teacher).json()
        self.assertEqual(imported["report"]["training_conflicts"], 0)
        config = {"key_file": "private-key", "ledger": "shared-budget-ledger",
                  "student023": "offline-student"}
        def offline_apply(source, package, folder, settings):
            def caller(config, messages, destination, **kwargs):
                # Valid syntax but the student omitted the semantic term.
                response = {"groups": [{"start": 1, "end": 1, "role": "body"}], "inline": []}
                kwargs["validate"](response)
                return response, {"model": "offline-student"}
            return runtime_apply(source, package, folder, settings, caller=caller)
        with patch("articleflow.web017.model_config", return_value=config), \
             patch("articleflow.webimport037.apply", side_effect=offline_apply):
            started = self.client.post(f"/api/v017/templates/{imported['id']}/verify-import",
                                       headers=self.headers, json={"consent_external": True})
        self.assertEqual(started.status_code, 200, started.text)
        report = self.client.get(f"/api/v017/template-runs/{started.json()['id']}").json()["report"]
        self.assertEqual(report["status"], "needs_review")
        self.assertFalse(report["samples"][0]["inline_exact"])
        self.assertEqual(report["samples"][0]["inline_correct"], 0)
        self.assertEqual(self.client.post(f"/api/v017/templates/{imported['id']}/approve-import",
                                          headers=self.headers, json={"run_id": started.json()["id"]}).status_code, 409)


if __name__ == "__main__":
    unittest.main()
