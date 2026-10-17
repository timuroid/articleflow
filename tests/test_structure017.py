"""Conservation checks for the structural DOCX pass."""
from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from lxml import etree as ET

from articleflow.docx_engine import NS, W, load_package, parse_xml, write_package, xml_bytes
from articleflow.structure017 import compile_document, inventory, suggest_roles


class Structure017Tests(unittest.TestCase):
    def test_adjacent_bookmark_marker_and_table_survive_reordering(self):
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / "source.docx"
            result = Path(folder) / "result.docx"
            doc = Document()
            paragraph = doc.add_paragraph("Название статьи")
            doc.add_paragraph("Авторы статьи")
            doc.add_paragraph("Аннотация")
            doc.add_paragraph("Постановка проблемы: задача.")
            doc.add_paragraph("Цель: проверка.")
            doc.add_paragraph("Результаты: результат.")
            doc.add_paragraph("Практическая значимость: польза.")
            doc.add_paragraph("Ключевые слова: тест.")
            table = doc.add_table(rows=1, cols=1)
            table.cell(0, 0).text = "Данные таблицы"
            doc.save(source)
            parts = load_package(source)
            root = parse_xml(parts["word/document.xml"])
            body = root.find("w:body", NS)
            start = ET.Element(f"{{{W}}}bookmarkStart")
            start.set(f"{{{W}}}id", "23")
            start.set(f"{{{W}}}name", "Test")
            body[0].append(start)
            end = ET.Element(f"{{{W}}}bookmarkEnd")
            end.set(f"{{{W}}}id", "23")
            body.insert(1, end)
            parts["word/document.xml"] = xml_bytes(root)
            write_package(source, parts)
            before = inventory(source)
            assignments, _ = suggest_roles(before["blocks"])
            report = compile_document(source, result, assignments=assignments, approved=True, fonts_ok=True)
            self.assertEqual(report["status"], "ready")
            with ZipFile(result) as zipped:
                xml = zipped.read("word/document.xml")
            self.assertEqual(xml.count(b"bookmarkStart"), 1)
            self.assertEqual(xml.count(b"bookmarkEnd"), 1)
            self.assertIn("Данные таблицы".encode(), xml)


if __name__ == "__main__":
    unittest.main()
