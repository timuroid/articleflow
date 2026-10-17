import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from docx import Document
from lxml import etree
from articleflow.docx_engine import NS,load_package,write_package,parse_xml,xml_bytes,DocumentError
from articleflow.agent028 import run,restyle,permute
from articleflow.adapters028 import normalize
from articleflow.agent023 import extract_input

class Agent028Tests(unittest.TestCase):
    def test_adapter_only_removes_nonsemantic_markers(self):
        with tempfile.TemporaryDirectory() as td:
            base=Path(td);src=base/'original.docx';d=Document();d.add_paragraph('Title');d.add_paragraph('Text');d.save(src)
            parts=load_package(src);root=parse_xml(parts['word/document.xml']);p=root.find('.//w:p',NS);p.insert(0,etree.Element('{'+NS['w']+'}proofErr'));parts['word/document.xml']=xml_bytes(root);write_package(src,parts)
            with self.assertRaisesRegex(DocumentError,'CATALOGUE_INLINE_UNSUPPORTED'):extract_input(src,base/'fail')
            report=normalize(src,base/'normal.docx');self.assertEqual(report['removed_nodes'],1)
            extract_input(base/'normal.docx',base/'okay')
            restyle(base/'normal.docx',base/'styled.docx')
            self.assertTrue(permute(base/'styled.docx',base/'moved.docx')['applied'])
            self.assertEqual([p.text for p in Document(base/'moved.docx').paragraphs],['Text','Title'])

    def test_stress_failure_prevents_certification_and_reaches_teacher_feedback(self):
        from articleflow.agent023 import run as learn
        with tempfile.TemporaryDirectory() as td:
            base=Path(td);src=base/'sample.docx';d=Document();d.add_paragraph('Body');d.save(src);feedback=[]
            def caller(config,messages,directory,**kw):
                if config['model']=='teacher':
                    data=json.loads(messages[1]['content']);feedback.append(data['feedback'])
                    return {'roles':{'body':'Body paragraph'},'groups':[{'start':1,'end':1,'role':'body'}],'questions':[]},{}
                return {'groups':[{'start':1,'end':1,'role':'body'}]},{}
            pixels={'pixel_exact':True,'matched_page_diagnostics':[]}
            with patch('articleflow.agent023.render_pdf',return_value=base/'dummy.pdf'),patch('articleflow.agent023.compare_all',return_value=pixels):
                result=learn([src],base/'loop',{'model':'teacher'},{'model':'student'},max_iterations=2,caller=caller,adaptive=True,acceptance_probe=lambda *a:[{'variant':'restyled','status':'failed','passed':False,'code':'TEST_MISMATCH'}])
            self.assertEqual(result['status'],'needs_review');self.assertEqual(len(result['iterations']),2)
            self.assertEqual(feedback[1]['stress_failures'][0]['code'],'TEST_MISMATCH')

    def test_invalid_teacher_question_is_visible_retry_not_silent_repair(self):
        from articleflow.agent023 import run as learn
        with tempfile.TemporaryDirectory() as td:
            base=Path(td);src=base/'sample.docx';d=Document();d.add_paragraph('Body');d.save(src)
            def caller(*a,**kw):raise DocumentError('QUESTION_SCHEMA')
            with patch('articleflow.agent023.render_pdf',return_value=base/'dummy.pdf'):
                result=learn([src],base/'loop',{'model':'teacher'},{'model':'student'},max_iterations=2,caller=caller,adaptive=True)
            self.assertEqual(result['status'],'needs_review');self.assertEqual([i['code'] for i in result['iterations']],['QUESTION_SCHEMA']*2)

    def test_controller_stops_at_budget_and_does_not_certify(self):
        with tempfile.TemporaryDirectory() as td:
            base=Path(td);src=base/'sample.docx';d=Document();d.add_paragraph('Body');d.save(src)
            with patch('articleflow.agent028.render_pdf',return_value=base/'dummy.pdf'),patch('articleflow.agent028.compare_pixels',return_value={'pixel_exact':True}),patch('articleflow.agent023.render_pdf',return_value=base/'dummy.pdf'):
                result=run([src],base/'run',{'model':'teacher'},{'model':'student'},max_calls=0)
            self.assertEqual(result['code'],'AGENT_CALL_LIMIT');self.assertNotIn('certification',result)
