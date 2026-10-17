"""Regressions from real manuscripts and our authored template-transfer trial."""
import copy
import tempfile
import unittest
from pathlib import Path
from docx import Document
from articleflow.docx_engine import load_package, write_package, parse_xml, xml_bytes, NS, DocumentError
from articleflow import typography019 as t, bookmarks020 as bm
from articleflow.structure017 import inventory, suggest_roles, compile_document

class Transfer020Tests(unittest.TestCase):
    def test_bibliography_not_heading_and_restarts_at_next_section(self):
        with tempfile.TemporaryDirectory() as d:
            path=Path(d)/'in.docx';doc=Document()
            for text,style in [('Title','Title'),('References','Heading 1'),('1. A. Author. Short title.','Normal'),('2. Other citation.','Normal'),('Appendix','Heading 1')]:doc.add_paragraph(text,style)
            doc.save(path);parts=load_package(path);body=parse_xml(parts['word/document.xml']).find('w:body',NS)
            blocks=inventory(path)['blocks']
            self.assertEqual(t.context('body',body[2],blocks[2]['format_hints']),'body.references')
            self.assertEqual(t.context('body',body[3],blocks[3]['format_hints']),'body.references')
            self.assertEqual(t.context('body',body[4],blocks[4]['format_hints']),'body.heading.1')

    def test_foreign_numbering_never_imported_and_source_numbering_retained(self):
        doc=Document();p=doc.add_paragraph('citation')._p
        p.get_or_add_pPr().append(t.prop('numPr'))
        p.find('w:pPr/w:numPr',NS).append(t.prop('numId',val='42'))
        with tempfile.TemporaryDirectory() as d:
            path=Path(d)/'in.docx';doc.save(path);r=t.Resolver(load_package(path))
            p=parse_xml(xml_bytes(p));profile=t.measure(p,r,'body.references')
            self.assertNotIn('numPr',profile['pPr'])
            foreign=parse_xml(profile['pPr'].encode());n=t.prop('numPr');n.append(t.prop('numId',val='999'));foreign.append(n);profile['pPr']=t.stable_xml(foreign)
            plain=copy.deepcopy(p);plain.remove(plain.find('w:pPr',NS))
            t.apply(plain,profile,r,'body.references');self.assertIsNone(plain.find('w:pPr/w:numPr',NS))
            t.apply(p,profile,r,'body.references');self.assertEqual(p.find('w:pPr/w:numPr/w:numId',NS).get(t.tag('val')),'42')
            self.assertEqual(len(p.findall('w:pPr/w:numPr',NS)),1)

    def test_first_table_row_is_measured_and_applied_separately(self):
        with tempfile.TemporaryDirectory() as d:
            path=Path(d)/'in.docx';doc=Document();table=doc.add_table(rows=2,cols=2)
            for i,row in enumerate(table.rows):
                for c in row.cells:c.text='Header' if i==0 else '42';c.paragraphs[0].runs[0].bold=i==0
            doc.save(path);r=t.Resolver(load_package(path));profile=t.table_profile(parse_xml(xml_bytes(table._tbl)),r)
            self.assertIn('first_row',profile)
            for row in table.rows:
                for c in row.cells:c.paragraphs[0].runs[0].bold=False
            node=parse_xml(xml_bytes(table._tbl));t.apply_table(node,profile,r)
            rows=node.findall('w:tr',NS)
            self.assertEqual(t.value(rows[0].find('.//w:rPr',NS),'b','1'),'1')
            self.assertEqual(t.value(rows[1].find('.//w:rPr',NS),'b'),'0')

    def test_interparagraph_bookmark_import_and_range_gate(self):
        with tempfile.TemporaryDirectory() as d:
            root=Path(d);path=root/'in.docx';doc=Document()
            for text in ['Title','Author','Body one','Body two']:doc.add_paragraph(text)
            doc.save(path);parts=load_package(path);xml=parse_xml(parts['word/document.xml']);body=xml.find('w:body',NS)
            body.insert(2,parse_xml(b'<w:bookmarkStart xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:id="7" w:name="section"/>'));body.insert(5,t.prop('bookmarkEnd',id='7'))
            parts['word/document.xml']=xml_bytes(xml);write_package(path,parts)
            before=bm.signature(body);blocks=inventory(path)['blocks'];roles={b['id']:'body' for b in blocks}
            compile_document(path,root/'out.docx',assignments=roles,required=[])
            out=parse_xml(load_package(root/'out.docx')['word/document.xml']).find('w:body',NS)
            self.assertEqual(before,bm.signature(out));self.assertEqual(len(out.findall('.//w:bookmarkStart',NS)),1)
            roles[blocks[-1]['id']]='title'
            with self.assertRaisesRegex(DocumentError,'BOOKMARK_RANGE_CHANGED|BOOKMARK_INVALID'):
                compile_document(path,root/'bad.docx',assignments=roles,required=[])
            self.assertFalse((root/'bad.docx').exists())

    def test_dangling_bookmark_is_rejected(self):
        node=t.prop('bookmarkStart',id='9')
        with self.assertRaisesRegex(DocumentError,'BOOKMARK_INVALID'):bm.signature([node])

    def test_bookmarked_drawing_ignores_unused_namespace_but_protects_payload(self):
        template='<w:p xmlns:w="'+NS['w']+'" {ns}><w:bookmarkStart w:id="1" w:name="image"/><w:r><w:drawing><x:graphic xmlns:x="urn:drawing" size="{size}"/></w:drawing></w:r><w:bookmarkEnd w:id="1"/></w:p>'
        a=parse_xml(template.format(ns='xmlns:unused="urn:unused"',size='42').encode())
        b=parse_xml(template.format(ns='',size='42').encode())
        changed=parse_xml(template.format(ns='',size='43').encode())
        self.assertEqual(bm.signature([a]),bm.signature([b]))
        self.assertNotEqual(bm.signature([a]),bm.signature([changed]))

class Label020Tests(unittest.TestCase):
    def test_label_bold_word_with_plain_punctuation(self):
        with tempfile.TemporaryDirectory() as d:
            path=Path(d)/'in.docx';doc=Document();p=doc.add_paragraph();p.add_run('Цель').bold=True;p.add_run('. Проверить перенос.').bold=False;doc.save(path)
            parts=load_package(path);r=t.Resolver(parts);node=parse_xml(parts['word/document.xml']).find('w:body/w:p',NS)
            profile=t.measure(node,r,'abstract_goal');self.assertEqual(profile['label_scope'],'words')
            target=parse_xml(f'<w:p xmlns:w="{NS["w"]}"><w:r><w:t>Цель. Новый текст.</w:t></w:r></w:p>'.encode())
            t.apply(target,profile,r,'abstract_goal');runs=target.findall('w:r',NS)
            self.assertEqual(runs[0].find('w:t',NS).text,'Цель');self.assertEqual(t.value(runs[0].find('w:rPr',NS),'b','1'),'1')
            self.assertEqual(t.value(runs[1].find('w:rPr',NS),'b'),'0')

class Preservation020Tests(unittest.TestCase):
    def test_bookmark_cannot_span_main_and_editorial_outputs(self):
        with tempfile.TemporaryDirectory() as d:
            root=Path(d);source=root/'source.docx';doc=Document()
            doc.add_paragraph('Title');doc.add_paragraph('Editorial');doc.save(source)
            parts=load_package(source);xml=parse_xml(parts['word/document.xml']);ps=xml.findall('w:body/w:p',NS)
            start=t.prop('bookmarkStart',id='1');start.set(t.tag('name'),'range');ps[0].insert(0,start);ps[1].append(t.prop('bookmarkEnd',id='1'))
            parts['word/document.xml']=xml_bytes(xml);write_package(source,parts)
            with self.assertRaisesRegex(DocumentError,'BOOKMARK_INVALID'):
                compile_document(source,root/'out.docx',assignments={'n-000001':'title','n-000002':'editorial'},required=[],editorial_output=root/'notes.docx')
            self.assertFalse((root/'out.docx').exists());self.assertFalse((root/'notes.docx').exists())

    def test_source_default_spacing_does_not_leak_into_target(self):
        with tempfile.TemporaryDirectory() as d:
            source=Path(d)/'in.docx';doc=Document();doc.add_paragraph('Text');doc.save(source)
            parts=load_package(source)
            profile={'pPr':f'<w:pPr xmlns:w="{NS["w"]}"><w:spacing w:after="0"/></w:pPr>',
                     'rPr':f'<w:rPr xmlns:w="{NS["w"]}"><w:rFonts w:ascii="Arial" w:hAnsi="Arial"/></w:rPr>'}
            t.apply_base(parts,profile)
            p,_=t.Resolver(parts).resolve(t.prop('p'))
            self.assertIsNone(p.find('w:spacing',NS).get(t.tag('line')))
            self.assertEqual(p.find('w:spacing',NS).get(t.tag('after')),'0')

    def test_editorial_output_must_exist_and_be_distinct(self):
        with tempfile.TemporaryDirectory() as d:
            root=Path(d);source=root/'source.docx';out=root/'out.docx';notes=root/'notes.docx'
            doc=Document();doc.add_paragraph('Scientific conclusion');doc.save(source)
            roles={inventory(source)['blocks'][0]['id']:'editorial'}
            for target,error in [(None,'EDITORIAL_OUTPUT_REQUIRED'),(out,'EDITORIAL_OUTPUT_CONFLICT'),(source,'EDITORIAL_OUTPUT_CONFLICT')]:
                with self.assertRaisesRegex(DocumentError,error):
                    compile_document(source,out,assignments=roles,required=[],editorial_output=target)
                self.assertFalse(out.exists())
            compile_document(source,out,assignments=roles,required=[],editorial_output=notes)
            self.assertEqual(Document(notes).paragraphs[0].text,'Scientific conclusion')

    def test_table_header_fonts_included_in_audit(self):
        from articleflow.fonts019 import required_fonts
        self.assertEqual(required_fonts({'body.table':{'first_row':{'rPr':f'<w:rPr xmlns:w="{NS["w"]}"><w:rFonts w:ascii="Header Font"/></w:rPr>'}}}),('Header Font',))

if __name__=='__main__':unittest.main()
