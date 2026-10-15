import copy
import json
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.shared import Pt

from articleflow.catalogue import extract_catalogue
from articleflow.context_rules import fit_rules
from articleflow.docx_engine import DocumentError, compare_content, digest, load_package, parse_xml, write_package, xml_bytes, NS
from articleflow.source_aware import ROLE_VERSION, choose_source_style, format_source


class SourceAwareTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)
        self.source = self.root / "source.docx"
        doc = Document()
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(7)
        r = p.add_run("Синтетический заголовок")
        r.font.name, r.font.size, r.bold = "Arial", Pt(14), True
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run("Основной текст.")
        r.font.name, r.font.size = "Arial", Pt(11)
        doc.save(self.source)
        self.roles = {"b-000001": "title", "b-000002": "body"}
        self.package = self.root / "template"
        extract_catalogue(self.source, self.roles, {"roles": {"title": "Заголовок", "body": "Основной текст"}}, self.package)
        self.template = json.loads((self.package / "template/rules.json").read_text())
        content = json.loads((self.package / "content/content.json").read_text())
        oracle = json.loads((self.package / "oracle-plan.json").read_text())
        self.pack, _ = fit_rules(self.template, content, oracle, oracle["template_sha256"])
        (self.package / "oracle-plan.json").unlink()

    def role_plan(self, source):
        return {"schema_version": ROLE_VERSION, "source_sha256": digest(source.read_bytes()),
                "blocks": [{"id": key, "role": role} for key, role in self.roles.items()]}

    def test_known_document_rebuild_needs_no_target_oracle(self):
        out = self.root / "same"
        report = format_source(self.package / "template", self.source, self.role_plan(self.source), out, self.pack)
        self.assertEqual(report["status"], "assembled_draft")
        self.assertFalse(report["target_oracle_used"])
        self.assertFalse(report["strict_I1_eligible"])
        self.assertTrue(report["source_formatting_used"])
        self.assertTrue(compare_content(self.source, out / "candidate.docx")["inline_exact"])
        self.assertFalse((out / "input/oracle-plan.json").exists())
        self.assertTrue((out / "input/source-observations.json").exists())

    def test_new_text_wrong_font_size_and_spacing_are_converted(self):
        changed = self.root / "new.docx"
        doc = Document(self.source)
        doc.paragraphs[0].runs[0].text = "Совсем другая статья"
        doc.paragraphs[0].runs[0].font.name = "Courier New"
        doc.paragraphs[0].runs[0].font.size = Pt(24)
        doc.paragraphs[0].paragraph_format.space_after = Pt(40)
        doc.paragraphs[1].runs[0].text = "Новое содержание и число 42."
        doc.save(changed)
        out = self.root / "new-result"
        report = format_source(self.package / "template", changed, self.role_plan(changed), out, self.pack)
        self.assertEqual(report["status"], "assembled_draft")
        result = Document(out / "candidate.docx")
        self.assertEqual(result.paragraphs[0].text, "Совсем другая статья")
        self.assertEqual(result.paragraphs[1].text, "Новое содержание и число 42.")
        self.assertEqual(result.paragraphs[0].runs[0].font.name, "Arial")
        self.assertEqual(result.paragraphs[0].runs[0].font.size.pt, 14)
        self.assertEqual(result.paragraphs[0].paragraph_format.space_after.pt, 7)
        self.assertTrue(compare_content(changed, out / "candidate.docx")["content_exact"])
        self.assertTrue(any(d["reason"] == "unique_typography_mapping" for d in report["decisions"]))

    def test_ambiguous_typography_does_not_choose_arbitrarily(self):
        template = copy.deepcopy(self.template)
        key = template["roles"]["title"]["character_styles"][0]
        raw = template["styles"]["character"][key]
        template["styles"]["character"]["second"] = [x.replace('w:val="28"', 'w:val="30"') for x in raw]
        source = {"character": {"unknown": [x.replace('Arial', 'Courier New') for x in raw]}}
        with self.assertRaisesRegex(DocumentError, "SOURCE_VARIANT_AMBIGUOUS"):
            choose_source_style("character", "unknown", source, template, [key, "second"])

    def test_resource_mismatch_blocks_output(self):
        parts = load_package(self.source)
        root = parse_xml(parts["word/styles.xml"])
        root.set("test-resource-version", "different")
        parts["word/styles.xml"] = xml_bytes(root)
        changed = self.root / "resources.docx"
        write_package(changed, parts)
        out = self.root / "blocked"
        report = format_source(self.package / "template", changed, self.role_plan(changed), out, self.pack)
        self.assertEqual(report["status"], "needs_review")
        self.assertEqual(report["unresolved"][0]["reason"], "SOURCE_RESOURCES_INCOMPATIBLE")
        self.assertFalse((out / "candidate.docx").exists())

    def test_source_hash_duplicate_ids_unknown_roles_rejected(self):
        plan = self.role_plan(self.source)
        bad = copy.deepcopy(plan); bad["source_sha256"] = "wrong"
        duplicate = copy.deepcopy(plan); duplicate["blocks"].append(duplicate["blocks"][0])
        unknown = copy.deepcopy(plan); unknown["blocks"][0]["role"] = "invented"
        for index, candidate in enumerate([bad, duplicate, unknown]):
            with self.subTest(index=index), self.assertRaises(DocumentError):
                format_source(self.package / "template", self.source, candidate, self.root / str(index), self.pack)
