"""Web 017 integration checks with disposable, synthetic DOCX files."""
from __future__ import annotations

import io
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

from docx import Document
from fastapi.testclient import TestClient

from articleflow.web017 import create_app


def article(*, authors: bool = True) -> bytes:
    doc = Document()
    doc.add_paragraph("Исследование синтетического корпуса")
    if authors:
        doc.add_paragraph("Иванов И. И.")
    doc.add_paragraph("Аннотация")
    doc.add_paragraph("Постановка проблемы: требуется проверка структуры.")
    doc.add_paragraph("Цель: проверить перестановку блоков.")
    doc.add_paragraph("Результаты: получен проверяемый файл.")
    doc.add_paragraph("Практическая значимость: уменьшается ручная работа.")
    doc.add_paragraph("Ключевые слова: структура, DOCX.")
    doc.add_paragraph("Основной текст остаётся неизменным.")
    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()


class Web017Tests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.client = TestClient(create_app(Path(self.temp.name)))
        self.csrf = self.client.get("/api/v017/config").json()["csrf_token"]
        self.headers = {"X-CSRF-Token": self.csrf}

    def tearDown(self):
        self.client.close()
        self.temp.cleanup()

    def template(self):
        response = self.client.post(
            "/api/v017/templates", headers=self.headers,
            data={"name": "Синтетический образец"},
            files=[("files", ("sample.docx", article(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"))],
        )
        self.assertEqual(response.status_code, 200, response.text)
        template_id = response.json()["id"]
        approved = self.client.post(f"/api/v017/templates/{template_id}/approve", headers=self.headers, json={"selected": {}})
        self.assertEqual(approved.status_code, 200, approved.text)
        return template_id

    def test_csrf_and_batch_review_zip(self):
        forbidden = self.client.post("/api/v017/templates", data={"name": "bad"},
                                     files=[("files", ("sample.docx", article(), "application/octet-stream"))])
        self.assertEqual(forbidden.status_code, 403)
        template_id = self.template()
        with patch("articleflow.web017.font_audit", return_value={"families": {"Arial": True, "Times New Roman": True}, "exact": True}):
            batch = self.client.post(
                "/api/v017/batches", headers=self.headers, data={"template_id": template_id},
                files=[("files", ("input.docx", article(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"))],
            )
            self.assertEqual(batch.status_code, 202, batch.text)
            batch_id = batch.json()["id"]
            job_id = batch.json()["jobs"][0]
            job = self.client.get(f"/api/v017/jobs/{job_id}").json()
            self.assertEqual(job["status"], "draft")
            self.assertEqual(self.client.get(f"/api/v017/jobs/{job_id}/files/final").status_code, 409)
            assignments = job["report"]["assignments"]
            self.assertEqual(set(assignments), {block["id"] for block in job["blocks"]})
            decided = self.client.patch(f"/api/v017/jobs/{job_id}/decisions", headers=self.headers,
                                        json={"expected_revision": 0, "assignments": assignments})
            self.assertEqual(decided.status_code, 200, decided.text)
            self.assertEqual(self.client.patch(f"/api/v017/jobs/{job_id}/decisions", headers=self.headers,
                                               json={"expected_revision": 0, "assignments": assignments}).status_code, 409)
            job = self.client.get(f"/api/v017/jobs/{job_id}").json()
            self.assertEqual(job["status"], "ready", job["report"])
            self.assertEqual(self.client.get(f"/api/v017/jobs/{job_id}/files/final").status_code, 200)
            exported = self.client.post(f"/api/v017/batches/{batch_id}/exports", headers=self.headers)
            self.assertEqual(exported.status_code, 200, exported.text)
            blob = self.client.get(exported.json()["url"]).content
            with ZipFile(io.BytesIO(blob)) as archive:
                names = archive.namelist()
                self.assertIn(f"results/{job_id}.docx", names)
                self.assertEqual(json.loads(archive.read("manifest.json"))["jobs"][0]["status"], "ready")

    def test_missing_required_data_stays_draft(self):
        template_id = self.template()
        with patch("articleflow.web017.font_audit", return_value={"families": {"Arial": True, "Times New Roman": True}, "exact": True}):
            batch = self.client.post("/api/v017/batches", headers=self.headers, data={"template_id": template_id},
                                     files=[("files", ("input.docx", article(authors=False), "application/octet-stream"))])
            self.assertEqual(batch.status_code, 202, batch.text)
            job_id = batch.json()["jobs"][0]
            job = self.client.get(f"/api/v017/jobs/{job_id}").json()
            assignments = job["report"]["assignments"]
            # Correct the deliberately weak title/author heuristic before review.
            for block in job["blocks"]:
                if block["text"].startswith("Аннотация"):
                    assignments[block["id"]] = "abstract_heading"
            self.client.patch(f"/api/v017/jobs/{job_id}/decisions", headers=self.headers,
                              json={"expected_revision": 0, "assignments": assignments})
            revised = self.client.get(f"/api/v017/jobs/{job_id}").json()
            self.assertEqual(revised["status"], "draft")
            self.assertIn("authors", revised["report"]["missing"])
            draft = self.client.get(f"/api/v017/jobs/{job_id}/files/draft").content
            with ZipFile(io.BytesIO(draft)) as docx:
                self.assertIn("[ЗАПОЛНИТЬ: authors]", docx.read("word/document.xml").decode())


if __name__ == "__main__":
    unittest.main()
