import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from docx import Document
from PIL import Image
from articleflow.agent023 import build_package, extract_input, apply
from articleflow.order024 import fit, arrange, expected_source
from articleflow.rules023 import semantic_spans
from articleflow.integrity023 import check
from articleflow.docx_engine import DocumentError


class OrderingTests(unittest.TestCase):
    def test_apply_repairs_order_and_preserves_payload(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            rows = [('title', 'Название'), ('abstract', 'Аннотация'), ('keywords', 'Ключевые слова')]
            def make(name, values):
                d = Document()
                for role, text in values:
                    p = d.add_paragraph(text)
                    if role == 'title': p.runs[0].bold = True
                d.save(root / name)
            make('reference.docx', rows)
            content = extract_input(root/'reference.docx', root/'training')
            labels = {b['id']: r for b, (r, _) in zip(content['blocks'], rows)}
            package, conflicts = build_package([{'content': content, 'extracted': root/'training',
                'semantics': semantic_spans(root/'reference.docx')}], [labels], {r: r for r, _ in rows}, root/'template')
            self.assertFalse(conflicts)
            # Independent text and deliberately incorrect order.
            incoming = [('keywords', 'Термины: другое'), ('title', 'Новая работа'), ('abstract', 'Новый метод')]
            make('incoming.docx', incoming)
            def caller(*args, **kwargs):
                return {'groups': [{'start': i, 'end': i, 'role': r} for i, (r, _) in enumerate(incoming, 1)]}, {'api_attempts': 0}
            with patch('articleflow.agent023.render_pdf', return_value=root/'mock.pdf'):
                result = apply(root/'incoming.docx', root/'template', root/'application', {}, caller=caller)
            self.assertEqual(result['status'], 'rendered_unverified', result)
            self.assertTrue(result['content_exact'])
            self.assertEqual(len(result['structure_moves']), 3)
            output = root/'application/result.docx'
            self.assertEqual([p.text for p in Document(output).paragraphs], ['Новая работа', 'Новый метод', 'Термины: другое'])
            self.assertFalse(check(root/'incoming.docx', output)['content_exact'])
            d = Document(output); d.paragraphs[0].text = 'Потеря содержания'; d.save(root/'tampered.docx')
            self.assertFalse(check(root/'application/expected-content-order.docx', root/'tampered.docx')['content_exact'])
            self.assertNotIn('Название', str(package))

    def test_conflicting_references_are_not_majority_voted(self):
        content = {'blocks': [{'id': str(i), 'text': str(i), 'atoms': []} for i in range(2)],
                   'structure': [{'kind': 'paragraph', 'id': str(i)} for i in range(2)]}
        contract, conflicts = fit([(content, {'0':'title','1':'abstract'}), (content, {'0':'abstract','1':'title'})])
        self.assertTrue(conflicts); self.assertEqual(contract['status'], 'conflicting')

    def test_repeated_section_roles_have_no_global_sort_rule(self):
        content = {'blocks': [{'id': str(i), 'text': 'text', 'atoms': []} for i in range(4)],
                   'structure': [{'kind': 'paragraph', 'id': str(i)} for i in range(4)]}
        contract, conflicts = fit([(content, dict(zip(['0','1','2','3'], ['heading','body','heading','body'])))])
        self.assertFalse(conflicts); self.assertEqual(contract['edges'], [])

    def test_single_object_caption_pair_can_repeat_without_global_order_violation(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            image = root/'figure.png'
            Image.new('RGB', (4, 4), 'black').save(image)
            def document(path, roles):
                doc = Document()
                for role in roles:
                    if role == 'object_anchor':
                        doc.add_picture(str(image))
                    else:
                        doc.add_paragraph(role)
                doc.save(path)
                content = extract_input(path, root/path.stem)
                return content, {block['id']: role for block, role in zip(content['blocks'], roles)}

            front = ['front_a', 'front_b', 'front_c']
            training = front + ['prose', 'object_anchor', 'following_label', 'prose']
            content, labels = document(root/'reference.docx', training)
            contract, conflicts = fit([(content, labels)])
            self.assertFalse(conflicts)
            edges = {(edge['before'], edge['after']) for edge in contract['edges']}
            self.assertNotIn(('object_anchor', 'following_label'), edges)
            self.assertFalse(any('object_anchor' in edge or 'following_label' in edge for edge in edges))
            self.assertTrue({('front_a', 'front_b'), ('front_b', 'front_c')} <= edges)

            heldout = ['front_c', 'front_a', 'front_b', 'prose',
                       'object_anchor', 'following_label', 'prose',
                       'object_anchor', 'following_label']
            content, labels = document(root/'heldout.docx', heldout)
            arranged, permutation, moves, issues = arrange(content, labels, contract, root/'heldout.docx')
            self.assertFalse(issues)
            self.assertEqual([labels[node['id']] for node in arranged['structure']],
                             ['front_a', 'front_b', 'front_c', 'prose',
                              'object_anchor', 'following_label', 'prose',
                              'object_anchor', 'following_label'])
            self.assertEqual(sorted(permutation), list(range(len(heldout))))
            self.assertTrue(moves)

    def test_cover_object_does_not_erase_front_matter_order(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            image = root/'logo.png'
            Image.new('RGB', (4, 4), 'black').save(image)
            source = root/'cover.docx'
            doc = Document()
            doc.add_picture(str(image))
            for value in ('title', 'authors', 'abstract', 'body first', 'body second'):
                doc.add_paragraph(value)
            doc.save(source)
            content = extract_input(source, root/'input')
            roles = ['cover_logo', 'front_a', 'front_b', 'front_c', 'prose', 'prose']
            labels = {block['id']: role for block, role in zip(content['blocks'], roles)}
            contract, conflicts = fit([(content, labels)])
            self.assertFalse(conflicts)
            edges = {(edge['before'], edge['after']) for edge in contract['edges']}
            self.assertTrue({('front_a', 'front_b'), ('front_b', 'front_c')} <= edges)

    def test_table_boundary_blocks_unsafe_move(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td); source = root/'source.docx'
            d = Document(); d.add_paragraph('Аннотация'); d.add_table(rows=1, cols=1).cell(0,0).text='42'; d.add_paragraph('Название'); d.save(source)
            content = extract_input(source, root/'input')
            labels = dict(zip([b['id'] for b in content['blocks']], ['abstract','cell','title']))
            contract = {'schema':'role-order-024','status':'observed','edges':[{'before':'title','after':'abstract'}]}
            _, permutation, moves, issues = arrange(content, labels, contract, source)
            self.assertTrue(issues); self.assertFalse(moves); self.assertEqual(permutation, [0,1,2])
            with self.assertRaisesRegex(DocumentError, 'ORDER_NOT_PERMUTATION'):
                expected_source(source, [0,0,2], root/'bad.docx')


if __name__ == '__main__': unittest.main()
