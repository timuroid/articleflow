import copy
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import pymupdf
from docx import Document
from articleflow import templateagent022 as agent
from articleflow.documentcheck022 import audit_pdf,repair_keep_links
from articleflow.docx_engine import NS,load_package,parse_xml,paragraph_text,digest
from articleflow.structure017 import inspect_templates


class AllPages022Tests(unittest.TestCase):
    def pdf(self,path,last='Last page',size=(595,842)):
        d=pymupdf.open()
        for text in ('First page',last):
            p=d.new_page(width=size[0],height=size[1]);p.insert_text((70,70),text)
        d.save(path);d.close()

    def test_detects_only_last_page_change_and_reports_region(self):
        with tempfile.TemporaryDirectory() as d:
            d=Path(d);self.pdf(d/'a.pdf');self.pdf(d/'b.pdf','Different last page')
            r=agent.compare_all(d/'a.pdf',d/'b.pdf',d/'diff')
            self.assertFalse(r['pixel_exact']);self.assertTrue(r['matched_page_diagnostics'][0]['pixel_exact'])
            self.assertIsNotNone(r['matched_page_diagnostics'][1]['difference_bbox_px'])
            self.assertTrue((d/'diff/diff-2.png').is_file())

    def test_detects_page_count_and_size_mismatch(self):
        with tempfile.TemporaryDirectory() as d:
            d=Path(d);self.pdf(d/'a.pdf');self.pdf(d/'b.pdf',size=(500,842))
            self.assertFalse(agent.compare_all(d/'a.pdf',d/'b.pdf',d/'diff')['page_sizes_equal'])
            doc=pymupdf.open(d/'a.pdf');doc.delete_page(1);doc.save(d/'c.pdf');doc.close()
            self.assertFalse(agent.compare_all(d/'a.pdf',d/'c.pdf',d/'diff2')['pages_equal'])

    def test_all_pages_checked_for_overlapping_text(self):
        with tempfile.TemporaryDirectory() as d:
            p=Path(d)/'a.pdf';doc=pymupdf.open();doc.new_page().insert_text((72,72),'Clear text')
            page=doc.new_page();page.insert_text((72,72),'First overlapping text',fontsize=20);page.insert_text((72,76),'Second overlapping text',fontsize=20);doc.save(p);doc.close()
            r=audit_pdf(p);self.assertEqual(len(r['pages']),2)
            self.assertTrue(any(i['code']=='TEXT_OVERLAP_REVIEW' and i['page']==2 for i in r['issues']))

    def test_blank_page_is_review_not_pass(self):
        with tempfile.TemporaryDirectory() as d:
            p=Path(d)/'a.pdf';doc=pymupdf.open();doc.new_page();doc.save(p);doc.close()
            self.assertEqual(audit_pdf(p)['issues'][0]['code'],'EMPTY_PAGE_REVIEW')

    def test_safe_orphan_repair_preserves_text_and_only_changes_layout(self):
        with tempfile.TemporaryDirectory() as d:
            d=Path(d);doc=Document();doc.add_paragraph('Введение');doc.add_paragraph('Scientific content 42.');doc.save(d/'a.docx')
            changes=repair_keep_links(d/'a.docx',d/'b.docx',{'issues':[{'code':'ORPHAN_LABEL','page':1,'text_sha256':digest('введение'.encode())}]})
            self.assertEqual(len(changes),1)
            self.assertEqual([p.text for p in Document(d/'a.docx').paragraphs],[p.text for p in Document(d/'b.docx').paragraphs])
            self.assertIsNotNone(parse_xml(load_package(d/'b.docx')['word/document.xml']).find('w:body/w:p/w:pPr/w:keepNext',NS))

    def test_agent_attempt_limit_is_not_success(self):
        with tempfile.TemporaryDirectory() as d:
            d=Path(d);doc=Document();doc.add_paragraph('Sample title');doc.save(d/'a.docx');spec=inspect_templates([d/'a.docx'])
            self.pdf(d/'a.pdf')
            with patch.object(agent,'render_pdf',return_value=d/'a.pdf'):
                r=agent.run([d/'a.docx'],spec,d/'run',max_attempts=1)
            self.assertEqual(r['status'],'needs_review');self.assertEqual(r['samples'][0]['stop_reason'],'attempt_limit')
            self.assertEqual(r['transfer_validation'],'not_measured')

    def test_renderer_error_is_unavailable_and_saves_report(self):
        with tempfile.TemporaryDirectory() as d:
            d=Path(d);doc=Document();doc.add_paragraph('Title');doc.save(d/'a.docx');s=inspect_templates([d/'a.docx'])
            with patch.object(agent,'render_pdf',side_effect=FileNotFoundError('private detail')):
                r=agent.run([d/'a.docx'],s,d/'run')
            self.assertEqual(r['status'],'unavailable');self.assertNotIn('private detail',json.dumps(r));self.assertTrue((d/'run/result.json').is_file())

    def test_failed_catalogue_never_overridden_by_exact_pixels(self):
        with tempfile.TemporaryDirectory() as d:
            d=Path(d);doc=Document();doc.add_paragraph('Title');doc.save(d/'a.docx');s=inspect_templates([d/'a.docx']);self.pdf(d/'a.pdf')
            with patch.object(agent,'render_pdf',return_value=d/'a.pdf'),patch.object(agent,'content_fingerprint',side_effect=[{'text':'a'},{'text':'b'}]),patch.object(agent,'_negative_controls',return_value=[{'detected':True}]):
                r=agent.run([d/'a.docx'],s,d/'run')
            self.assertEqual(r['status'],'needs_review');self.assertFalse(next(a for a in r['samples'][0]['attempts'] if a['strategy']=='complete_property_catalogue')['content_exact'])


class Web022Tests(unittest.TestCase):
    def test_background_run_artifacts_and_no_content_in_runtime_zip(self):
        from fastapi.testclient import TestClient
        from articleflow.web017 import create_app
        from zipfile import ZipFile
        from io import BytesIO
        with tempfile.TemporaryDirectory() as d:
            d=Path(d);doc=Document();doc.add_paragraph('PRIVATE SCIENTIFIC TEXT');doc.save(d/'a.docx')
            def fake_run(paths,spec,folder,progress):
                folder.mkdir(parents=True);p=folder/'sample-1/catalogue/template';p.mkdir(parents=True)
                (p/'rules.json').write_text('{}');(p/'resources.zip').write_bytes(b'resources')
                r={'status':'passed','samples':[{'runtime_catalogue':{'rules_sha256':'test'}}]};progress(r);return r
            with TestClient(create_app(d/'data')) as c,patch.object(agent,'run',side_effect=fake_run):
                h={'X-CSRF-Token':c.get('/api/v017/config').json()['csrf_token']}
                tid=c.post('/api/v017/templates',headers=h,data={'name':'sample'},files={'files':('sample.docx',(d/'a.docx').read_bytes())}).json()['id']
                self.assertEqual(c.post(f'/api/v017/templates/{tid}/refine').status_code,403)
                rid=c.post(f'/api/v017/templates/{tid}/refine',headers=h).json()['id']
                self.assertEqual(c.get('/api/v017/template-runs/'+rid).json()['status'],'passed')
                response=c.get(f'/api/v017/template-runs/{rid}/catalogue')
                with ZipFile(BytesIO(response.content)) as z:
                    self.assertEqual(set(z.namelist()),{'sample-1/rules.json','sample-1/resources.zip','manifest.json'})
                    self.assertFalse(any(b'PRIVATE SCIENTIFIC TEXT' in z.read(n) for n in z.namelist()))
                self.assertEqual(c.get(f'/api/v017/template-runs/{rid}/preview/0').status_code,404)

if __name__=='__main__':unittest.main()

class RuntimeAcceptance022Tests(unittest.TestCase):
    def test_exact_catalogue_cannot_mask_failed_production_formatter(self):
        with tempfile.TemporaryDirectory() as d:
            d=Path(d);doc=Document();doc.add_paragraph('Title');doc.save(d/'a.docx');s=inspect_templates([d/'a.docx'])
            pdf=pymupdf.open();pdf.new_page().insert_text((70,70),'Title');pdf.save(d/'a.pdf');pdf.close()
            with patch.object(agent,'render_pdf',return_value=d/'a.pdf'),patch.object(agent,'_negative_controls',return_value=[{'detected':True}]),patch('articleflow.roundtrip022.verify',return_value={'strategy':'production_roundtrip','status':'failed','rows':[]}):
                r=agent.run([d/'a.docx'],s,d/'run')
            self.assertEqual(r['samples'][0]['attempts'][1]['status'],'passed')
            self.assertEqual(r['status'],'needs_review')
            self.assertEqual(r['samples'][0]['stop_reason'],'production_template_not_verified')

    def test_neutral_input_removes_table_gaps_and_geometry(self):
        from articleflow.roundtrip022 import neutralize
        from docx.shared import Pt
        with tempfile.TemporaryDirectory() as d:
            d=Path(d);doc=Document();doc.add_paragraph('Title');table=doc.add_table(rows=2,cols=2)
            for row in table.rows:
                for cell in row.cells:cell.text='Cell'
            doc.add_paragraph();p=doc.add_paragraph('Scientific number 42');p.paragraph_format.space_before=Pt(24);p.runs[0].font.size=Pt(30);doc.save(d/'a.docx')
            mapping=neutralize(d/'a.docx',d/'b.docx')
            self.assertNotIn(2,mapping);self.assertEqual(mapping[3],2)
            out=Document(d/'b.docx');self.assertEqual([p.text for p in out.paragraphs],['Title','Scientific number 42'])
            self.assertIsNone(out.paragraphs[-1].paragraph_format.space_before)
            self.assertIsNone(out.paragraphs[-1].runs[0].font.size)

    def test_table_boundary_is_extracted_applied_once_and_not_invented(self):
        from articleflow import layout021,typography019 as t
        from articleflow.structure017 import compile_document,inventory
        with tempfile.TemporaryDirectory() as d:
            d=Path(d);doc=Document();table=doc.add_table(rows=2,cols=2)
            for row in table.rows:
                for cell in row.cells:cell.text='value'
            doc.add_paragraph();doc.add_paragraph('After first table');doc.add_table(rows=1,cols=1).cell(0,0).text='Second';doc.add_paragraph('After second table');doc.save(d/'ref.docx')
            spec=inspect_templates([d/'ref.docx']);obs=[o for o in spec['samples'][0]['observations'] if o['role']=='body.table']
            profiles=[spec['styles']['body.table'][o['style_key']] for o in obs]
            self.assertEqual([len(p['trailing_spacing']) for p in profiles],[1,0])
            parts=load_package(d/'ref.docx');resolver=t.Resolver(parts);root=parse_xml(parts['word/document.xml']);body=root.find('w:body',NS)
            ordered=[copy.deepcopy(body[0]),copy.deepcopy(body[2])];pending=[(ordered[0],{'id':'table'},'body','body.table',profiles[0])]
            changes,warnings=layout021.apply_boundaries(ordered,pending,resolver)
            self.assertEqual(len(ordered),3);self.assertEqual(changes[0]['kind'],'table_spacing');self.assertEqual(warnings,[])
            changes,warnings=layout021.apply_boundaries(ordered,pending,resolver)
            self.assertEqual(len(ordered),3);self.assertEqual(changes,[])

    def test_approval_rejects_catalogue_only_success(self):
        from fastapi.testclient import TestClient
        from articleflow.web017 import create_app
        with tempfile.TemporaryDirectory() as d:
            d=Path(d);doc=Document();doc.add_paragraph('Title');doc.save(d/'a.docx')
            def catalogue_only(paths,spec,folder,progress,**kwargs):
                r={'status':'passed','samples':[{'attempts':[{'strategy':'complete_property_catalogue','status':'passed'}]}]};progress(r);return r
            with TestClient(create_app(d/'data')) as c,patch.object(agent,'run',side_effect=catalogue_only):
                h={'X-CSRF-Token':c.get('/api/v017/config').json()['csrf_token']}
                tid=c.post('/api/v017/templates',headers=h,data={'name':'sample'},files={'files':('sample.docx',(d/'a.docx').read_bytes())}).json()['id']
                response=c.post(f'/api/v017/templates/{tid}/approve',headers=h,json={'selected':{}})
                self.assertEqual(response.status_code,409)
                self.assertEqual(response.json()['detail']['code'],'FULL_DOCUMENT_CHECK_FAILED')
                self.assertFalse(c.get(f'/api/v017/templates/{tid}').json()['approved'])
