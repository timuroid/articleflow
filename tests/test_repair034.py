import io,json,tarfile,tempfile,unittest
from contextlib import nullcontext
from pathlib import Path
from unittest.mock import patch
from docx import Document
from docx.shared import Pt
from articleflow.docx_engine import DocumentError
from articleflow.sandbox034 import propose,unpack,snapshot,file_hashes
from articleflow.managed034 import load_engine


def host_compile(engine,source,template,observation,assignments,marks,destination):
    """Exercise the real compiler without Docker in focused offline tests."""
    from articleflow.leakage034 import compile_probe
    return compile_probe(source,template,observation[0],observation[1],assignments,marks)

class RepairIsolationTests(unittest.TestCase):
    def test_patch_cannot_edit_evaluator_or_escape_directory(self):
        with tempfile.TemporaryDirectory() as td:
            base=Path(td)/'base';snapshot(base);before=file_hashes(base)
            for path in ['format032.py','../rules023.py','worker034.py','agent023.py']:
                with self.assertRaisesRegex(DocumentError,'PATCH_FILE_FORBIDDEN'):
                    propose(base,Path(td)/'out',{'diagnosis':'test','edits':[{'file':path,'old':'x','new':'y'}]})
            self.assertEqual(file_hashes(base),before)

    def test_invalid_patch_never_mutates_base(self):
        with tempfile.TemporaryDirectory() as td:
            base=Path(td)/'base';snapshot(base);before=file_hashes(base)
            for old,new,code in [('not present','replacement','PATCH_ANCHOR_NOT_UNIQUE'),('def fit(examples):','def :','PATCH_SYNTAX')]:
                with self.assertRaisesRegex(DocumentError,code):propose(base,Path(td)/'out',{'diagnosis':'test','edits':[{'file':'rules023.py','old':old,'new':new}]})
            self.assertEqual(file_hashes(base),before)

    def test_worker_archive_cannot_write_symlink_or_parent(self):
        for name,link in [('../outside',False),('/absolute',False),('leak',True)]:
            data=io.BytesIO()
            with tarfile.open(fileobj=data,mode='w') as t:
                item=tarfile.TarInfo(name)
                if link:item.type=tarfile.SYMTYPE;item.linkname='/tmp/private'
                t.addfile(item)
            with tempfile.TemporaryDirectory() as td:
                with self.assertRaisesRegex(DocumentError,'WORKER_UNSAFE_ARCHIVE'):unpack(data.getvalue(),Path(td)/'out')

    def test_runtime_rejects_changed_engine_before_docker(self):
        with tempfile.TemporaryDirectory() as td:
            root=Path(td).resolve();snapshot(root/'engine');profile={'path':'engine','hashes':file_hashes(root/'engine'),'image':'sha256:test'}
            with (root/'engine/order024.py').open('a') as f:f.write('\n# changed\n')
            with patch('articleflow.managed034.Engine') as docker:
                with self.assertRaisesRegex(DocumentError,'ENGINE_PROFILE_CHANGED'):load_engine(root,profile)
                docker.assert_not_called()

    def test_sealed_probe_details_never_reach_patch_model(self):
        from articleflow.repair034 import messages
        with tempfile.TemporaryDirectory() as td:
            base=Path(td);(base/'order024.py').write_text('pass\n')
            prompt=messages(base,[{'id':'development','passed':False}],{'code':'SEALED_REGRESSION'},['order024.py'])
            self.assertNotIn('golden.docx',json.dumps(prompt));self.assertIn('SEALED_REGRESSION',json.dumps(prompt))
            self.assertIn('Incoming DOCX typography',prompt[0]['content'])

    def test_typography_perturbation_preserves_semantic_structure(self):
        from articleflow.docx_engine import NS,load_package,parse_xml
        from articleflow.integrity023 import check
        from articleflow.leakage034 import perturb
        with tempfile.TemporaryDirectory() as td:
            source=Path(td)/'source.docx';document=Document()
            document.add_heading('Heading',level=1)
            paragraph=document.add_paragraph();paragraph.add_run('Emphasis').bold=True
            paragraph.add_run(' and ordinary text')
            table=document.add_table(rows=1,cols=2)
            table.cell(0,0).text='Left';table.cell(0,1).text='Right'
            document.save(source)
            outputs=[]
            source_parts=load_package(source)
            original_colors=len(parse_xml(source_parts['word/document.xml']).findall('.//w:rPr/w:color',NS))
            for variant in (0,1):
                target=Path(td)/f'corrupted-{variant}.docx'
                self.assertGreater(perturb(source,target,variant),0)
                self.assertTrue(all(check(source,target).values()))
                target_parts=load_package(target)
                self.assertEqual(set(target_parts),set(source_parts))
                self.assertGreater(len(parse_xml(target_parts['word/document.xml']).findall('.//w:rPr/w:color',NS)),
                                   original_colors)
                outputs.append(target.read_bytes())
            self.assertNotEqual(outputs[0],outputs[1])

    def test_typography_gate_rejects_added_unstable_selector_and_allows_structure(self):
        from articleflow.leakage034 import compare
        def observation(source, changed_color, changed_bold):
            block={'id':'b-000001','text':'Sample','atoms':[]}
            paragraph={'role':'body','list_level':-1,'language':'latin'}
            semantic={'b-000001':[{'start':0,'end':6,'features':{
                'bold':changed_bold,'color':changed_color,
                'shading':changed_color,'highlight':changed_color}}]}
            return {'blocks':[block]},semantic,{'paragraph':{'b-000001':paragraph},'layout':{},'drawing':{}}
        with tempfile.TemporaryDirectory() as td:
            root=Path(td);source=root/'reference.docx';document=Document();document.add_paragraph('Sample');document.save(source)
            baseline=root/'baseline';baseline.mkdir()
            baseline.joinpath('semantic.json').write_text(json.dumps({'rules':[
                {'kind':'character','when':{'role':'body','bold':False}}]}))
            training={'samples':[{'source':source}], 'labels':[{'b-000001':'body'}]}
            def fake_observe(engine, path, extracted, assignments):
                self.assertEqual(assignments,training['labels'][0])
                corrupted=Path(path).name.startswith('input-')
                return observation(path, 'blue' if corrupted else 'red', corrupted)
            with patch('articleflow.leakage034._observe',side_effect=fake_observe),\
                 patch('articleflow.leakage034._compile',return_value={'status':'planned','fingerprint':'stable'}):
                leaking=root/'leaking';leaking.mkdir()
                leaking.joinpath('semantic.json').write_text(json.dumps({'rules':[
                    {'kind':'character','when':{'role':'body','bold':False,
                                                'color':'red','shading':'red','highlight':'red'}}]}))
                result=compare('base','candidate',training,baseline,leaking,root/'probe-leaking')
                self.assertFalse(result['passed']);self.assertEqual(result['code'],'TYPOGRAPHY_LEAKAGE')
                self.assertEqual([(x['kind'],x['selector']) for x in result['unstable_selectors']],
                                 [('character','color'),('character','highlight'),('character','shading')])
                structural=root/'structural';structural.mkdir()
                structural.joinpath('semantic.json').write_text(json.dumps({'rules':[
                    {'kind':'character','when':{'role':'body','bold':False,'list_level':-1}}]}))
                result=compare('base','candidate',training,baseline,structural,root/'probe-structural')
                self.assertTrue(result['passed']);self.assertEqual(result['unstable_selectors'],[])

    def test_typography_gate_rejects_new_decision_on_existing_bold_selector(self):
        from articleflow.leakage034 import compare
        with tempfile.TemporaryDirectory() as td:
            root=Path(td);source=root/'reference.docx'
            document=Document();document.add_paragraph('Title');document.add_paragraph('Body')
            document.save(source)
            base=root/'base';base.mkdir()
            candidate=root/'candidate';candidate.mkdir()
            common=[{'kind':'character','when':{'role':'title','bold':False},
                     'set':{'style':'title'}},
                    {'kind':'character','when':{'role':'body'},'set':{'style':'body'}}]
            base.joinpath('semantic.json').write_text(json.dumps({'rules':common}))
            candidate.joinpath('semantic.json').write_text(json.dumps({'rules':common+[
                {'kind':'character','when':{'role':'body','bold':False},
                 'set':{'style':'body_from_source_bold'}}]}))
            training={'samples':[{'source':source}],
                      'labels':[{'b-000001':'title','b-000002':'body'}]}
            def observe(engine,path,extracted,assignments):
                bold=Path(path).name.startswith('input-')
                blocks=[{'id':'b-000001','text':'Title'}, {'id':'b-000002','text':'Body'}]
                semantics={b['id']:[{'start':0,'end':len(b['text']),
                                     'features':{'bold':bold}}] for b in blocks}
                paragraphs={b['id']:{'role':assignments[b['id']], 'language':'latin'}
                            for b in blocks}
                return {'blocks':blocks},semantics,{'paragraph':paragraphs,'layout':{},'drawing':{}}
            with patch('articleflow.leakage034._observe',side_effect=observe),\
                 patch('articleflow.leakage034._compile',return_value={'status':'planned','fingerprint':'stable'}):
                result=compare('base','candidate',training,base,candidate,root/'probe')
            self.assertFalse(result['passed'])
            self.assertEqual(result['code'],'TYPOGRAPHY_LEAKAGE')
            self.assertEqual(result['unstable_selectors'],[])
            self.assertEqual(result['unstable_decisions'],[
                {'kind':'character','changed':8,'examined':18}])

    def test_typography_gate_allows_existing_semantic_bold_dependence(self):
        from articleflow.leakage034 import compare
        with tempfile.TemporaryDirectory() as td:
            root=Path(td);source=root/'reference.docx'
            document=Document();document.add_paragraph('Term');document.save(source)
            rules=[{'kind':'character','when':{'role':'term','bold':False},'set':{'style':'plain'}},
                   {'kind':'character','when':{'role':'term','bold':True},'set':{'style':'bold'}}]
            for name in ('base','candidate'):
                folder=root/name;folder.mkdir()
                folder.joinpath('semantic.json').write_text(json.dumps({'rules':rules}))
            training={'samples':[{'source':source}], 'labels':[{'b-000001':'term'}]}
            def observe(engine,path,extracted,assignments):
                changed=Path(path).name.startswith('input-')
                return ({'blocks':[{'id':'b-000001','text':'Term'}]},
                        {'b-000001':[{'start':0,'end':4,'features':{'bold':changed}}]},
                        {'paragraph':{'b-000001':{'role':'term','language':'latin'}},
                         'layout':{},'drawing':{}})
            def compile_result(engine,path,*args):
                return {'status':'planned','fingerprint':'bold' if Path(path).name.startswith('input-') else 'plain'}
            with patch('articleflow.leakage034._observe',side_effect=observe),\
                 patch('articleflow.leakage034._compile',side_effect=compile_result):
                result=compare('base','candidate',training,root/'base',root/'candidate',root/'probe')
            self.assertTrue(result['passed'])
            self.assertEqual(result['unstable_decisions'],[])
            self.assertEqual(result['unstable_compilations'],0)

    def test_typography_gate_detects_new_actual_compiler_instability(self):
        from articleflow.leakage034 import compare
        with tempfile.TemporaryDirectory() as td:
            root=Path(td);source=root/'reference.docx'
            document=Document();document.add_paragraph('Body');document.save(source)
            for name in ('base','candidate'):
                folder=root/name;folder.mkdir()
                folder.joinpath('semantic.json').write_text(json.dumps({'rules':[
                    {'kind':'paragraph','when':{'role':'body'},'set':{'style':'body'}}]}))
            training={'samples':[{'source':source}], 'labels':[{'b-000001':'body'}]}
            def observe(engine,path,extracted,assignments):
                return ({'blocks':[{'id':'b-000001','text':'Body'}]},
                        {'b-000001':[{'start':0,'end':4,'features':{'bold':False}}]},
                        {'paragraph':{'b-000001':{'role':'body','language':'latin'}},
                         'layout':{},'drawing':{}})
            class Compiler:
                def __init__(self,name):self.name=name
                def invoke(self,operation,payload,inputs,destination):
                    self.assert_operation(operation)
                    changed=(self.name=='candidate' and
                             Path(inputs['source.docx']).name.startswith('input-'))
                    return {'status':'planned','fingerprint':'new' if changed else 'original'}
                def assert_operation(self,operation):
                    if operation!='compile':raise AssertionError(operation)
            with patch('articleflow.leakage034._observe',side_effect=observe):
                result=compare(Compiler('base'),Compiler('candidate'),training,
                               root/'base',root/'candidate',root/'probe')
            self.assertFalse(result['passed'])
            self.assertEqual(result['unstable_selectors'],[])
            self.assertEqual(result['unstable_decisions'],[])
            self.assertEqual(result['unstable_compilations'],2)

    def test_compiler_fingerprint_ignores_equal_style_span_segmentation(self):
        from articleflow.leakage034 import _plan_fingerprint
        common={'schema_version':'1','section_style':'section',
                'layout':[],'objects':[]}
        single={**common,'blocks':[{'id':'b-000001','role':'body',
                 'paragraph_style':'paragraph','spans':[
                     {'start':0,'end':10,'style':'normal'}]}]}
        split={**common,'blocks':[{'id':'b-000001','role':'body',
                'paragraph_style':'paragraph','spans':[
                    {'start':0,'end':4,'style':'normal'},
                    {'start':4,'end':10,'style':'normal'}]}]}
        self.assertEqual(_plan_fingerprint(single),_plan_fingerprint(split))
        split['schema_version']='new-version'
        split['provenance']='different-extraction-path'
        split['template_sha256']='different-template-metadata'
        self.assertEqual(_plan_fingerprint(single),_plan_fingerprint(split))
        split['blocks'][0]['spans'][1]['style']='emphasis'
        self.assertNotEqual(_plan_fingerprint(single),_plan_fingerprint(split))

    def test_typography_leakage_rejects_before_sealed_probe(self):
        from articleflow.repair034 import run
        class FakeEngine:
            def __init__(self,source,image):self.hashes=file_hashes(source);self.image='sha256:fixture'
        with tempfile.TemporaryDirectory() as td:
            root=Path(td);snapshot(root/'source');source=root/'reference.docx';Document().save(source)
            training={'samples':[{'source':source}], 'labels':[{'b-000001':'body'}], 'roles':{'body':'Body'}}
            cases=[{'id':'document','split':'development','training':training}, {'id':'guard','split':'sealed'}]
            calls=[]
            def evaluate(engine,case,folder):
                calls.append(case['id'])
                if case.get('kind')=='source_gate':return {'id':case['id'],'split':'development','passed':True}
                if case['split']=='sealed':return {'id':case['id'],'split':'sealed','passed':True}
                return {'id':case['id'],'split':'development','passed':False,
                        'conflicts':35 if 'baseline' in folder.parts else 25}
            def caller(config,messages,directory,**kwargs):
                return {'diagnosis':'fixture','edits':[{'file':'rules023.py','old':'def fit(examples):',
                        'new':'def fit(examples):\n    # fixture'}]}, {'api_attempts':0}
            with patch('articleflow.repair034.Engine',FakeEngine),\
                 patch('articleflow.repair034.evaluate',side_effect=evaluate),\
                 patch('articleflow.leakage034.compare',return_value={
                     'passed':False,'code':'TYPOGRAPHY_LEAKAGE','unstable_selectors':[{'kind':'character','selector':'color'}]}):
                report=run(root/'source',cases,root/'run',{},caller=caller,max_attempts=1)
            self.assertEqual(report['attempts'][0]['code'],'TYPOGRAPHY_LEAKAGE')
            self.assertEqual(calls.count('guard'),1)  # baseline only

    def test_sealed_failure_is_redacted_in_real_repair_loop_feedback(self):
        from articleflow.repair034 import run
        prompts=[]
        class FakeEngine:
            def __init__(self,source,image):self.hashes=file_hashes(source);self.image='sha256:fixture'
        def evaluate(engine,case,folder):
            baseline='baseline' in folder.parts
            return {'id':case['id'],'split':case['split'],'passed':not baseline and case['split']=='development','private_detail':case['private']}
        def caller(config,messages,directory,**kwargs):
            prompts.append(json.dumps(messages))
            return {'diagnosis':'fixture','edits':[{'file':'rules023.py','old':'def fit(examples):','new':'def fit(examples):\n    # probe'}]}, {'api_attempts':0}
        with tempfile.TemporaryDirectory() as td:
            root=Path(td);snapshot(root/'source')
            cases=[{'id':'development','split':'development','private':'allowed'}, {'id':'sealed_only_sentinel','split':'sealed','private':'hidden_expected_text_sentinel'}]
            with patch('articleflow.repair034.Engine',FakeEngine),patch('articleflow.repair034.evaluate',side_effect=evaluate):
                result=run(root/'source',cases,root/'run',{},caller=caller)
            self.assertEqual(result['status'],'needs_review');self.assertEqual(len(prompts),2)
            for prompt in prompts:
                self.assertNotIn('sealed_only_sentinel',prompt);self.assertNotIn('hidden_expected_text_sentinel',prompt)

    def test_partial_progress_requires_no_case_regression(self):
        from articleflow.repair034 import improves
        before=[{'id':'a','split':'development','conflicts':33},{'id':'b','split':'development','passed':True}]
        self.assertTrue(improves(before,[{'id':'a','conflicts':25},{'id':'b','passed':True}]))
        self.assertFalse(improves(before,[{'id':'a','conflicts':25},{'id':'b','conflicts':1}]))
        self.assertFalse(improves(before,[{'id':'a','conflicts':33},{'id':'b','passed':True}]))
        self.assertFalse(improves(before,[{'id':'a','conflicts':1}]))
        measured=[{'id':'a','split':'development','conflicts':33,'runtime_unresolved_count':5},
                  {'id':'b','split':'development','passed':True}]
        self.assertFalse(improves(measured,[{'id':'a','conflicts':25,'runtime_unresolved_count':6},
                                           {'id':'b','passed':True}]))
        self.assertTrue(improves(measured,[{'id':'a','conflicts':25,'runtime_unresolved_count':4},
                                          {'id':'b','passed':True}]))

    def test_partial_conflict_improvement_requires_exact_source_instances(self):
        from articleflow.repair034 import run
        class FakeEngine:
            def __init__(self,source,image):self.hashes=file_hashes(source);self.image='sha256:fixture'
        for source_exact in (False,True):
            with self.subTest(source_exact=source_exact),tempfile.TemporaryDirectory() as td:
                root=Path(td);snapshot(root/'source');source=root/'reference.docx'
                Document().save(source)
                training={'samples':[{'source':source}], 'labels':[{'b-000001':'body'}],
                          'roles':{'body':'Body'}}
                cases=[{'id':'document','split':'development','training':training},
                       {'id':'guard','split':'sealed'}]
                calls=[]
                def evaluate(engine,case,folder):
                    calls.append((case['id'],str(folder)))
                    if case.get('kind')=='source_gate':
                        self.assertEqual(case['goldens'],[source])
                        self.assertIs(case['training'],training)
                        return {'id':case['id'],'split':'development','passed':source_exact,
                                'instances':[{'format':{'property_exact':source_exact},
                                              'pixels':{'pixel_exact':source_exact}}]}
                    if case['split']=='sealed':
                        return {'id':case['id'],'split':'sealed','passed':True}
                    return {'id':case['id'],'split':'development','passed':False,
                            'conflicts':33 if 'baseline' in folder.parts else 30}
                def caller(config,messages,directory,**kwargs):
                    return {'diagnosis':'reduce conflicts','edits':[{'file':'rules023.py',
                            'old':'def fit(examples):','new':'def fit(examples):\n    # source gate fixture'}]}, {'api_attempts':0}
                with patch('articleflow.repair034.Engine',FakeEngine),\
                     patch('articleflow.repair034.evaluate',side_effect=evaluate),\
                     patch('articleflow.leakage034.compare',return_value={'passed':True,'unstable_selectors':[]}):
                    result=run(root/'source',cases,root/'repair',{},caller=caller,max_attempts=1)
                self.assertEqual(result['status'],'needs_review')
                item=result['attempts'][0]
                self.assertEqual(len(item['source_instances']),1)
                self.assertEqual(item['status'],'candidate_progress' if source_exact else 'rejected')
                self.assertEqual(item.get('code'),None if source_exact else 'SOURCE_INSTANCE_REGRESSION')
                candidate_guards=[p for name,p in calls if name=='guard' and 'attempt-1' in p]
                self.assertEqual(bool(candidate_guards),source_exact)

    def test_accepted_patch_requires_fresh_template_recheck(self):
        from articleflow.managed034 import run
        for fresh_status in ['needs_review','reference_verified']:
            with self.subTest(fresh_status=fresh_status),tempfile.TemporaryDirectory() as td:
                root=Path(td);folder=root/'run';calls=[]
                def learn(references,directory,*args,**kwargs):
                    calls.append(directory);directory.mkdir(parents=True)
                    return {'schema':'teacher-student-023','status':'needs_review' if len(calls)==1 else fresh_status,'iterations':[]}
                def repair(base,cases,directory,*args,**kwargs):
                    directory.mkdir();engine=directory/'engine';engine.mkdir();profile=directory/'profile.json'
                    profile.write_text(json.dumps({'hashes':{},'image':'sha256:fixture'}))
                    return {'status':'accepted','engine':str(engine),'profile':str(profile)}
                with patch('articleflow.agent028.run',side_effect=learn),patch('articleflow.managed034.collect',return_value=[{'id':'dev','split':'development'}]),patch('articleflow.probes034.create',return_value=[{'id':'guard','split':'sealed'}]),patch('articleflow.repair034.run',side_effect=repair),patch('articleflow.managed034.load_engine',return_value=None):
                    result=run([],folder,{}, {},repair_image='fixture')
                self.assertEqual(len(calls),2);self.assertEqual(result['status'],fresh_status)
                self.assertEqual(result['engine_learning_root'],'engine-replay/learning')
                self.assertTrue((folder/'before-engine-repair.json').exists())

    def test_import_failure_generates_instance_probe_without_student_response(self):
        from articleflow.managed034 import collect
        with tempfile.TemporaryDirectory() as td:
            root=Path(td);source=root/'reference.docx'
            cases=collect([source],root/'run',{'iterations':[]})
            self.assertEqual(len(cases),1);self.assertEqual(cases[0]['kind'],'instance')
            self.assertNotIn('response',cases[0]);self.assertEqual(cases[0]['golden'],source)

    def test_training_conflict_without_student_is_replayable(self):
        from articleflow.agent023 import extract_input
        from articleflow.managed034 import collect
        from articleflow.repair034 import evaluate
        with tempfile.TemporaryDirectory() as td:
            root=Path(td);source=root/'reference.docx';document=Document()
            for text,size in [('Opening body',11),('First middle body',11),
                              ('Second middle body',16),('Closing body',11)]:
                run=document.add_paragraph(text).runs[0];run.font.size=Pt(size)
            document.save(source)
            run_root=root/'run';extract_input(source,run_root/'learning/reference-1/extracted')
            teacher=run_root/'learning/iteration-1/teacher';teacher.mkdir(parents=True)
            teacher.joinpath('validated.json').write_text(json.dumps({
                'roles':{'body':'Body paragraph'},'groups':[{'start':1,'end':4,'role':'body'}],
                'questions':[]}))
            cases=collect([source],run_root,{'iterations':[{}]})
            self.assertEqual(len(cases),1);self.assertEqual(cases[0]['kind'],'training')
            self.assertNotIn('response',cases[0]);self.assertNotIn('template',cases[0])
            with patch('articleflow.repair034.use',side_effect=lambda engine:nullcontext()),\
                 patch('articleflow.leakage034._compile',side_effect=host_compile),\
                 patch('articleflow.agent023.apply') as student_apply,\
                 patch('articleflow.engine034.render') as render:
                row=evaluate(object(),cases[0],root/'probe')
            self.assertFalse(row['passed']);self.assertGreater(row['conflicts'],0)
            student_apply.assert_not_called();render.assert_not_called()

    def test_training_probe_checks_compiled_reference_before_student_exists(self):
        from articleflow.agent023 import extract_input
        from articleflow.catalogue import render_catalogue
        from articleflow.managed034 import collect
        from articleflow.repair034 import evaluate
        with tempfile.TemporaryDirectory() as td:
            root=Path(td);source=root/'reference.docx';document=Document()
            for text,size in [('Title',16),('Body',11)]:
                run=document.add_paragraph(text).runs[0];run.font.size=Pt(size)
            document.save(source)
            run_root=root/'run';extract_input(source,run_root/'learning/reference-1/extracted')
            teacher=run_root/'learning/iteration-1/teacher';teacher.mkdir(parents=True)
            teacher.joinpath('validated.json').write_text(json.dumps({
                'roles':{'title':'Document title','body':'Body paragraph'},
                'groups':[{'start':1,'end':1,'role':'title'},{'start':2,'end':2,'role':'body'}],
                'questions':[]}))
            # The last teacher attempt may fail validation; retain the latest
            # earlier complete semantic answer for a controlled replay.
            cases=collect([source],run_root,{'iterations':[{},{}]})
            self.assertEqual(cases[0]['kind'],'training')
            def trusted_render(engine,template,content,plan,target):
                render_catalogue(template,content,plan,target)
            with patch('articleflow.repair034.use',side_effect=lambda engine:nullcontext()),\
                 patch('articleflow.leakage034._compile',side_effect=host_compile),\
                 patch('articleflow.engine034.render',side_effect=trusted_render),\
                 patch('articleflow.research.render_pdf',side_effect=lambda source,folder:source),\
                 patch('articleflow.research.compare_pixels',return_value={'pixel_exact':True}),\
                 patch('articleflow.agent023.apply') as student_apply:
                row=evaluate(object(),cases[0],root/'probe')
            self.assertTrue(row['passed']);self.assertEqual(row['conflicts'],0)
            self.assertEqual(len(row['instances']),1)
            self.assertTrue(row['instances'][0]['format']['property_exact'])
            self.assertTrue(all(row['instances'][0]['integrity'].values()))
            self.assertTrue(row['runtime_exact'])
            student_apply.assert_not_called()

    def test_training_probe_rejects_oracle_replay_with_empty_runtime_rules(self):
        from articleflow.agent023 import extract_input,build_package
        from articleflow.catalogue import render_catalogue
        from articleflow.managed034 import collect
        from articleflow.repair034 import evaluate
        with tempfile.TemporaryDirectory() as td:
            root=Path(td);source=root/'reference.docx';document=Document()
            for title,size in [('Title',16),('Body',11)]:
                document.add_paragraph(title).runs[0].font.size=Pt(size)
            document.save(source)
            run_root=root/'run';extract_input(source,run_root/'learning/reference-1/extracted')
            teacher=run_root/'learning/iteration-1/teacher';teacher.mkdir(parents=True)
            teacher.joinpath('validated.json').write_text(json.dumps({
                'roles':{'title':'Document title','body':'Body paragraph'},
                'groups':[{'start':1,'end':1,'role':'title'},
                          {'start':2,'end':2,'role':'body'}], 'questions':[]}))
            case=collect([source],run_root,{'iterations':[{}]})[0]
            def empty_rules(samples,labels,roles,folder,inline_roles=None):
                package,conflicts=build_package(samples,labels,roles,folder,inline_roles)
                self.assertEqual(conflicts,[])
                path=folder/'semantic.json';stored=json.loads(path.read_text())
                stored['rules']=[];path.write_text(json.dumps(stored))
                return package,conflicts
            def trusted_render(engine,template,content,plan,target):
                render_catalogue(template,content,plan,target)
            with patch('articleflow.repair034.use',side_effect=lambda engine:nullcontext()),\
                 patch('articleflow.agent023.build_package',side_effect=empty_rules),\
                 patch('articleflow.leakage034._compile',side_effect=host_compile),\
                 patch('articleflow.engine034.render',side_effect=trusted_render),\
                 patch('articleflow.research.render_pdf',side_effect=lambda source,folder:source),\
                 patch('articleflow.research.compare_pixels',return_value={'pixel_exact':True}):
                row=evaluate(object(),case,root/'probe')
            self.assertEqual(row['conflicts'],0)
            self.assertEqual(len(row['instances']),1)
            self.assertTrue(row['instances'][0]['passed'])
            self.assertFalse(row['runtime_exact'])
            self.assertGreater(row['runtime_unresolved_count'],0)
            self.assertFalse(row['passed'])
            # A candidate that returns a complete but incorrect style plan
            # must fail even though the separate oracle replay is pixel-exact.
            with patch('articleflow.repair034.use',side_effect=lambda engine:nullcontext()),\
                 patch('articleflow.agent023.build_package',side_effect=empty_rules),\
                 patch('articleflow.leakage034._compile',return_value={
                     'status':'planned','fingerprint':'wrong-style-plan'}),\
                 patch('articleflow.engine034.render',side_effect=trusted_render),\
                 patch('articleflow.research.render_pdf',side_effect=lambda source,folder:source),\
                 patch('articleflow.research.compare_pixels',return_value={'pixel_exact':True}):
                wrong=evaluate(object(),case,root/'wrong-style')
            self.assertTrue(wrong['instances'][0]['passed'])
            self.assertEqual(wrong['runtime_wrong_style_count'],1)
            self.assertFalse(wrong['passed'])

    def test_source_gate_reconstructs_even_with_conflicts_and_rejects_damage(self):
        from articleflow.agent023 import extract_input
        from articleflow.catalogue import render_catalogue
        from articleflow.managed034 import collect
        from articleflow.repair034 import evaluate
        with tempfile.TemporaryDirectory() as td:
            root=Path(td);source=root/'reference.docx';document=Document()
            for text,size in [('Opening body',11),('First middle body',11),
                              ('Second middle body',16),('Closing body',11)]:
                run=document.add_paragraph(text).runs[0];run.font.size=Pt(size)
            document.save(source)
            run_root=root/'run';extract_input(source,run_root/'learning/reference-1/extracted')
            teacher=run_root/'learning/iteration-1/teacher';teacher.mkdir(parents=True)
            teacher.joinpath('validated.json').write_text(json.dumps({
                'roles':{'body':'Body paragraph'},'groups':[{'start':1,'end':4,'role':'body'}],
                'questions':[]}))
            training=collect([source],run_root,{'iterations':[{}]})[0]['training']
            case={'id':'source-check','kind':'source_gate','split':'development',
                  'training':training,'goldens':[source]}
            def damaged_render(engine,template,content,plan,target):
                render_catalogue(template,content,plan,target)
                damaged=Document(target)
                damaged.paragraphs[2].runs[0].font.size=Pt(9)
                damaged.save(target)
            with patch('articleflow.repair034.use',side_effect=lambda engine:nullcontext()),\
                 patch('articleflow.leakage034._compile',side_effect=host_compile),\
                 patch('articleflow.engine034.render',side_effect=damaged_render),\
                 patch('articleflow.research.render_pdf',side_effect=lambda source,folder:source),\
                 patch('articleflow.research.compare_pixels',return_value={'pixel_exact':False,'differing_pixel_ratio':0.1}):
                row=evaluate(object(),case,root/'probe')
            self.assertGreater(row['conflicts'],0)
            self.assertEqual(len(row['instances']),1)
            self.assertFalse(row['instances'][0]['format']['property_exact'])
            self.assertFalse(row['instances'][0]['pixels']['pixel_exact'])
            self.assertFalse(row['passed'])
