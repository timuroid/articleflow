import json
import tempfile
from pathlib import Path
from unittest.mock import patch

from behave import given, when, then
from docx import Document

from articleflow.docx_engine import (
    NS, compare_content, load_package, lossless_split, lossless_rebuild,
    parse_xml, write_package, xml_bytes,
)
from articleflow.docx_engine import DocumentError
from articleflow.research import one_model_call, plan_schema, finalize_hybrid_trial, validate_model_choice


def temporary_workspace(context):
    temp = tempfile.TemporaryDirectory()
    context.add_cleanup(temp.cleanup)
    context.root = Path(temp.name)


@given("создан синтетический DOCX с выделением и верхним индексом")
def source_document(context):
    temporary_workspace(context)
    context.source = context.root / "source.docx"
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("Определение").bold = True
    p.add_run(". Значение 1.25 м")
    p.add_run("2").font.superscript = True
    doc.save(context.source)


@when("исходные свойства разделены и восстановлены без модели")
def reconstruct(context):
    context.manifest = lossless_split(context.source, context.root / "split")
    context.target = context.root / "result.docx"
    lossless_rebuild(context.root / "split", context.target)


@then("содержимое документа совпадает")
def content_matches(context):
    assert compare_content(context.source, context.target)["content_exact"]


@then("контроль не допускается в модельный benchmark")
def not_a_trial(context):
    assert context.manifest["model_calls"] == 0
    assert not context.manifest["eligible_for_model_benchmark"]


@when("из результата удалён верхний индекс без изменения букв")
def remove_superscript(context):
    parts = load_package(context.source)
    tree = parse_xml(parts["word/document.xml"])
    for element in tree.findall(".//w:vertAlign", NS):
        element.getparent().remove(element)
    parts["word/document.xml"] = xml_bytes(tree)
    context.target = context.root / "changed.docx"
    write_package(context.target, parts)


@then("проверка содержимого не проходит")
def content_fails(context):
    result = compare_content(context.source, context.target)
    assert result["checks"]["paragraphs"]
    assert not result["checks"]["protected_vertical"]
    assert not result["content_exact"]


@given("тестовый провайдер возвращает невалидный JSON")
def bad_provider(context):
    temporary_workspace(context)
    context.secret = "synthetic-not-a-real-key"


@when("выполняется одна попытка планирования")
def request_once(context):
    with patch.dict("os.environ", {"OPENROUTER_API_KEY": context.secret}), patch("articleflow.research.httpx.Client") as cls:
        client = cls.return_value.__enter__.return_value
        client.post.return_value.status_code = 200
        client.post.return_value.json.return_value = {"usage": {"cost": 0.01}, "choices": [{"finish_reason": "stop", "message": {"content": "not json"}}]}
        context.meta = one_model_call(model="synthetic", messages=[], schema=plan_schema(["body"]), pricing={"prompt": "0.000001", "completion": "0.000001"}, key_file=None, run_dir=context.root / "run", ledger=context.root / "ledger.json", cap_usd="1")
        context.call_count = client.post.call_count


@then("провайдер вызван ровно один раз без repair")
def no_retry(context):
    assert context.call_count == 1
    assert context.meta["api_attempts"] == 1
    assert context.meta["repair_calls"] == 0
    assert context.meta["status"] == "invalid_model_output"


@then("секрет отсутствует в сохранённых артефактах")
def secret_not_saved(context):
    assert all(context.secret not in path.read_text() for path in context.root.rglob("*.json"))


@given("гибридный renderer сохранил исходную геометрию")
def hybrid(context):
    context.row = {"render_protocol": "I1-text-to-model/I3-object-layout-to-renderer"}


@when("все измеренные проверки совпадения успешны")
def successful_checks(context):
    context.row.update(content_exact=True, inline_exact=True, pixel_exact=True, roles_exact=True)
    finalize_hybrid_trial(context.row)


@then("наблюдаемое совпадение отмечено отдельно")
def observed_match(context):
    assert context.row["observed_full_match"]


@then("строгий I1 не пройден")
def strict_not_passed(context):
    assert not context.row["eligible_for_strict_I1"]
    assert not context.row["strict_I1_pass"]


@given("подтверждён выпуск модели 30 июня 2026 года")
def old_release(context):
    context.registry = {"release_after": "2026-06-30", "models": [{"id": "synthetic", "release_date": "2026-06-30", "source_type": "official_release_log", "source": "https://example.org/releases"}]}


@when("проверяется допуск с границей после июня 2026 года")
def check_release(context):
    context.error = None
    with patch("articleflow.research.httpx.Client") as client:
        try:
            validate_model_choice("synthetic", context.registry)
        except DocumentError as exc:
            context.error = str(exc)
        context.http_client_created = client.called


@then("модель отклоняется до платного запроса")
def old_release_rejected(context):
    assert context.error == "MODEL_TOO_OLD"
    assert not context.http_client_created
