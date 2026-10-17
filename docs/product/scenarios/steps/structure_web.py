"""Executable Web 019 recovery scenarios; synthetic documents, no network."""
import io
import json
import tempfile
from pathlib import Path
from zipfile import ZipFile
from unittest.mock import patch
from behave import given, when, then
from docx import Document
from fastapi.testclient import TestClient
from articleflow.structure017 import inventory, suggest_roles, compile_document, inspect_templates
from articleflow.web017 import create_app


def setup(c):
    if hasattr(c, 'work019'): return
    temp=tempfile.TemporaryDirectory();c.add_cleanup(temp.cleanup);c.work019=Path(temp.name)
    c.doc019=Document();c.doc019.add_paragraph('Исследование тестового материала')
    c.doc019.add_paragraph('Иванов И. И.');c.doc019.add_paragraph('Аннотация')
    for text in ['Цель. Проверить структуру.', 'Результаты. Получено 42.', 'Практическая значимость. Контроль качества.', 'Постановка проблемы. Различия оформления.']:
        c.doc019.add_paragraph(text)
    c.doc019.add_paragraph('Ключевые слова: оформление');c.doc019.add_paragraph('Основной научный текст.')


def compile_test(c):
    setup(c);src=c.work019/'source.docx';c.doc019.save(src);c.source019=src.read_bytes()
    c.blocks019=inventory(src)['blocks'];c.roles019=suggest_roles(c.blocks019)[0]
    c.report019=compile_document(src,c.work019/'result.docx',assignments=c.roles019,approved=True,fonts_ok=True,editorial_output=c.work019/'editorial.docx')


@given('утверждённый шаблон требует постановку проблемы перед целью внутри аннотации')
@given('утверждённый шаблон требует постановку проблемы внутри аннотации')
@given('утверждённый шаблон требует авторов и практическую значимость')
def standard(c): setup(c)

@given('статья содержит подписанную постановку после аннотации и до ключевых слов')
def misplaced(c): setup(c)

@when('система строит структурный план статьи')
@when('система определяет роли')
@when('система собирает документ')
def build(c): compile_test(c)

@then('план связывает постановку с её исходным фрагментом')
def binding(c):
    b=next(b for b in c.blocks019 if b['text'].startswith('Постановка проблемы'))
    assert c.report019['assignments'][b['id']]=='abstract_problem'

@then('ставит её перед целью внутри аннотации')
def order(c):
    texts=[p.text for p in Document(c.work019/'result.docx').paragraphs]
    assert next(i for i,t in enumerate(texts) if t.startswith('Постановка')) < next(i for i,t in enumerate(texts) if t.startswith('Цель'))

@then('сохраняет научный текст и остальные исходные узлы')
def conserved(c):
    from collections import Counter
    assert Counter(p.text for p in c.doc019.paragraphs)==Counter(p.text for p in Document(c.work019/'result.docx').paragraphs)
    assert (c.work019/'source.docx').read_bytes()==c.source019

@given('статья содержит заголовок основной части «Проблема и постановка»')
def body_heading(c):
    setup(c)
    for p in c.doc019.paragraphs:
        if p.text.startswith('Постановка проблемы'):p._element.getparent().remove(p._element)
    c.doc019.add_paragraph('Проблема и постановка')

@then('заголовок основной части остаётся в основной части')
def stays(c):
    b=next(b for b in c.blocks019 if b['text']=='Проблема и постановка');assert c.roles019[b['id']]=='body'

@then('отсутствие постановки в аннотации отражается отдельным вопросом')
def absent_problem(c): assert 'abstract_problem' in c.report019['missing']

@given('в статье нет этих сведений')
def absent(c):
    setup(c)
    for p in c.doc019.paragraphs:
        if p.text.startswith(('Иванов','Практическая значимость')):p._element.getparent().remove(p._element)

@then('она вставляет явно помеченные места для заполнения')
def placeholders(c):
    texts=[p.text for p in Document(c.work019/'result.docx').paragraphs]
    assert '[ЗАПОЛНИТЬ: authors]' in texts and '[ЗАПОЛНИТЬ: abstract_practical]' in texts

@then('сообщает редактору, какие сведения нужны')
def needs(c): assert {'authors','abstract_practical'}<=set(c.report019['missing'])

@then('запрещает выдачу документа как готового')
def draft(c): assert c.report019['status']=='draft'

@given('статья содержит строку «Рабочая рукопись» до заголовка')
def editorial(c):
    setup(c);c.doc019.paragraphs[0].insert_paragraph_before('Рабочая рукопись')

@when('редактор подтверждает, что это редакционная заметка')
def confirm(c): compile_test(c)

@then('система не заменяет ею сведения об авторах')
def authors(c):
    b=next(b for b in c.blocks019 if b['text']=='Рабочая рукопись');assert c.roles019[b['id']]=='editorial'
    assert 'authors' not in c.report019['missing']

@then('сохраняет её в отдельном редакционном результате')
def separate(c): assert Document(c.work019/'editorial.docx').paragraphs[0].text=='Рабочая рукопись'

@then('оставляет исходный DOCX доступным')
def unchanged(c): assert (c.work019/'source.docx').read_bytes()==c.source019


def web(c):
    setup(c);mock=patch('articleflow.web017.font_audit',return_value={'exact':True,'families':{}});mock.start();c.add_cleanup(mock.stop)
    c.client019=TestClient(create_app(c.work019/'web'));c.add_cleanup(c.client019.close)
    c.headers019={'X-CSRF-Token':c.client019.get('/api/v017/config').json()['csrf_token']}


def uploaded(c, variants=False):
    web(c);files=[]
    for i in range(2 if variants else 1):
        c.doc019.styles['Normal'].font.name='Arial' if i==0 else 'Times New Roman'
        buffer=io.BytesIO();c.doc019.save(buffer)
        template=Document(io.BytesIO(buffer.getvalue()))
        problem=next(p for p in template.paragraphs if p.text.startswith('Постановка'))
        goal=next(p for p in template.paragraphs if p.text.startswith('Цель'))
        goal._p.addprevious(problem._p)
        buffer=io.BytesIO();template.save(buffer);files.append(('files',(str(i)+'.docx',buffer.getvalue(),'application/octet-stream')))
    response=c.client019.post('/api/v017/templates',headers=c.headers019,data={'name':'Synthetic'},files=files)
    assert response.status_code==200;c.template019=response.json()['id']
    c.spec019=c.client019.get('/api/v017/templates/'+c.template019).json()['spec']

@given('редактор загрузил несколько эталонных DOCX одного шаблона')
def multi(c): uploaded(c,True)

@given('в эталонах различаются параметры одного обязательного слота')
def conflict(c): assert c.spec019['conflicts']

@when('система сравнивает наблюдения')
def inspect(c): assert c.spec019['samples']

@then('она показывает оба варианта с источником')
def evidence(c):
    variants=c.spec019['variants']['__base__'];assert len(variants)==2 and all(variants.values())

@then('не утверждает спорный вариант автоматически без достаточного правила')
def not_approved(c):
    assert not c.client019.get('/api/v017/templates/'+c.template019).json()['approved']
    assert c.client019.post('/api/v017/templates/'+c.template019+'/approve',headers=c.headers019,json={'selected':{}}).status_code==409

@then('выбор варианта не отменяет несовпадение с эталонами')
def conflicting_choice_rejected(c):
    selected={r:next(iter(v)) for r,v in c.spec019['variants'].items()}
    response=c.client019.post('/api/v017/templates/'+c.template019+'/approve',headers=c.headers019,json={'selected':selected})
    assert response.status_code==409
    assert response.json()['detail']['code']=='FULL_DOCUMENT_CHECK_FAILED'
    assert c.client019.get('/api/v017/templates/'+c.template019).json()['version']==1

def version(c):
    selected={r:next(iter(v)) for r,v in c.spec019['variants'].items()}
    response=c.client019.post('/api/v017/templates/'+c.template019+'/approve',headers=c.headers019,json={'selected':selected})
    assert response.status_code==200 and response.json()['version']==2

@given('утверждённый шаблон и пакет из нескольких статей')
def batch(c):
    uploaded(c);version(c)
    buffer=io.BytesIO();c.doc019.save(buffer);c.complete019=buffer.getvalue()

@given('в одной статье недостаёт обязательного поля')
def missing_batch(c): absent(c)

@when('система обрабатывает пакет')
def process(c):
    buffer=io.BytesIO();c.doc019.save(buffer)
    response=c.client019.post('/api/v017/batches',headers=c.headers019,data={'template_id':c.template019},files=[('files',('complete.docx',c.complete019)),('files',('missing.docx',buffer.getvalue()))])
    assert response.status_code==202;c.batch019=response.json()['id']
    c.jobs019=c.client019.get('/api/v017/batches/'+c.batch019).json()['jobs']

@then('остальные статьи продолжают обрабатываться')
def continued(c): assert len(c.jobs019)==2 and all(j['status']=='draft' for j in c.jobs019)

@then('архив содержит отдельный статус и отчёт для каждого документа')
def zipped(c):
    response=c.client019.post('/api/v017/batches/'+c.batch019+'/exports',headers=c.headers019)
    assert response.status_code==200
    data=c.client019.get(response.json()['url']).content
    with ZipFile(io.BytesIO(data)) as z:
        manifest=json.loads(z.read('manifest.json'));assert len(manifest['jobs'])==2
        assert len([n for n in z.namelist() if n.endswith('.json')])>=3

@then('статья с недостающим полем остаётся черновиком')
def incomplete(c):
    job=next(j for j in c.jobs019 if j['source_name']=='missing.docx');assert job['status']=='draft'
    assert c.client019.get('/api/v017/jobs/'+job['id']+'/files/final').status_code==409

@given('подтверждённый план выносит исходный абзац в редакционные заметки')
def editorial_without_destination(c):
    setup(c);c.doc019=Document();c.doc019.add_paragraph('Содержательный исходный абзац.')
    c.source020=c.work019/'source020.docx';c.doc019.save(c.source020)
    c.bytes020=c.source020.read_bytes();c.out020=c.work019/'out020.docx'
    c.roles020={b['id']:'editorial' for b in inventory(c.source020)['blocks']}

@when('сборщик вызван без отдельного пути редакционного результата')
def missing_destination(c):
    from articleflow.docx_engine import DocumentError
    try:compile_document(c.source020,c.out020,assignments=c.roles020,required=[])
    except DocumentError as error:c.error020=str(error)
    else:c.error020=None

@then('сборка отклонена до записи выходного документа')
def no_output(c):
    assert c.error020=='EDITORIAL_OUTPUT_REQUIRED'
    assert not c.out020.exists()

@then('исходный научный абзац остаётся доступным без изменений')
def same_source(c):assert c.source020.read_bytes()==c.bytes020

@given('в статье есть только заголовок ключевых слов без самого списка')
def empty_keywords021(c):
    setup(c);c.doc019=Document();c.doc019.add_paragraph('Ключевые слова:')
    c.source021=c.work019/'source021.docx';c.doc019.save(c.source021)

@when('сборщик проверяет обязательные ключевые слова')
def require_keywords021(c):
    c.out021=c.work019/'out021.docx'
    c.report021=compile_document(c.source021,c.out021,assignments={'n-000001':'keywords'},required=['keywords'],approved=True,fonts_ok=True)

@then('отсутствие списка явно отражается в отчёте и заполнителе черновика')
def report_keywords021(c):
    assert c.report021['missing']==['keywords'] and c.report021['status']=='draft'
    assert any('keywords' in q for q in c.report021['questions'])
    assert '[ЗАПОЛНИТЬ: keywords]' in [p.text for p in Document(c.out021).paragraphs]

@given('извлечённый шаблон ожидает проверку интервалов рендером')
def visual_template021(c):
    uploaded(c)

@when('визуальный контроль сообщает несовпадение')
def visual_mismatch021(c):
    from unittest.mock import patch
    selected={r:next(iter(v)) for r,v in c.spec019['variants'].items()}
    with patch('articleflow.templatecheck021.validate_samples',return_value={'status':'failed','rows':[]}):
        c.visual_response021=c.client019.post('/api/v017/templates/'+c.template019+'/approve',headers=c.headers019,json={'selected':selected})

@then('версия шаблона остаётся неутверждённой с причиной ошибки')
def visual_block021(c):
    assert c.visual_response021.status_code==409
    assert c.visual_response021.json()['detail']['code']=='TEMPLATE_LAYOUT_MISMATCH'
    assert not c.client019.get('/api/v017/templates/'+c.template019).json()['approved']

@given('два контрольных PDF отличаются только на последней странице')
def pdf_pair022(c):
    import pymupdf
    setup(c)
    for name,last in [('a','Original last page'),('b','Changed last page')]:
        doc=pymupdf.open();doc.new_page().insert_text((70,70),'First page')
        doc.new_page().insert_text((70,70),last);doc.save(c.work019/(name+'.pdf'));doc.close()

@when('полный визуальный оценщик сравнивает документы')
def compare022(c):
    from articleflow.templateagent022 import compare_all
    c.metric022=compare_all(c.work019/'a.pdf',c.work019/'b.pdf',c.work019/'diff')

@then('последняя страница отмечена как несовпадающая с областью различий')
def last_page022(c):
    assert not c.metric022['pixel_exact']
    assert c.metric022['matched_page_diagnostics'][0]['pixel_exact']
    assert c.metric022['matched_page_diagnostics'][1]['difference_bbox_px']

@given('контроль обнаружил заголовок без следующего блока')
def orphan022(c):
    setup(c);doc=Document();doc.add_paragraph('Введение');doc.add_paragraph('Текст 42.');doc.save(c.work019/'orphan.docx')

@when('выполняется разрешённый ремонт связи переноса')
def repair022(c):
    from articleflow.documentcheck022 import repair_keep_links
    from articleflow.docx_engine import digest
    c.changes022=repair_keep_links(c.work019/'orphan.docx',c.work019/'fixed.docx',{'issues':[{'code':'ORPHAN_LABEL','text_sha256':digest('введение'.encode())}]})

@then('текст документа неизменен и у заголовка установлена связь со следующим блоком')
def repair_checked022(c):
    assert [p.text for p in Document(c.work019/'orphan.docx').paragraphs]==[p.text for p in Document(c.work019/'fixed.docx').paragraphs]
    assert Document(c.work019/'fixed.docx').paragraphs[0].paragraph_format.keep_with_next

@given('цикл извлечения имеет лимит одной компактной гипотезы')
def limited022(c):
    setup(c);c.source022=c.work019/'source022.docx';c.doc019.save(c.source022)

@when('он завершает доступный эксперимент')
def limited_run022(c):
    from articleflow.templateagent022 import run
    c.result022=run([c.source022],inspect_templates([c.source022]),c.work019/'agent',max_attempts=1)

@then('результат требует проверки и перенос на новые статьи не объявлен измеренным')
def stopped022(c):
    assert c.result022['status']=='needs_review'
    assert c.result022['samples'][0]['stop_reason']=='attempt_limit'
    assert c.result022['transfer_validation']=='not_measured'

@given('полный каталог прошёл сравнение, а рабочая сборка не совпала')
def mismatched_production022(c):
    setup(c);c.source022=c.work019/'runtime-source.docx';c.doc019.save(c.source022)

@when('проверяется допуск рабочего шаблона 022')
def runtime_gate022(c):
    from unittest.mock import patch
    from articleflow import templateagent022 as agent
    from articleflow.structure017 import inspect_templates
    import pymupdf
    p=c.work019/'fixed.pdf';doc=pymupdf.open();doc.new_page().insert_text((72,72),'Test');doc.save(p);doc.close()
    with patch.object(agent,'render_pdf',return_value=p),patch.object(agent,'_negative_controls',return_value=[{'detected':True}]),patch('articleflow.roundtrip022.verify',return_value={'strategy':'production_roundtrip','status':'failed','rows':[]}):
        c.result022=agent.run([c.source022],inspect_templates([c.source022]),c.work019/'runtime-gate')

@then('шаблон остаётся не прошедшим проверку')
def runtime_rejected022(c):
    assert c.result022['samples'][0]['attempts'][1]['status']=='passed'
    assert c.result022['status']=='needs_review'

@given('три наблюдения одного смыслового контекста содержат два разных отступа')
def conflict023(c):
    c.examples023=[{'kind':'paragraph','when':{'role':'body'},'set':{'style':s},'id':str(i)} for i,s in enumerate(['a','a','b'])]

@when('агент компилирует правила выбора оформления')
def fit023(c):
    from articleflow.rules023 import fit
    c.rules023,c.conflicts023=fit(c.examples023)

@then('контекст остаётся неоднозначным без правила большинства')
def ambiguous023(c):
    assert not c.rules023 and len(c.conflicts023)==1

@given('малая модель назначила роль только первому из двух абзацев')
def incomplete023(c):
    c.response023={'groups':[{'start':1,'end':1,'role':'body'}]}

@when('проверяется полное покрытие смыслового плана')
def coverage023(c):
    from articleflow.semantic023 import expand_groups
    from articleflow.docx_engine import DocumentError
    try:expand_groups(c.response023,[{'id':'a'},{'id':'b'}],{'body':'Основной текст'})
    except DocumentError as e:c.error023=str(e)
    else:c.error023=None

@then('неполный план отклонён до сборки')
def rejected023(c):
    assert c.error023=='SEMANTIC_GROUP_COVERAGE'

@given('образец задаёт порядок название аннотация ключевые слова')
def reference_order024(c):
    setup(c)

@when('вход с неверным порядком обрабатывается с правильными ролями')
def apply_order024(c):
    import importlib.util
    path=Path(__file__).resolve().parents[4]/'scripts/probe_agent023_order.py'
    spec=importlib.util.spec_from_file_location('probe_order024',path)
    module=importlib.util.module_from_spec(spec);spec.loader.exec_module(module)
    c.order024=module.probe(c.work019/'order024')

@then('порядок соответствует образцу без изменения исходного содержания')
def assert_order024(c):
    assert c.order024['target_order_repaired']
    assert c.order024['status']=='rendered_unverified'
    assert c.order024['unresolved']==0

@given('цикл извлечения находится в состоянии выполнения')
def learning_live024(c):
    setup(c);c.app024=create_app(c.work019/'web024');store=c.app024.state.store
    store.execute('INSERT INTO templates VALUES (?,?,?)',('synthetic','Synthetic','now'))
    store.execute('INSERT INTO template_runs VALUES (?,?,?,?,?,?)',('run','synthetic',1,'running',json.dumps({'schema':'teacher-student-023'}),'now'))

@when('пользователь сохраняет уточнение обязательности блока')
def feedback024(c):
    with TestClient(c.app024) as client:
        headers={'X-CSRF-Token':client.get('/api/v017/config').json()['csrf_token']}
        body={'expected_revision':0,'text':'Сведения о финансировании необязательны'}
        response=client.post('/api/v017/template-runs/run/feedback',json=body,headers=headers)
        assert response.status_code==200
        c.feedback024=response.json()
        c.stale024=client.post('/api/v017/template-runs/run/feedback',json=body,headers=headers).status_code

@then('уточнение принято для следующей итерации и защищено номером версии')
def assert_feedback024(c):
    assert c.feedback024=={'revision':1,'applies':'next_iteration','already_applied':False}
    assert c.stale024==409

@given('модель выделила переменную с несовпадающей цитатой')
def inline_bad026(c):
    c.inline026={'inline':[{'paragraph':1,'start':0,'end':1,'role':'variable','quote':'y'}]}

@when('проверяются диапазоны смысловых ролей внутри абзаца')
def validate_inline026(c):
    from articleflow.inline026 import validate
    from articleflow.docx_engine import DocumentError
    try:validate(c.inline026,[{'id':'b','text':'x'}],{'variable':'Обозначение переменной'})
    except DocumentError as exc:c.inline_error026=str(exc)
    else:c.inline_error026=None

@then('план отклонён без исправления цитаты и изменения текста')
def assert_inline026(c):
    assert c.inline_error026=='INLINE_QUOTE_MISMATCH'
    assert c.inline026['inline'][0]['quote']=='y'
